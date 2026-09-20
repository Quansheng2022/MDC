"""P11-MNT-007 / ISSUE-001: Markdown link targets must be preserved in DOCX.

Focused tests for the renderer → writer hyperlink path. No corpus re-run, no
network access: every document is compiled from an in-memory Markdown snippet.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

from md_converter.compiler import compile_markdown


def _external_hyperlink_targets(doc: Document) -> list[str]:
    """Return every external hyperlink target stored in the package."""
    return sorted(
        rel.target_ref
        for rel in doc.part.rels.values()
        if rel.is_external and rel.reltype == RT.HYPERLINK
    )


def _hyperlink_elements(doc: Document):
    """Yield every ``w:hyperlink`` element in document order."""
    yield from doc.element.body.iter(qn("w:hyperlink"))


def _anchor_targets(doc: Document) -> list[str]:
    """Return every internal anchor target (``w:anchor``) in document order."""
    return [hl.get(qn("w:anchor")) for hl in _hyperlink_elements(doc) if hl.get(qn("w:anchor"))]


def _compile(text: str, out: Path) -> Document:
    """Compile ``text`` into ``out`` and reopen the result."""
    compile_markdown(
        text,
        config={"word_com": False, "enable_cover": False, "toc": False},
        output_path=out,
    )
    return Document(str(out))


def test_ut_render_link_001_link_becomes_real_hyperlink(temp_dir: Path) -> None:
    """UT-RENDER-LINK-001: Link → w:hyperlink + external relationship."""
    doc = _compile("See [OpenAI](https://openai.com) now.\n", temp_dir / "ut_link.docx")

    assert _external_hyperlink_targets(doc) == ["https://openai.com"]

    hyperlinks = list(_hyperlink_elements(doc))
    assert len(hyperlinks) == 1
    hyperlink = hyperlinks[0]
    assert hyperlink.get(qn("r:id"))
    assert "OpenAI" in "".join(t.text or "" for t in hyperlink.iter(qn("w:t")))

    # 可见文本与既有样式保持不变（蓝色 + 下划线）
    paragraph = next(p for p in doc.paragraphs if "OpenAI" in p.text)
    assert paragraph.text == "See OpenAI now."
    run = hyperlink.find(qn("w:r"))
    rpr = run.find(qn("w:rPr"))
    assert rpr is not None
    underline = rpr.find(qn("w:u"))
    assert underline is not None and underline.get(qn("w:val")) == "single"
    color = rpr.find(qn("w:color"))
    assert color is not None and color.get(qn("w:val")).upper() == "0000FF"


def test_it_link_001_http_https_mailto_and_anchor_targets_preserved(temp_dir: Path) -> None:
    """IT-LINK-001: http/https/mailto/anchor targets are all recoverable."""
    markdown = (
        "# Links\n\n"
        "- [http link](http://example.com/a)\n"
        "- [https link](https://example.com/b)\n"
        "- [mail](mailto:test@example.com)\n"
        "- [relative](./other.md)\n"
        "- [anchor](#links)\n"
    )
    doc = _compile(markdown, temp_dir / "it_link.docx")

    assert _external_hyperlink_targets(doc) == [
        "./other.md",
        "http://example.com/a",
        "https://example.com/b",
        "mailto:test@example.com",
    ]
    assert _anchor_targets(doc) == ["links"]

    # 链接文本仍然出现在文档文本流中（No Content Loss）
    text = "\n".join(p.text for p in doc.paragraphs)
    for needle in ("http link", "https link", "mail", "relative", "anchor"):
        assert needle in text


def test_ut_render_link_002_empty_target_keeps_text_without_hyperlink(temp_dir: Path) -> None:
    """UT-RENDER-LINK-002: a degenerate empty target must not create a relation."""
    doc = _compile("See [nothing]() here.\n", temp_dir / "ut_link_empty.docx")

    assert _external_hyperlink_targets(doc) == []
    assert list(_hyperlink_elements(doc)) == []
    assert any("nothing" in p.text for p in doc.paragraphs)
