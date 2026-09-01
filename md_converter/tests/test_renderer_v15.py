"""Integration tests for V1.5 renderer behavior."""

# ruff: noqa: I001  # 与项目 [tool.isort] 配置（multi_line_output=3）排序结果不一致

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from md_converter.ast.nodes import (
    CodeBlock,
)
from md_converter.ast.nodes import Document as AstDocument
from md_converter.ast.nodes import (
    HardBreak,
    Heading,
    Paragraph,
    SoftBreak,
    Table,
    TableCell,
    TableRow,
    Text,
)
from md_converter.renderer.render_context import RenderContext
from md_converter.renderer.themes.v15_theme import V15Theme
from md_converter.renderer.word_renderer import WordRenderer
from md_converter.renderer.word_writer import WordWriter


def _render(document: AstDocument) -> Document:
    writer = WordWriter.create()
    renderer = WordRenderer(RenderContext.create(theme=V15Theme.load_default()), writer)
    return renderer.render(document)


def _run_fonts(run) -> dict:
    rPr = run._element.rPr
    rFonts = rPr.rFonts if rPr is not None else None
    return {
        "ascii": rFonts.get(qn("w:ascii")) if rFonts is not None else None,
        "eastAsia": rFonts.get(qn("w:eastAsia")) if rFonts is not None else None,
    }


def test_body_font_mapping_applied() -> None:
    """Body runs use Calibri (ascii) and DengXian (eastAsia)."""
    doc = _render(AstDocument(children=[Paragraph(content=[Text("Hello 世界")])]))

    run = doc.paragraphs[0].runs[0]
    fonts = _run_fonts(run)
    assert fonts["ascii"] == "Calibri"
    assert fonts["eastAsia"] == "DengXian"


def test_body_font_size_is_10_5() -> None:
    """Body text renders at 10.5pt per the frozen spec."""
    doc = _render(AstDocument(children=[Paragraph(content=[Text("正文")])]))

    run = doc.paragraphs[0].runs[0]
    assert run.font.size.pt == 10.5


def test_cjk_paragraph_left_aligned() -> None:
    """Chinese paragraphs align left (not justified)."""
    doc = _render(AstDocument(children=[Paragraph(content=[Text("中文段落两端对齐测试")])]))

    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_latin_paragraph_left_aligned() -> None:
    """Latin paragraphs align left."""
    doc = _render(AstDocument(children=[Paragraph(content=[Text("English paragraph")])]))

    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_line_break_paragraph_left_aligned() -> None:
    """Line-broken label blocks (e.g. metadata lines) align left, not justified."""
    doc = _render(
        AstDocument(
            children=[
                Paragraph(
                    content=[
                        Text("模式: DEEP RESEARCH"),
                        SoftBreak(),
                        Text("报告日期: 2026-08-10"),
                    ]
                )
            ]
        )
    )

    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_heading_keep_with_next() -> None:
    """Headings keep with the next paragraph."""
    doc = _render(
        AstDocument(
            children=[
                Heading(level=1, content=[Text("标题")]),
                Paragraph(content=[Text("正文")]),
            ]
        )
    )

    heading = doc.paragraphs[0]
    assert heading.style.name == "Heading 1"
    assert heading.paragraph_format.keep_with_next is True
    assert heading.paragraph_format.keep_together is True


def test_heading_style_sizes() -> None:
    """Heading styles follow frozen sizes."""
    doc = _render(
        AstDocument(
            children=[
                Heading(level=1, content=[Text("H1")]),
                Heading(level=4, content=[Text("H4")]),
            ]
        )
    )

    assert doc.styles["Heading 1"].font.size.pt == 20
    assert doc.styles["Heading 2"].font.size.pt == 16
    assert doc.styles["Heading 3"].font.size.pt == 14
    assert doc.styles["Heading 4"].font.size.pt == 14
    assert doc.styles["Heading 4"].font.italic is True


def test_table_repeat_header_and_font_size() -> None:
    """Tables repeat headers across pages and use 9.5pt."""
    table = Table(
        rows=[
            TableRow(cells=[TableCell(content=[Text("Head")])]),
            TableRow(cells=[TableCell(content=[Text("Cell")])]),
        ]
    )
    doc = _render(AstDocument(children=[table]))

    docx_table = doc.tables[0]
    trPr = docx_table.rows[0]._tr.trPr
    assert trPr is not None
    assert trPr.find(qn("w:tblHeader")) is not None

    cell_run = docx_table.rows[1].cells[0].paragraphs[0].runs[0]
    assert cell_run.font.size.pt == 9.5


def test_number_column_right_aligned() -> None:
    """Numeric table columns align right (chapter 9.2)."""
    table = Table(
        rows=[
            TableRow(cells=[TableCell(content=[Text("Name")]), TableCell(content=[Text("Value")])]),
            TableRow(cells=[TableCell(content=[Text("A")]), TableCell(content=[Text("1,250.5")])]),
        ]
    )
    doc = _render(AstDocument(children=[table]))

    value_cell = doc.tables[0].rows[1].cells[1]
    assert value_cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_code_block_style() -> None:
    """Code blocks use Consolas 10pt with light background."""
    doc = _render(
        AstDocument(
            children=[
                CodeBlock(
                    language="python",
                    text="def hello():\n    print('hi')",
                )
            ]
        )
    )

    run = doc.paragraphs[0].runs[0]
    assert run.font.size.pt == 10
    assert _run_fonts(run)["ascii"] == "Consolas"
    assert doc.paragraphs[0].paragraph_format.keep_together is True


def test_ascii_block_left_aligned_monospace() -> None:
    """ASCII diagrams render left-aligned at 8.5pt Consolas."""
    doc = _render(
        AstDocument(
            children=[
                CodeBlock(
                    language="",
                    text="┌─────┐\n│  A  │\n└─────┘",
                )
            ]
        )
    )

    p = doc.paragraphs[0]
    assert p.alignment == WD_ALIGN_PARAGRAPH.LEFT
    run = p.runs[0]
    assert run.font.size.pt == 8.5
    assert _run_fonts(run)["ascii"] == "Consolas"


def test_page_margins_one_inch() -> None:
    """Page margins are 1in per the frozen spec."""
    doc = _render(AstDocument(children=[Paragraph(content=[Text("x")])]))

    section = doc.sections[0]
    assert section.top_margin.inches == 1.0
    assert section.bottom_margin.inches == 1.0
    assert section.left_margin.inches == 1.0
    assert section.right_margin.inches == 1.0


def test_hard_break_renders_as_line_break() -> None:
    """Markdown 硬换行应渲染为行内换行（w:br），而非分页符。"""
    doc = _render(
        AstDocument(children=[Paragraph(content=[Text("第一行"), HardBreak(), Text("第二行")])])
    )
    p = doc.paragraphs[0]
    br_types = [br.get(qn("w:type")) for br in p._element.iter(qn("w:br"))]
    assert br_types == [None]
