"""Table fitting renderer integration (Program D / WP-D03).

Proves the D-02 policy is applied through the existing conversion path
(``CompilerContext.compile`` -> Renderer -> ``WordWriter``), end to end:

* rendered tables now carry deterministic explicit column widths instead of
  relying on Word autofit;
* cell values, row/column counts and the Grid-family style are preserved;
* wide tables stay inside the effective content width of the section that
  actually renders them;
* profile geometry reaches the fitting decision without any profile-ID
  branching;
* an irreducible wide table keeps every cell and reports ``RENDER006``;
* repeated conversion is deterministic.
"""

from __future__ import annotations

import contextlib
import io
from pathlib import Path
from typing import Any, Dict, List, Tuple

import pytest
from docx import Document

from md_converter.compiler import CompilerContext

#: Effective content width of A4 portrait with the frozen 1 in margins.
BASELINE_CONTENT_WIDTH_CM = 21.0 - 2.54 - 2.54

#: Effective content widths produced by two Professional Output Profiles.
BUSINESS_CONTENT_WIDTH_CM = 21.0 - 2.0 - 2.0
ACADEMIC_CONTENT_WIDTH_CM = 21.0 - 3.0 - 3.0

BASE_CONFIG: Dict[str, Any] = {
    "word_com": False,
    "enable_cover": False,
    "toc": False,
    "style_tables": True,
    "verbose": False,
}

NORMAL_TABLE = (
    "# Fitting\n\n"
    "| Identifier | Description | Owner | Status |\n"
    "|------------|-------------|-------|--------|\n"
    "| DOC-0001 | Pipeline normalisation pass | platform | accepted |\n"
    "| DOC-0002 | Renderer table fitting | renderer | accepted |\n"
)

#: A denser table that fills the content width of every portfolio profile.
PROFILE_TABLE = (
    "# Profile fitting\n\n"
    "| Identifier | Description | Owner | Status |\n"
    "|------------|-------------|-------|--------|\n"
    "| DOC-0001 | Pipeline normalisation pass for renderer tables | platform | accepted |\n"
    "| DOC-0002 | Deterministic column width allocation | renderer | accepted |\n"
)


def _table_markdown(columns: int, cell_text: str) -> str:
    """Build a Markdown table with ``columns`` columns of ``cell_text``."""
    header = "|" + "|".join(f" H{index} " for index in range(columns)) + "|"
    separator = "|" + "|".join(" --- " for _ in range(columns)) + "|"
    row = "|" + "|".join(f" {cell_text} " for _ in range(columns)) + "|"
    return "\n".join(["# Wide", "", header, separator, row, ""])


def _compile(
    body: str, tmp_path: Path, name: str, profile_id: str | None = None
) -> Tuple[CompilerContext, Path]:
    """Compile ``body`` through the public compiler context."""
    config = dict(BASE_CONFIG)
    if profile_id is not None:
        config["output_profile"] = profile_id
    context = CompilerContext.create(config)
    output = tmp_path / f"{name}.docx"
    with contextlib.redirect_stdout(io.StringIO()):
        context.compile(body, {"title": name}, output)
    return context, output


def _grid_widths_cm(table: Any) -> List[float]:
    """Return the explicit ``w:tblGrid`` column widths of a table (cm)."""
    return [column.width.cm for column in table.columns]


def _first_table(path: Path) -> Any:
    """Return the first table of a rendered artifact."""
    document = Document(str(path))
    assert document.tables, "expected the rendered artifact to contain a table"
    return document.tables[0]


def _codes(context: CompilerContext) -> List[str]:
    """Return every diagnostic code raised by a conversion."""
    return [diagnostic.code for diagnostic in context.diag.diagnostics]


def _cell_matrix(table: Any) -> List[List[str]]:
    """Return the visible cell text matrix of a table."""
    return [[cell.text for cell in row.cells] for row in table.rows]


def test_normal_table_is_fitted_to_the_effective_content_width(tmp_path: Path) -> None:
    """Representative normal table: explicit widths filling the content width."""
    context, output = _compile(NORMAL_TABLE, tmp_path, "normal")
    table = _first_table(output)

    assert table.autofit is False
    widths = _grid_widths_cm(table)

    assert len(widths) == 4
    assert sum(widths) == pytest.approx(BASELINE_CONTENT_WIDTH_CM, abs=0.01)
    assert min(widths) >= 1.2 - 1e-9
    assert "RENDER006" not in _codes(context)


def test_normal_table_content_and_structure_are_unchanged(tmp_path: Path) -> None:
    """Fitting must not touch cell values, rows or columns."""
    _, output = _compile(NORMAL_TABLE, tmp_path, "meaning")
    table = _first_table(output)

    assert _cell_matrix(table) == [
        ["Identifier", "Description", "Owner", "Status"],
        ["DOC-0001", "Pipeline normalisation pass", "platform", "accepted"],
        ["DOC-0002", "Renderer table fitting", "renderer", "accepted"],
    ]
    assert "Grid" in table.style.name


def test_wide_table_is_fitted_within_its_effective_content_width(tmp_path: Path) -> None:
    """8-column wide table: total width inside the content box, floor respected."""
    context, output = _compile(_table_markdown(8, "value0001"), tmp_path, "wide8")
    table = _first_table(output)

    widths = _grid_widths_cm(table)
    assert len(widths) == 8
    assert sum(widths) == pytest.approx(BASELINE_CONTENT_WIDTH_CM, abs=0.01)
    assert min(widths) >= 1.2 - 1e-9
    assert context.rendered_qa_result is not None
    assert context.rendered_qa_result.metrics["table_overflow"] == 0


def test_long_text_column_receives_more_width(tmp_path: Path) -> None:
    """Content-aware allocation: the long column is wider in the artifact."""
    body = (
        "# Proportions\n\n"
        "| ID | Explanation |\n"
        "|----|-------------|\n"
        f"| 1 | {'detail ' * 8}| \n"
    )
    _, output = _compile(body, tmp_path, "proportion")
    widths = _grid_widths_cm(_first_table(output))

    assert widths[1] > widths[0]
    assert widths[0] >= 1.2 - 1e-9


def test_irreducible_wide_table_preserves_content_and_warns(tmp_path: Path) -> None:
    """Irreducible wide table: content kept, floor kept, RENDER006 recorded."""
    context, output = _compile(_table_markdown(25, "a"), tmp_path, "irreducible")
    table = _first_table(output)
    widths = _grid_widths_cm(table)

    assert len(widths) == 25
    assert all(width == pytest.approx(1.2, abs=1e-3) for width in widths)
    assert _cell_matrix(table)[0][0].strip() == "H0"

    codes = _codes(context)
    assert "RENDER006" in codes
    assert "RENDER007" not in codes

    report = context.get_quality_gate_report()
    for stage in ("static_qa", "rendered_qa", "post_processor", "final_artifact_qa"):
        assert (report.get(stage) or {}).get("status") != "FAIL", stage


def test_irreducible_wide_table_uses_the_bounded_font_step(tmp_path: Path) -> None:
    """The bounded compactness step lowers the table font, never below 8.5pt."""
    _, output = _compile(_table_markdown(25, "a"), tmp_path, "font_step")
    table = _first_table(output)
    header_run = table.rows[0].cells[0].paragraphs[0].runs[0]

    assert header_run.font.size.pt == pytest.approx(9.0)
    assert header_run.font.size.pt >= 8.5


def test_repeated_conversion_is_deterministic(tmp_path: Path) -> None:
    """Same input and config -> identical fitted widths (SPEC-INV-003)."""
    _, first = _compile(NORMAL_TABLE, tmp_path, "det1")
    _, second = _compile(NORMAL_TABLE, tmp_path, "det2")

    assert _grid_widths_cm(_first_table(first)) == _grid_widths_cm(_first_table(second))
    assert _cell_matrix(_first_table(first)) == _cell_matrix(_first_table(second))


@pytest.mark.parametrize(
    ("profile_id", "expected_content_width_cm"),
    [
        ("professional_report", BASELINE_CONTENT_WIDTH_CM),
        ("business_report", BUSINESS_CONTENT_WIDTH_CM),
        ("academic", ACADEMIC_CONTENT_WIDTH_CM),
    ],
)
def test_profile_geometry_drives_fitting_without_profile_branching(
    tmp_path: Path, profile_id: str, expected_content_width_cm: float
) -> None:
    """Profile differences reach the fitting decision as resolved geometry only."""
    _, output = _compile(PROFILE_TABLE, tmp_path, profile_id, profile_id=profile_id)
    table = _first_table(output)

    assert sum(_grid_widths_cm(table)) == pytest.approx(expected_content_width_cm, abs=0.01)
    # Meaning is independent of the presentation profile.
    assert _cell_matrix(table)[1][0] == "DOC-0001"


def test_narrow_table_is_not_stretched_to_the_full_content_width(tmp_path: Path) -> None:
    """A short two-column table keeps its natural width instead of filling the page."""
    body = "# Narrow\n\n| K | V |\n|---|---|\n| a | b |\n"
    _, output = _compile(body, tmp_path, "narrow")
    widths = _grid_widths_cm(_first_table(output))

    assert sum(widths) < BASELINE_CONTENT_WIDTH_CM * 0.6
    assert all(width == pytest.approx(1.2, abs=1e-3) for width in widths)


def test_fitted_tables_keep_the_grid_style_gate_green(tmp_path: Path) -> None:
    """FinalArtifactQA table-styling gate stays green after fitting."""
    context, output = _compile(NORMAL_TABLE, tmp_path, "gate")
    table = _first_table(output)

    assert "Grid" in table.style.name
    assert table.rows, "header row must be preserved"
    report = context.get_quality_gate_report()
    assert (report.get("final_artifact_qa") or {}).get("status") != "FAIL"
