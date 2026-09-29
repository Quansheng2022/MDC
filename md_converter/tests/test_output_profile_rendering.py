"""Renderer/presentation integration for output profiles (Program C, WP-POP-04).

Proves the frozen architecture end to end without a GUI:

* every profile passes the full quality gate chain on a representative
  document (plain text, headings, lists, table, code block);
* the document meaning is identical across profiles - only presentation
  differs;
* the default profile produces an artifact byte-identical to the accepted
  baseline (product specification 7: no silent output-wide restyling);
* each declared profile value actually reaches the artifact;
* no renderer module knows about profiles, so profile logic cannot spread as
  duplicated renderer branching.
"""

from __future__ import annotations

import ast
import contextlib
import hashlib
import io
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Tuple

import pytest
from docx import Document
from docx.oxml.ns import qn

from md_converter.compiler import CompilerContext
from md_converter.profiles import ALL_PROFILES, DEFAULT_PROFILE_ID, profile_ids, resolve_profile

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RENDERER_DIR = PROJECT_ROOT / "md_converter" / "renderer"

#: Representative document (product specification 17: the representative matrix).
REPRESENTATIVE_MARKDOWN = """---
title: Representative Document
author: MD Converter
date: 2026-09-29
---

# Section One

Plain body text for the profile matrix.

Body text with **bold**, *italic* and `inline code` for a mixed document.

## Section Two

- First item
- Second item

| Metric | Value |
|--------|-------|
| Alpha  | 12.5  |
| Beta   | 8.25  |

```python
def hello():
    return "world"
```
"""

BODY_PARAGRAPH_PREFIX = "Plain body text"
CODE_PARAGRAPH_PREFIX = "def hello"

#: Configuration keeping the matrix deterministic and free of optional Word
#: automation (the accepted regression configuration).
_BASE_CONFIG: Dict[str, Any] = {
    "word_com": False,
    "enable_cover": False,
    "toc": False,
    "style_tables": True,
    "verbose": False,
}


def _compile(tmp_path: Path, profile_id: str | None, name: str) -> Tuple[CompilerContext, Any]:
    """Compile the representative document for ``profile_id``."""
    config = dict(_BASE_CONFIG)
    if profile_id is not None:
        config["output_profile"] = profile_id
    context = CompilerContext.create(config)
    output = tmp_path / f"{name}.docx"
    with contextlib.redirect_stdout(io.StringIO()):
        context.compile(REPRESENTATIVE_MARKDOWN, {"title": "Representative Document"}, output)
    return context, output


def _find_paragraph(doc: Any, prefix: str) -> Any:
    """Return the first paragraph whose text starts with ``prefix``."""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise AssertionError(f"no paragraph starts with {prefix!r}")


def _run_sizes(paragraph: Any) -> List[float]:
    """Return the explicit font sizes of a paragraph's runs."""
    return [run.font.size.pt for run in paragraph.runs if run.font.size is not None]


def _style_fonts(doc: Any, style_name: str) -> Dict[str, str]:
    """Return the ``w:rFonts`` slots of one Word style (empty when unset)."""
    style = doc.styles[style_name]
    r_pr = style._element.find(qn("w:rPr"))
    r_fonts = r_pr.find(qn("w:rFonts")) if r_pr is not None else None
    if r_fonts is None:
        return {}
    return {
        slot: r_fonts.get(qn(f"w:{slot}"))
        for slot in ("ascii", "hAnsi", "eastAsia", "cs")
        if r_fonts.get(qn(f"w:{slot}")) is not None
    }


def _artifact_facts(path: Path) -> Dict[str, Any]:
    """Return the presentation and meaning facts of one compiled artifact."""
    doc = Document(str(path))
    section = doc.sections[0]
    table = doc.tables[0]
    header_run = table.rows[0].cells[0].paragraphs[0].runs[0]
    return {
        "margins_cm": (
            round(section.top_margin.cm, 3),
            round(section.bottom_margin.cm, 3),
            round(section.left_margin.cm, 3),
            round(section.right_margin.cm, 3),
        ),
        "normal_size": doc.styles["Normal"].font.size.pt,
        "heading_sizes": [doc.styles[f"Heading {level}"].font.size.pt for level in range(1, 5)],
        "heading_space_before": doc.styles["Heading 1"].paragraph_format.space_before.pt,
        "heading_space_after": doc.styles["Heading 1"].paragraph_format.space_after.pt,
        "normal_fonts": _style_fonts(doc, "Normal"),
        "heading_fonts": _style_fonts(doc, "Heading 1"),
        "table_style": table.style.name,
        "table_font_size": header_run.font.size.pt,
        "body_run_sizes": _run_sizes(_find_paragraph(doc, BODY_PARAGRAPH_PREFIX)),
        "code_run_sizes": _run_sizes(_find_paragraph(doc, CODE_PARAGRAPH_PREFIX)),
        "body_line_spacing": _find_paragraph(
            doc, BODY_PARAGRAPH_PREFIX
        ).paragraph_format.line_spacing,
        "body_space_after": _find_paragraph(
            doc, BODY_PARAGRAPH_PREFIX
        ).paragraph_format.space_after.pt,
        "meaning_paragraphs": [paragraph.text for paragraph in doc.paragraphs],
        "meaning_table_cells": [cell.text for row in table.rows for cell in row.cells],
        "meaning_inline_shapes": len(doc.inline_shapes),
    }


def _meaning(facts: Dict[str, Any]) -> Dict[str, Any]:
    """Return the document-meaning subset of ``facts``."""
    return {
        "paragraphs": facts["meaning_paragraphs"],
        "table_cells": facts["meaning_table_cells"],
        "inline_shapes": facts["meaning_inline_shapes"],
    }


def _presentation(facts: Dict[str, Any]) -> Dict[str, Any]:
    """Return the presentation subset of ``facts``."""
    return {key: value for key, value in facts.items() if not key.startswith("meaning_")}


def _package_digest(path: Path) -> str:
    """Return a digest of the document *parts* of a DOCX package.

    The raw file bytes of a DOCX are not stable across saves: the ZIP container
    records the local write time per entry, so two identical documents can
    differ byte-for-byte.  Comparing the package parts is the deterministic
    form of "the artifact is identical", which is what WP-POP-04 requires.
    """
    with zipfile.ZipFile(path) as archive:
        parts = sorted(
            (name, hashlib.sha256(archive.read(name)).hexdigest()) for name in archive.namelist()
        )
    return hashlib.sha256(repr(parts).encode("utf-8")).hexdigest()


@pytest.fixture
def matrix(tmp_path: Path) -> Dict[str, Any]:
    """Compile the representative document for the baseline and every profile."""
    compiled: Dict[str, Any] = {}
    _, baseline_path = _compile(tmp_path, None, "baseline")
    compiled["baseline"] = {
        "path": baseline_path,
        "facts": _artifact_facts(baseline_path),
        "package": _package_digest(baseline_path),
    }
    for profile in ALL_PROFILES:
        context, path = _compile(tmp_path, profile.id, profile.id)
        compiled[profile.id] = {
            "path": path,
            "facts": _artifact_facts(path),
            "package": _package_digest(path),
            "report": context.get_quality_gate_report(),
            "codes": [diagnostic.code for diagnostic in context.diag.diagnostics],
        }
    return compiled


# ============================================================
# Quality gates
# ============================================================


def test_every_profile_passes_every_quality_gate(matrix: Dict[str, Any]) -> None:
    """Presentation differences must never break the QA chain."""
    for name, entry in matrix.items():
        if name == "baseline":
            continue
        report = entry["report"]
        for stage in ("static_qa", "rendered_qa", "post_processor", "final_artifact_qa"):
            status = (report.get(stage) or {}).get("status")
            assert status != "FAIL", (name, stage, status)
        assert entry["codes"] == [], (name, entry["codes"])


def test_readability_floors_hold_for_every_profile(matrix: Dict[str, Any]) -> None:
    """No profile may render text below the absolute 8pt floor."""
    for name, entry in matrix.items():
        facts = entry["facts"]
        sizes = [
            facts["normal_size"],
            facts["table_font_size"],
            *facts["heading_sizes"],
            *facts["body_run_sizes"],
            *facts["code_run_sizes"],
        ]
        assert min(sizes) >= 8.0, (name, sizes)


# ============================================================
# Meaning preservation and default compatibility
# ============================================================


def test_document_meaning_is_identical_across_profiles(matrix: Dict[str, Any]) -> None:
    """Product specification 14: profiles are presentation configuration only."""
    baseline = _meaning(matrix["baseline"]["facts"])

    for name, entry in matrix.items():
        if name == "baseline":
            continue
        assert _meaning(entry["facts"]) == baseline, name


def test_default_profile_artifact_is_identical_to_the_baseline(
    matrix: Dict[str, Any],
) -> None:
    """Product specification 7: no silent output-wide restyling."""
    assert matrix[DEFAULT_PROFILE_ID]["package"] == matrix["baseline"]["package"]


def test_every_non_default_profile_changes_the_artifact(matrix: Dict[str, Any]) -> None:
    """Product specification 17: output differences are intentional and visible."""
    baseline = _presentation(matrix["baseline"]["facts"])

    for name, entry in matrix.items():
        if name in ("baseline", DEFAULT_PROFILE_ID):
            continue
        presentation = _presentation(entry["facts"])
        differences = [key for key, value in presentation.items() if value != baseline[key]]
        assert len(differences) >= 2, (name, differences)
        assert entry["package"] != matrix["baseline"]["package"], name


# ============================================================
# Declared values reach the artifact
# ============================================================


@pytest.mark.parametrize("profile_id", profile_ids())
def test_declared_profile_values_reach_the_artifact(
    matrix: Dict[str, Any], profile_id: str
) -> None:
    """Every bounded property the profile declares is observable in the DOCX."""
    profile = resolve_profile(profile_id)
    typography = profile.presentation.typography
    margins = profile.presentation.page_margins
    facts = matrix[profile_id]["facts"]

    assert facts["margins_cm"] == pytest.approx(
        (margins.top_cm, margins.bottom_cm, margins.left_cm, margins.right_cm), abs=0.01
    )
    assert facts["normal_size"] == pytest.approx(typography.body_size_pt)
    assert facts["body_run_sizes"], profile_id
    for size in facts["body_run_sizes"]:
        assert size == pytest.approx(typography.body_size_pt, abs=0.01), profile_id
    assert facts["body_line_spacing"] == pytest.approx(typography.line_spacing, abs=0.01)
    assert facts["body_space_after"] == pytest.approx(typography.paragraph_space_after_pt, abs=0.01)
    assert facts["heading_sizes"] == pytest.approx(list(typography.heading_sizes_pt))
    assert facts["heading_space_before"] == pytest.approx(typography.heading_space_before_pt)
    assert facts["heading_space_after"] == pytest.approx(typography.heading_space_after_pt)
    assert facts["table_style"] == typography.table_style
    assert facts["table_font_size"] == pytest.approx(typography.table_font_size_pt)
    assert facts["code_run_sizes"] == pytest.approx([typography.code_size_pt], abs=0.01)
    assert facts["normal_fonts"]["ascii"] == typography.body_font
    assert facts["normal_fonts"]["eastAsia"] == typography.body_east_asia_font
    assert facts["normal_fonts"]["cs"] == typography.body_font_complex
    assert facts["heading_fonts"]["ascii"] == typography.heading_font
    assert facts["heading_fonts"]["eastAsia"] == typography.heading_east_asia_font


def test_table_style_stays_in_the_grid_family(matrix: Dict[str, Any]) -> None:
    """FinalArtifactQA requires the Grid family while styling tables."""
    for name, entry in matrix.items():
        if name == "baseline":
            continue
        assert "Grid" in entry["facts"]["table_style"], name


# ============================================================
# Architecture guard: no renderer branching
# ============================================================


def test_renderer_never_references_profiles() -> None:
    """The renderer consumes a resolved theme, never a profile identity."""
    modules = sorted(RENDERER_DIR.rglob("*.py"))

    assert modules
    for module in modules:
        source = module.read_text(encoding="utf-8")
        assert "output_profile" not in source, module.name
        assert "theme_presentation_overrides" not in source, module.name
        tree = ast.parse(source)
        for node in ast.walk(tree):
            if isinstance(node, ast.ImportFrom):
                assert "profiles" not in (node.module or ""), module.name
            elif isinstance(node, ast.Import):
                assert all("profiles" not in alias.name for alias in node.names), module.name
