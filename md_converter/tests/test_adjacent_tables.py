"""P11-MNT-009 / ISSUE-004: adjacent Markdown tables must stay separate.

Word merges directly adjacent ``w:tbl`` elements on open/save, which destroyed
the boundary between two blank-line-separated source tables in 5 real Cycle-01
documents. These tests pin the separator, the Word COM round trip and the
FinalArtifactQA structural signal.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from md_converter.compiler import CompilerContext, compile_markdown
from md_converter.renderer.word_writer import WordWriter
from md_converter.tests.test_word_com_final_artifact import _word_com_probe

ADJACENT_TABLES = (
    "# Tables\n\n"
    "| A | B |\n"
    "| --- | --- |\n"
    "| 1 | 2 |\n"
    "\n"
    "| C | D |\n"
    "| --- | --- |\n"
    "| 3 | 4 |\n"
)

SEPARATED_TABLES = (
    "# Tables\n\n"
    "| A | B |\n"
    "| --- | --- |\n"
    "| 1 | 2 |\n"
    "\n"
    "中间段落\n"
    "\n"
    "| C | D |\n"
    "| --- | --- |\n"
    "| 3 | 4 |\n"
)


def _body_sequence(docx_path: Path) -> list[str]:
    """Return the local names of the body's content elements (sectPr excluded)."""
    from lxml import etree

    with zipfile.ZipFile(docx_path) as zf:
        body = etree.fromstring(zf.read("word/document.xml")).find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body"
        )
    return [
        etree.QName(child).localname for child in body if etree.QName(child).localname != "sectPr"
    ]


def _adjacent_table_pairs(doc: Document) -> list[int]:
    """Return the positions of directly adjacent body-level tables."""
    content = [child for child in doc.element.body if child.tag != qn("w:sectPr")]
    return [
        idx
        for idx in range(len(content) - 1)
        if content[idx].tag == qn("w:tbl") and content[idx + 1].tag == qn("w:tbl")
    ]


@pytest.fixture(scope="module")
def word_com_ready() -> None:
    """Runtime gate for the Word COM round trip (Windows release gate)."""
    state, detail = _word_com_probe()
    if state == "skip":
        pytest.skip(detail)
    if state == "fail":
        pytest.fail(detail, pytrace=False)
    assert state == "ok", f"Unexpected Word COM probe state: {state!r}"


def test_ut_table_adj_001_adjacent_tables_get_stable_separator(temp_dir: Path) -> None:
    """UT-TABLE-ADJ-001: two adjacent source tables → two w:tbl + separator."""
    out = temp_dir / "adjacent.docx"
    ctx = CompilerContext.create({"word_com": False, "enable_cover": False, "toc": False})
    ctx.compile(ADJACENT_TABLES, {}, out)

    sequence = _body_sequence(out)
    assert sequence.count("tbl") == 2
    assert sequence == ["p", "tbl", "p", "tbl"]

    doc = Document(str(out))
    assert _adjacent_table_pairs(doc) == []
    assert ctx.final_artifact_qa_result is not None
    assert ctx.final_artifact_qa_result.metrics.get("adjacent_tables") == 0
    assert all(issue["code"] != "table_adjacent" for issue in ctx.final_artifact_qa_result.warnings)


def test_ut_table_adj_002_paragraph_separated_tables_unchanged(temp_dir: Path) -> None:
    """UT-TABLE-ADJ-002: tables already separated by content gain no paragraph."""
    out = temp_dir / "separated.docx"
    compile_markdown(
        SEPARATED_TABLES,
        config={"word_com": False, "enable_cover": False, "toc": False},
        output_path=out,
    )

    sequence = _body_sequence(out)
    assert sequence == ["p", "tbl", "p", "tbl"]
    assert _adjacent_table_pairs(Document(str(out))) == []


def test_qa_signals_adjacent_tables_when_separator_is_missing(
    monkeypatch: pytest.MonkeyPatch, temp_dir: Path
) -> None:
    """Negative proof: without the separator the QA signal is raised (not silent)."""
    monkeypatch.setattr(WordWriter, "ensure_table_separator", lambda self: False, raising=True)
    out = temp_dir / "no_separator.docx"
    ctx = CompilerContext.create({"word_com": False, "enable_cover": False, "toc": False})
    ctx.compile(ADJACENT_TABLES, {}, out)

    assert _body_sequence(out).count("tbl") == 2
    assert _adjacent_table_pairs(Document(str(out))) == [1]

    result = ctx.final_artifact_qa_result
    assert result is not None
    assert result.metrics.get("adjacent_tables") == 1
    assert any(issue["code"] == "table_adjacent" for issue in result.warnings)


@pytest.mark.integration
def test_it_wordcom_001_table_count_survives_word_roundtrip(
    temp_dir: Path, word_com_ready: None
) -> None:
    """IT-WORDCOM-001: the real production path keeps two logical tables."""
    out = temp_dir / "wordcom_tables.docx"
    ctx = CompilerContext.create({"word_com": True})
    ctx.compile(ADJACENT_TABLES, {}, out)

    assert ctx.final_artifact_qa_result is not None
    assert ctx.final_artifact_qa_result.status != "FAIL"

    # Word 确实参与（TOC 刷新重存）时，两个独立表格必须仍然独立
    with zipfile.ZipFile(out) as zf:
        doc_xml = zf.read("word/document.xml").decode("utf-8")
        app_xml = zf.read("docProps/app.xml").decode("utf-8")
    if "rsidR" not in doc_xml and "Microsoft Office Word" not in app_xml:
        pytest.skip("Word COM round trip did not run in this environment")

    doc = Document(str(out))
    assert len(doc.tables) == 2
    assert _adjacent_table_pairs(doc) == []
