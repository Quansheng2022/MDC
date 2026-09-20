"""
Canonical Acceptance Corpus - 验收测试（P1-09）

对应 CANONICAL_SPEC.md SPEC-AC-001..005。

每个 Release Candidate 必须通过完整 corpus。评价维度:
    - Semantic Fidelity（语义保真）
    - Structural Fidelity（结构保真）
    - Formatting Validity（格式有效性，由质量门保证）
    - No Content Loss（内容零丢失，token 覆盖率 == 1.0）
    - Deterministic Output（确定性抽查）
"""

import json
from pathlib import Path
from typing import Any, Dict, List

import pytest

from md_converter.compiler import CompilerContext
from md_converter.utils.helpers import parse_frontmatter

ACCEPTANCE_DIR = Path(__file__).parent / "acceptance"
MANIFEST_PATH = ACCEPTANCE_DIR / "manifest.json"


def load_manifest() -> List[Dict[str, Any]]:
    """加载 Acceptance manifest。"""
    with open(MANIFEST_PATH, "r", encoding="utf-8") as f:
        return json.load(f)["cases"]


def _compile_case(case: Dict[str, Any], tmp_path: Path) -> Any:
    """编译单个 Acceptance 用例，返回 (ctx, output_path)。"""
    source_path = ACCEPTANCE_DIR / case["file"]
    text = source_path.read_text(encoding="utf-8")
    body, meta = parse_frontmatter(text)
    output_path = tmp_path / f'{case["id"]}.docx'
    config = {"word_com": False, "verbose": False}
    ctx = CompilerContext.create(config)
    ctx.compile(body, meta, output_path)
    return ctx, output_path


def _extract_signature(docx_path: Path) -> Dict[str, Any]:
    """提取 DOCX 结构签名（段落 / 标题 / 表格 / 图片）。"""
    from docx import Document

    doc = Document(str(docx_path))
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
    table_cells = [cell for table in tables for row in table for cell in row]
    return {
        "headings": [
            p.text
            for p in doc.paragraphs
            if p.style is not None and p.style.name.startswith("Heading")
        ],
        "paragraphs": [p.text for p in doc.paragraphs],
        "tables": tables,
        "table_cells": table_cells,
        "inline_shapes": len(doc.inline_shapes),
        "hyperlink_targets": _extract_hyperlink_targets(doc),
    }


def _extract_hyperlink_targets(doc: Any) -> List[str]:
    """提取超链接目标：外部关系目标 + 文档内锚点（AC010 / P11-MNT-007）。

    参数:
        doc: 已打开的 python-docx Document

    返回:
        List[str]: 排序后的目标列表（外部目标为原样引用，锚点以 ``#`` 前缀）
    """
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn

    targets = {
        rel.target_ref
        for rel in doc.part.rels.values()
        if rel.is_external and rel.reltype == RT.HYPERLINK
    }
    for hyperlink in doc.element.body.iter(qn("w:hyperlink")):
        anchor = hyperlink.get(qn("w:anchor"))
        if anchor:
            targets.add(f"#{anchor}")
    return sorted(targets)


def test_acceptance_ac010_hyperlink_targets(tmp_path: Path) -> None:
    """AC010: 链接文本与目标都必须保留在最终 DOCX 中（P11-MNT-007）。"""
    case = next(c for c in load_manifest() if c["id"] == "AC010")
    _, output_path = _compile_case(case, tmp_path)
    signature = _extract_signature(output_path)

    targets = signature["hyperlink_targets"]
    assert "https://openai.com" in targets
    assert "https://example.com/architecture" in targets
    assert "mailto:support@example.com" in targets

    # 可见文本仍然存在（No Content Loss）
    corpus = signature["paragraphs"] + signature["table_cells"]
    for expected_text in ("OpenAI", "support", "Architecture Document"):
        assert any(expected_text in text for text in corpus)


def test_acceptance_ac010_assertion_detects_missing_target(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """负向证明：若链接目标不再写入 DOCX，AC010 断言必须失败（ISSUE-002）。"""
    from docx.shared import RGBColor

    from md_converter.renderer.word_renderer import WordRenderer

    def _legacy_render_link(self: Any, node: Any, font_kind: str = "body") -> None:
        """P11-MNT-007 之前的实现：仅设置样式，不产生超链接关系。"""
        self.inline_state.push(color=RGBColor(0x00, 0x00, 0xFF), underline=True)
        self._render_inline(node.content, font_kind)
        self.inline_state.pop()

    monkeypatch.setattr(WordRenderer, "_render_link", _legacy_render_link)

    case = next(c for c in load_manifest() if c["id"] == "AC010")
    _, output_path = _compile_case(case, tmp_path)
    targets = _extract_signature(output_path)["hyperlink_targets"]

    assert "https://openai.com" not in targets
    assert "mailto:support@example.com" not in targets


@pytest.mark.parametrize("case", load_manifest(), ids=[c["id"] for c in load_manifest()])
def test_acceptance_structural_fidelity(case: Dict[str, Any], tmp_path: Path) -> None:
    """结构签名满足 manifest 最低要求，且全部质量门通过。"""
    ctx, output_path = _compile_case(case, tmp_path)
    signature = _extract_signature(output_path)
    expected = case["signature"]

    assert len(signature["headings"]) >= expected["min_headings"]
    assert len(signature["paragraphs"]) >= expected["min_paragraphs"]
    assert len(signature["tables"]) >= expected["min_tables"]
    assert signature["inline_shapes"] >= expected["min_inline_shapes"]

    # 质量门（CANONICAL_SPEC.md SPEC-QA-001..004）
    assert ctx.static_qa_result is not None
    assert ctx.static_qa_result.status != "FAIL"
    assert ctx.rendered_qa_result is not None
    assert ctx.rendered_qa_result.status != "FAIL"
    assert ctx.final_artifact_qa_result is not None
    assert ctx.final_artifact_qa_result.status != "FAIL"
    assert ctx.final_artifact_sha256

    # No Content Loss（SPEC-INV-001，最高优先级）
    coverage = ctx.final_artifact_qa_result.metrics.get("content_coverage", 1.0)
    assert coverage == 1.0, f"Content loss in {case['id']}: {coverage}"


@pytest.mark.parametrize("case", load_manifest(), ids=[c["id"] for c in load_manifest()])
def test_acceptance_semantic_fidelity(case: Dict[str, Any], tmp_path: Path) -> None:
    """要求的标题与关键文本必须出现在最终 DOCX 中。"""
    _, output_path = _compile_case(case, tmp_path)
    signature = _extract_signature(output_path)

    for expected_heading in case["semantics"].get("headings", []):
        assert any(
            expected_heading in heading for heading in signature["headings"]
        ), f"Missing heading {expected_heading!r} in {case['id']}"

    search_corpus = signature["paragraphs"] + signature["table_cells"]
    for expected_text in case["semantics"].get("contains", []):
        assert any(
            expected_text in text for text in search_corpus
        ), f"Missing text {expected_text!r} in {case['id']}"


@pytest.mark.parametrize(
    "case",
    [c for c in load_manifest() if c["signature"].get("deterministic")],
    ids=[c["id"] for c in load_manifest() if c["signature"].get("deterministic")],
)
def test_acceptance_deterministic_output(case: Dict[str, Any], tmp_path: Path) -> None:
    """确定性抽查：同一输入两次编译的结构签名一致（SPEC-INV-003）。"""
    source_path = ACCEPTANCE_DIR / case["file"]
    body, meta = parse_frontmatter(source_path.read_text(encoding="utf-8"))
    config = {"word_com": False, "verbose": False}

    first_path = tmp_path / "first.docx"
    second_path = tmp_path / "second.docx"
    CompilerContext.create(config).compile(body, meta, first_path)
    CompilerContext.create(config).compile(body, meta, second_path)

    assert _extract_signature(first_path) == _extract_signature(second_path)
