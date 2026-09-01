"""Tests for the V1.5 LayoutPlan and layout decision modules."""

from md_converter.ast.nodes import (
    CodeBlock,
    Diagram,
    Document,
    Heading,
    ListBlock,
    ListItem,
    Paragraph,
    Table,
    TableCell,
    TableRow,
    Text,
)
from md_converter.renderer.layout.content_analyzer import (
    ContentAnalyzer,
    detect_cell_data_type,
    infer_table_kind,
)
from md_converter.renderer.layout.decision_engine import DecisionEngine
from md_converter.renderer.layout.layout_plan import (
    BlockPlan,
    ContentType,
    LayoutPlan,
    SectionPlan,
)
from md_converter.renderer.layout.pagination import PaginationPolicyResolver
from md_converter.renderer.layout.rendered_qa import RenderedQA
from md_converter.renderer.layout.repair_strategy import RepairStrategy
from md_converter.renderer.layout.static_qa import StaticQA
from md_converter.renderer.layout.validity_gate import ValidityGate


def _table(rows: list) -> Table:
    return Table(
        rows=[TableRow(cells=[TableCell(content=[Text(cell)]) for cell in row]) for row in rows]
    )


def test_layout_plan_json_roundtrip() -> None:
    """LayoutPlan survives JSON serialization and reload."""
    plan = LayoutPlan(
        sections=[SectionPlan(id="sec_1", orientation="portrait")],
        blocks=[
            BlockPlan(
                id="blk_001",
                type=ContentType.HEADING,
                original_ast_id="blk_001",
                keep_with_next=True,
                keep_together=True,
            )
        ],
    )

    restored = LayoutPlan.from_json(plan.to_json())

    assert restored.version == "1.5"
    assert restored.sections[0].orientation == "portrait"
    assert restored.blocks[0].type == ContentType.HEADING
    assert restored.blocks[0].keep_with_next is True


def test_content_analyzer_classifies_blocks() -> None:
    """Content types follow the V1.5 classifier (chapter 6)."""
    doc = Document(
        children=[
            Heading(level=1, content=[Text("标题")]),
            Paragraph(content=[Text("用户通过浏览器访问系统。")]),
            Paragraph(content=[Text("Plain English paragraph.")]),
            ListBlock(
                ordered=False, items=[ListItem(children=[Paragraph(content=[Text("item")])])]
            ),
            _table([["A", "1.5", "3"], ["B", "2.5", "4"]]),
            CodeBlock(language="python", text="print('hi')"),
            CodeBlock(language="", text="\n".join(f"line{i}" for i in range(30))),
            Diagram(diagram_type="ascii", content="┌─────┐\n│  A  │\n└─────┘"),
        ]
    )

    analyzer = ContentAnalyzer()
    profiles = analyzer.analyze(doc)
    types = [p.content_type for p in profiles]

    assert types[0] == ContentType.HEADING
    assert types[1] == ContentType.PROSE_CJK
    assert types[2] == ContentType.PROSE_LATIN
    assert types[3] == ContentType.LIST
    assert types[4] == ContentType.TABLE_FINANCIAL
    assert types[5] == ContentType.CODE_SHORT
    assert types[6] == ContentType.CODE_LONG
    assert types[7] == ContentType.ASCII_DIAGRAM


def test_cell_data_type_detection() -> None:
    """Column alignment data types follow chapter 9.2."""
    assert detect_cell_data_type("1,250.5") == "number"
    assert detect_cell_data_type("35.6%") == "percentage"
    assert detect_cell_data_type("$12.8B") == "currency"
    assert detect_cell_data_type("2026-08-08") == "date"
    assert detect_cell_data_type("✅") == "boolean"
    assert detect_cell_data_type("https://example.com") == "url"
    assert detect_cell_data_type("Description") == "text"


def test_infer_table_kind() -> None:
    """Financial tables are detected by numeric columns."""
    assert infer_table_kind(["number", "percentage", "text"]) == ContentType.TABLE_FINANCIAL
    assert infer_table_kind(["text", "text", "text"]) == ContentType.TABLE_GENERAL


def test_pagination_resolver() -> None:
    """Pagination policies are content-type specific (chapter 8)."""
    resolver = PaginationPolicyResolver()

    assert resolver._default_policy(ContentType.HEADING)["keep_with_next"] is True
    assert resolver._default_policy(ContentType.ASCII_DIAGRAM)["no_wrap"] is True
    assert resolver._default_policy(ContentType.CODE_LONG)["max_lines_per_page"] == 45
    assert resolver._default_policy(ContentType.CODE_SHORT)["keep_together"] is True


def test_validity_gate_rejects_unsafe_intents() -> None:
    """User intents below readability floors are rejected with fallback."""
    gate = ValidityGate(
        body_min_pt=10,
        heading_min_pt=10,
        table_min_pt=8.5,
        margin_min_in=0.5,
    )

    result = gate.validate(
        {
            "body_font_size": 9,
            "table_font_size": 7,
            "page_orientation": "landscape",
            "content_deletion": True,
        }
    )

    assert "body_font_size" in result.rejected
    assert result.fallbacks["body_font_size"] == 10.5
    assert "table_font_size" in result.rejected
    assert "content_deletion" in result.rejected
    # 合法意图（landscape）被接受
    assert result.accepted["page_orientation"] == "landscape"
    assert len(result.warnings) == 3


def test_decision_engine_builds_plan() -> None:
    """DecisionEngine produces a serializable LayoutPlan from AST."""
    doc = Document(
        children=[
            Heading(level=1, content=[Text("第一章")]),
            Paragraph(content=[Text("中文段落内容，用于测试两端对齐。")]),
            _table([["A", "B"], ["C", "D"]]),
        ]
    )

    engine = DecisionEngine()
    plan = engine.build_plan(doc)

    assert plan.version == "1.5"
    assert len(plan.blocks) == 3
    assert plan.blocks[0].type == ContentType.HEADING
    assert plan.blocks[0].keep_with_next is True
    assert plan.blocks[0].keep_together is True
    assert plan.blocks[2].repeat_header is True


def test_static_qa_detects_issues() -> None:
    """Static QA reports semantic and structural issues."""
    qa = StaticQA()
    doc = Document(children=[])
    result = qa.run(doc)

    assert result.status == "FAIL"
    assert any(e["code"] == "ast_empty" for e in result.errors)

    bad_doc = Document(children=[Heading(level=7, content=[Text("bad")])])
    result2 = qa.run(bad_doc)
    assert any(e["code"] == "semantic_heading_level" for e in result2.errors)


def test_rendered_qa_metrics(tmp_path) -> None:
    """Rendered QA returns the documented metric shape."""
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("Normal text")
    qa = RenderedQA()
    result = qa.run(doc)

    assert result.status == "PASS"
    assert set(result.metrics) >= {
        "overflow",
        "empty_pages",
        "orphan_headings",
        "table_overflow",
        "figure_overflow",
        "font_violations",
    }


def test_repair_strategy_marks_wide_table_landscape() -> None:
    """Repair marks overflowing tables for landscape sections."""
    from docx import Document as DocxDocument

    plan = LayoutPlan(
        blocks=[
            BlockPlan(
                id="blk_1",
                type=ContentType.TABLE_GENERAL,
                landscape=False,
            )
        ]
    )
    static = StaticQA().run(Document(children=[]))
    rendered = RenderedQA().run(DocxDocument())
    rendered.warnings.append(
        {
            "code": "table_clipping",
            "message": "Table 1: width exceeds content width",
        }
    )

    repaired = RepairStrategy().repair(plan, static, rendered)

    assert repaired.blocks[0].landscape is True
    assert repaired.blocks[0].allow_split is True


def test_table_autofit_detection() -> None:
    """autofit 表格由 Word 自动适配页宽，不应触发宽度溢出警告。"""
    from docx import Document as DocxDocument

    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=6)
    assert RenderedQA._table_is_autofit(table) is True
    table.autofit = False
    assert RenderedQA._table_is_autofit(table) is False
