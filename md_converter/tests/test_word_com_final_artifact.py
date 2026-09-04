"""
Windows Word COM round-trip regression（P0-TBL-04）。

真实生产路径:
    Renderer -> PostProcessor -> python-docx save -> Word COM Open
    -> TOC update -> Word Save -> FinalArtifactQA

仅在 Windows + Word 已安装 + pywin32 可真实启动 Word 时运行；
否则 skip（Windows Release Gate，不是日常跨平台强制测试）。

P11 maintenance hardening:
- 禁止在 pytest collection/import 阶段直接启动 Word COM；
- Word COM 可达性探测在独立 subprocess 中执行，隔离 COM/RPC 状态；
- probe 出现 Windows fatal exception / RPC fatal signature 时 fail-closed；
- 实际 round-trip 测试仍走真实生产路径，不降低 FinalArtifactQA 断言。
"""

from __future__ import annotations

import platform
import subprocess
import sys
from pathlib import Path

import pytest

from md_converter.compiler import CompilerContext
from md_converter.renderer.layout.final_artifact_qa import FinalArtifactQA


WORD_COM_PROBE_TIMEOUT_SECONDS = 30
WORD_COM_FATAL_SIGNATURES = (
    "Windows fatal exception",
    "0x800706be",
    "0x800706ba",
)

# IMPORTANT:
# This code runs in a child Python process, not in the pytest collection process.
# The sentinel prefixes allow the parent fixture to distinguish a legitimately
# unavailable Word environment from an unstable/fatal COM lifecycle failure.
_WORD_COM_PROBE_CODE = r'''
import sys

try:
    import pythoncom
    import win32com.client
except Exception as exc:
    print(f"WORD_COM_UNAVAILABLE|import|{type(exc).__name__}|{exc}")
    raise SystemExit(10)

word = None
co_initialized = False
exit_code = 0

try:
    try:
        pythoncom.CoInitialize()
        co_initialized = True
    except Exception as exc:
        print(f"WORD_COM_UNSTABLE|coinitialize|{type(exc).__name__}|{exc}")
        exit_code = 20

    if exit_code == 0:
        try:
            word = win32com.client.DispatchEx("Word.Application")
        except Exception as exc:
            print(f"WORD_COM_UNAVAILABLE|dispatch|{type(exc).__name__}|{exc}")
            exit_code = 11

    if exit_code == 0:
        try:
            # Keep probe initialization aligned with the production COM path.
            word.Visible = False
            word.DisplayAlerts = 0
        except Exception as exc:
            print(f"WORD_COM_UNSTABLE|initialize|{type(exc).__name__}|{exc}")
            exit_code = 21

finally:
    if word is not None:
        try:
            word.Quit()
        except Exception as exc:
            print(f"WORD_COM_UNSTABLE|quit|{type(exc).__name__}|{exc}")
            if exit_code == 0:
                exit_code = 22
        finally:
            # Explicitly release the application proxy before COM uninitialize.
            word = None

    if co_initialized:
        try:
            pythoncom.CoUninitialize()
        except Exception as exc:
            print(f"WORD_COM_UNSTABLE|couninitialize|{type(exc).__name__}|{exc}")
            if exit_code == 0:
                exit_code = 23

if exit_code == 0:
    print("WORD_COM_OK")

raise SystemExit(exit_code)
'''


def _word_com_probe() -> tuple[str, str]:
    """
    Probe Word COM in an isolated subprocess.

    Returns:
        ("ok", detail)   - Word COM started and shut down cleanly.
        ("skip", detail) - Non-Windows or Word/pywin32 is unavailable.
        ("fail", detail) - Probe timed out, emitted fatal signatures, or showed
                           an unstable COM lifecycle.

    The probe never instantiates Word COM in the pytest collection process.
    """
    if platform.system() != "Windows":
        return "skip", "Requires Windows + installed Word + pywin32 (Windows Release Gate)"

    try:
        cp = subprocess.run(
            [sys.executable, "-c", _WORD_COM_PROBE_CODE],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=WORD_COM_PROBE_TIMEOUT_SECONDS,
            check=False,
        )
    except subprocess.TimeoutExpired:
        return (
            "fail",
            f"Word COM probe timed out after {WORD_COM_PROBE_TIMEOUT_SECONDS}s",
        )
    except OSError as exc:
        return "fail", f"Word COM probe process could not start: {exc}"

    output = (cp.stdout or "").strip()
    output_lower = output.lower()
    fatal_hits = [
        signature
        for signature in WORD_COM_FATAL_SIGNATURES
        if signature.lower() in output_lower
    ]

    if fatal_hits:
        return (
            "fail",
            "Word COM probe emitted fatal signature(s): "
            + ", ".join(fatal_hits)
            + (f"\nProbe output:\n{output}" if output else ""),
        )

    if cp.returncode == 0 and "WORD_COM_OK" in output:
        return "ok", "Word COM probe passed in isolated subprocess"

    # Preserve the original test semantics: environments where Word/pywin32
    # cannot be started are skipped rather than treated as product failures.
    if "WORD_COM_UNAVAILABLE|" in output:
        return (
            "skip",
            "Requires Windows + installed Word + pywin32 (Windows Release Gate)"
            + (f"\nProbe detail: {output}" if output else ""),
        )

    # COM initialization/cleanup instability is not equivalent to a missing
    # dependency. Fail closed so a release gate cannot silently pass an RPC or
    # lifecycle problem.
    if "WORD_COM_UNSTABLE|" in output:
        return (
            "fail",
            f"Word COM probe reported lifecycle instability (rc={cp.returncode})"
            + (f"\nProbe output:\n{output}" if output else ""),
        )

    return (
        "fail",
        f"Word COM probe returned unexpected result (rc={cp.returncode})"
        + (f"\nProbe output:\n{output}" if output else ""),
    )


@pytest.fixture(scope="session")
def require_word_com() -> None:
    """Runtime gate for the Windows Word COM integration test."""
    state, detail = _word_com_probe()

    if state == "skip":
        pytest.skip(detail)
    if state == "fail":
        pytest.fail(detail, pytrace=False)

    assert state == "ok", f"Unexpected Word COM probe state: {state!r}"


@pytest.mark.integration
def test_word_com_roundtrip_final_artifact_qa_pass(
    tmp_path: Path,
    require_word_com: None,
) -> None:
    """Word 保存后，Table Grid style-only 边框仍通过 FinalArtifactQA。"""
    markdown = "# Roadmap\n\n" "| A | B |\n" "| --- | --- |\n" "| 1 | 2 |\n"
    out_path = tmp_path / "word_com_roundtrip.docx"

    # Keep the real production path under test.  The isolated probe above only
    # validates environment reachability; it does not replace this round trip.
    ctx = CompilerContext.create({"word_com": True})
    ctx.compile(markdown, None, out_path)

    assert ctx.final_artifact_qa_result is not None
    assert ctx.final_artifact_qa_result.status == "PASS"
    assert ctx.final_artifact_sha256

    from docx import Document

    doc = Document(str(out_path))
    assert doc.tables
    for table in doc.tables:
        assert FinalArtifactQA._has_effective_table_borders(table)
