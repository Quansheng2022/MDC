"""
Quality Gate 测试（P0-05 / P0-06 / P0-07 / P0-08）。
"""

from pathlib import Path

import pytest

from md_converter.compiler import CompilerContext
from md_converter.quality_gate import (
    QualityGateDecision,
    QualityGateError,
    QualityGatePolicy,
    decide_qa,
)
from md_converter.renderer.layout.artifact_contract import ArtifactContract
from md_converter.renderer.layout.final_artifact_qa import FinalArtifactQA
from md_converter.renderer.layout.rendered_qa import RenderedQA, RenderedQAResult
from md_converter.renderer.layout.static_qa import StaticQAResult
from md_converter.renderer.post_processor import DocxPostProcessor
from md_converter.renderer.themes.v15_theme import V15Theme


def _static_result(status: str = "PASS") -> StaticQAResult:
    return StaticQAResult(status=status)


def _rendered_result(status: str = "PASS", warnings=None, errors=None) -> RenderedQAResult:
    return RenderedQAResult(
        status=status,
        errors=list(errors or []),
        warnings=list(warnings or []),
    )


# ============================================================
# QualityGatePolicy
# ============================================================


def test_policy_from_config_defaults() -> None:
    """默认策略：fail_on_error=True, fail_on_warning=False, iterations=2。"""
    policy = QualityGatePolicy.from_config(None)
    assert policy.fail_on_error is True
    assert policy.fail_on_warning is False
    assert policy.max_repair_iterations == 2


def test_policy_from_config_section() -> None:
    """quality_gate 配置节优先于顶层键。"""
    policy = QualityGatePolicy.from_config(
        {
            "fail_on_error": False,
            "quality_gate": {
                "fail_on_error": True,
                "fail_on_warning": True,
                "max_repair_iterations": 4,
            },
        }
    )
    assert policy.fail_on_error is True
    assert policy.fail_on_warning is True
    assert policy.max_repair_iterations == 4


def test_policy_from_config_top_level_fallback() -> None:
    """无 quality_gate 节时回退到顶层键。"""
    policy = QualityGatePolicy.from_config(
        {"fail_on_error": False, "fail_on_warning": True, "max_repair_iterations": 0}
    )
    assert policy.fail_on_error is False
    assert policy.fail_on_warning is True
    assert policy.max_repair_iterations == 0


# ============================================================
# decide_qa 决策表
# ============================================================


def test_decide_qa_decision_table() -> None:
    """PASS / PASS_WITH_WARN / FAIL 的决策映射。"""
    policy = QualityGatePolicy()
    assert decide_qa("PASS", repairable=False, policy=policy) == QualityGateDecision.CONTINUE
    assert (
        decide_qa("PASS_WITH_WARN", repairable=False, policy=policy) == QualityGateDecision.CONTINUE
    )
    assert decide_qa("PASS_WITH_WARN", repairable=True, policy=policy) == QualityGateDecision.REPAIR
    assert decide_qa("FAIL", repairable=True, policy=policy) == QualityGateDecision.REPAIR
    assert decide_qa("FAIL", repairable=False, policy=policy) == QualityGateDecision.ABORT


def test_decide_qa_fail_on_warning() -> None:
    """fail_on_warning=True 时 PASS_WITH_WARN -> ABORT。"""
    policy = QualityGatePolicy(fail_on_warning=True)
    assert decide_qa("PASS_WITH_WARN", repairable=False, policy=policy) == QualityGateDecision.ABORT


def test_decide_qa_fail_on_error_false() -> None:
    """fail_on_error=False 时 FAIL -> CONTINUE（记录警告，非静默）。"""
    policy = QualityGatePolicy(fail_on_error=False)
    assert decide_qa("FAIL", repairable=False, policy=policy) == QualityGateDecision.CONTINUE


# ============================================================
# Compiler 集成：StaticQA Gate（P0-05）
# ============================================================


def test_compiler_rejects_empty_document() -> None:
    """空 AST -> build rejected（QualityGateError, stage=static_qa）。"""
    ctx = CompilerContext.create({"word_com": False})
    with pytest.raises(QualityGateError) as exc_info:
        ctx.compile("")
    assert exc_info.value.stage == "static_qa"


def test_compiler_rejects_font_below_minimum(tmp_path: Path) -> None:
    """body_font < minimum -> build rejected。"""
    theme = V15Theme.load_default()
    theme.data["typography"]["body"]["size"] = "8pt"
    ctx = CompilerContext.create({"theme": theme, "word_com": False})
    with pytest.raises(QualityGateError) as exc_info:
        ctx.compile("# Title\n\nbody text", None, tmp_path / "out.docx")
    assert exc_info.value.stage == "static_qa"


def test_compiler_allows_warning_only_document(tmp_path: Path) -> None:
    """warning-only issue（空标题）-> build allowed（默认 fail_on_warning=False）。"""
    ctx = CompilerContext.create({"word_com": False})
    doc = ctx.compile("#\n\nSome body text.", None, tmp_path / "warn.docx")
    assert doc is not None
    assert ctx.static_qa_result.status == "PASS_WITH_WARN"


def test_compiler_fail_on_warning_aborts(tmp_path: Path) -> None:
    """fail_on_warning=True 时 warning-only issue -> QualityGateError。"""
    ctx = CompilerContext.create({"word_com": False, "fail_on_warning": True})
    with pytest.raises(QualityGateError) as exc_info:
        ctx.compile("#\n\nSome body text.", None, tmp_path / "warn2.docx")
    assert exc_info.value.stage == "static_qa"


# ============================================================
# Compiler 集成：RenderedQA Gate + Repair 闭环（P0-06 / P0-07）
# ============================================================


def test_compiler_rendered_qa_fail_aborts(tmp_path: Path, monkeypatch) -> None:
    """RenderedQA FAIL 且不可修复 -> QualityGateError（stage=rendered_qa）。"""
    failing = _rendered_result(
        status="FAIL",
        errors=[{"code": "font_violation", "message": "run below minimum"}],
    )

    def fake_run(self, doc, layout_plan):
        return failing

    monkeypatch.setattr(
        "md_converter.renderer.layout.rendered_qa.RenderedQA.run",
        fake_run,
    )
    ctx = CompilerContext.create({"word_com": False})
    with pytest.raises(QualityGateError) as exc_info:
        ctx.compile("# Title\n\nbody text", None, tmp_path / "fail.docx")
    assert exc_info.value.stage == "rendered_qa"


def test_compiler_repair_loop_reinspects(tmp_path: Path, monkeypatch) -> None:
    """Repair -> Re-render -> Re-inspect：修复后必须再经过一次 RenderedQA。"""
    calls = {"count": 0}
    original_run = RenderedQA.run

    def fake_run(self, doc, layout_plan):
        if self.diag is None:
            # FinalArtifactQA 内部复用的 RenderedQA 走真实逻辑
            return original_run(self, doc, layout_plan)
        calls["count"] += 1
        if calls["count"] == 1:
            return _rendered_result(
                status="PASS_WITH_WARN",
                warnings=[
                    {"code": "table_clipping", "message": "Table 1: width exceeds content width"}
                ],
            )
        return _rendered_result(status="PASS")

    monkeypatch.setattr(
        "md_converter.renderer.layout.rendered_qa.RenderedQA.run",
        fake_run,
    )
    markdown = "# Table\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n"
    ctx = CompilerContext.create({"word_com": False})
    ctx.compile(markdown, None, tmp_path / "repair.docx")

    assert len(ctx.repair_evidence) == 1
    record = ctx.repair_evidence[0]
    assert record.selected_strategy == "table_landscape"
    assert record.before_status == "PASS_WITH_WARN"
    assert record.after_status == "PASS"
    assert record.changed_blocks
    assert calls["count"] == 2


def test_compiler_repair_iterations_bounded(tmp_path: Path, monkeypatch) -> None:
    """修复闭环有界：超过 max_repair_iterations 后 FAIL -> QualityGateError。"""
    failing = _rendered_result(
        status="FAIL",
        errors=[{"code": "font_violation", "message": "run below minimum"}],
        warnings=[{"code": "table_clipping", "message": "Table 1: width exceeds content width"}],
    )

    def fake_run(self, doc, layout_plan):
        return failing

    def fake_can_repair(self, plan, static_result, rendered_result):
        return True

    def fake_repair(self, plan, static_result, rendered_result):
        from dataclasses import replace

        return replace(plan, metadata={**plan.metadata, "repaired": True})

    monkeypatch.setattr(
        "md_converter.renderer.layout.rendered_qa.RenderedQA.run",
        fake_run,
    )
    monkeypatch.setattr(
        "md_converter.renderer.layout.repair_strategy.RepairStrategy.can_repair",
        fake_can_repair,
    )
    monkeypatch.setattr(
        "md_converter.renderer.layout.repair_strategy.RepairStrategy.repair",
        fake_repair,
    )
    markdown = "# Table\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n"
    ctx = CompilerContext.create({"word_com": False, "max_repair_iterations": 2})
    with pytest.raises(QualityGateError) as exc_info:
        ctx.compile(markdown, None, tmp_path / "bounded.docx")
    assert exc_info.value.stage == "repair"


# ============================================================
# Compiler 集成：FinalArtifactQA Gate（P0-08）
# ============================================================


def test_compiler_final_artifact_qa_fail_aborts(tmp_path: Path, monkeypatch) -> None:
    """Post-Processor 后 Final QA FAIL -> QualityGateError。"""

    def fake_run(self, docx_path, **kwargs):
        from md_converter.renderer.layout.final_artifact_qa import FinalArtifactQAResult

        return FinalArtifactQAResult(
            status="FAIL",
            errors=[{"code": "body_empty", "message": "Final artifact body is empty"}],
        )

    monkeypatch.setattr(
        "md_converter.renderer.layout.final_artifact_qa.FinalArtifactQA.run",
        fake_run,
    )
    ctx = CompilerContext.create({"word_com": False})
    with pytest.raises(QualityGateError) as exc_info:
        ctx.compile("# Title\n\nbody text", None, tmp_path / "final.docx")
    assert exc_info.value.stage == "final_artifact_qa"


# ============================================================
# FinalArtifactQA 单元测试
# ============================================================


def test_final_artifact_qa_missing_file_fails(tmp_path: Path) -> None:
    """文件无法打开 -> FAIL（artifact_unreadable）。"""
    qa = FinalArtifactQA()
    result = qa.run(
        tmp_path / "missing.docx",
        config={"enable_cover": False, "toc": False, "style_tables": False},
    )
    assert result.status == "FAIL"
    assert any(e["code"] == "artifact_unreadable" for e in result.errors)


def test_final_artifact_qa_content_loss_detected(tmp_path: Path) -> None:
    """源文本 token 缺失 -> FAIL（content_loss），并记录 artifact hash。"""
    from docx import Document

    path = tmp_path / "lossy.docx"
    doc = Document()
    doc.add_paragraph("Only partial content.")
    doc.save(str(path))

    qa = FinalArtifactQA()
    result = qa.run(
        path,
        source_text="Only partial content. MISSING_TOKEN_XYZ",
        config={"enable_cover": False, "toc": False, "style_tables": False},
    )
    assert result.status == "FAIL"
    assert any(e["code"] == "content_loss" for e in result.errors)
    assert result.metrics["artifact_sha256"]
    assert result.metrics["content_coverage"] < 1.0


def test_final_artifact_qa_passes_intact_document(tmp_path: Path) -> None:
    """完整文档（含封面/TOC 关闭）-> PASS。"""
    from docx import Document

    path = tmp_path / "ok.docx"
    doc = Document()
    doc.add_paragraph("完整内容 intact content.")
    doc.save(str(path))

    qa = FinalArtifactQA()
    result = qa.run(
        path,
        source_text="完整内容 intact content.",
        config={"enable_cover": False, "toc": False, "style_tables": False},
    )
    assert result.status == "PASS"
    assert result.metrics["content_coverage"] == 1.0


# ============================================================
# Post-Processor Contract Gate（修改项 1）
# ============================================================


def test_post_processor_failure_fail_closed(tmp_path: Path, monkeypatch) -> None:
    """POST-01: PostProcessor 基础设施异常 -> QualityGateError(stage=post_processor)。"""

    def boom(*args, **kwargs):
        raise RuntimeError("injected post-processor failure")

    monkeypatch.setattr(DocxPostProcessor, "process", staticmethod(boom))
    ctx = CompilerContext.create({"word_com": False})
    with pytest.raises(QualityGateError) as exc_info:
        ctx.compile("# Title\n\nbody text", None, tmp_path / "pp.docx")
    assert exc_info.value.stage == "post_processor"
    assert any(d.code == "POST001" for d in ctx.diag.diagnostics)


def test_post_processor_not_run_when_nothing_required(tmp_path: Path, monkeypatch) -> None:
    """enable_cover/toc/style_tables 全 False 时不运行 PostProcessor。"""
    called = {"count": 0}

    def fake_process(*args, **kwargs):
        called["count"] += 1

    monkeypatch.setattr(DocxPostProcessor, "process", staticmethod(fake_process))
    ctx = CompilerContext.create(
        {"word_com": False, "enable_cover": False, "toc": False, "style_tables": False}
    )
    ctx.compile("# Title\n\nbody text", None, tmp_path / "nopost.docx")
    assert called["count"] == 0


def test_final_qa_required_toc_missing_is_error(tmp_path: Path) -> None:
    """POST-02: toc_required=True 且 TOC 缺失 -> FAIL（toc_missing）。"""
    from docx import Document

    path = tmp_path / "no_toc.docx"
    doc = Document()
    doc.add_paragraph("Content without TOC field.")
    doc.save(str(path))

    qa = FinalArtifactQA()
    result = qa.run(
        path,
        config={"enable_cover": False, "toc": True, "style_tables": False},
    )
    assert result.status == "FAIL"
    assert any(e["code"] == "toc_missing" for e in result.errors)


def test_final_qa_required_cover_missing_is_error(tmp_path: Path) -> None:
    """POST-03: cover_required=True 且封面缺失 -> FAIL（cover_missing）。"""
    from docx import Document

    path = tmp_path / "no_cover.docx"
    doc = Document()
    doc.add_paragraph("Body content without cover.")
    doc.save(str(path))

    qa = FinalArtifactQA()
    result = qa.run(
        path,
        metadata={"title": "Expected Cover Title"},
        config={"enable_cover": True, "toc": False, "style_tables": False},
    )
    assert result.status == "FAIL"
    assert any(e["code"] == "cover_missing" for e in result.errors)


def test_final_qa_toc_disabled_absent_is_pass(tmp_path: Path) -> None:
    """POST-04: toc=False 且 TOC 缺失 -> PASS（不产生 warning/error）。"""
    from docx import Document

    path = tmp_path / "ok_no_toc.docx"
    doc = Document()
    doc.add_paragraph("Content without TOC.")
    doc.save(str(path))

    qa = FinalArtifactQA()
    result = qa.run(
        path,
        config={"enable_cover": False, "toc": False, "style_tables": False},
    )
    assert result.status == "PASS"


def test_final_qa_cover_disabled_absent_is_pass(tmp_path: Path) -> None:
    """POST-05: enable_cover=False 且封面缺失 -> PASS。"""
    from docx import Document

    path = tmp_path / "ok_no_cover.docx"
    doc = Document()
    doc.add_paragraph("Body content without cover.")
    doc.save(str(path))

    qa = FinalArtifactQA()
    result = qa.run(
        path,
        metadata={"title": "Some Title"},
        config={"enable_cover": False, "toc": False, "style_tables": False},
    )
    assert result.status == "PASS"


def test_final_qa_table_styling_required_missing_is_error(tmp_path: Path) -> None:
    """POST-06: style_tables=True 且表格存在但样式缺失 -> FAIL。"""
    from docx import Document

    path = tmp_path / "unstyled_table.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    doc.save(str(path))

    qa = FinalArtifactQA()
    result = qa.run(
        path,
        config={"enable_cover": False, "toc": False, "style_tables": True},
    )
    assert result.status == "FAIL"
    assert any(e["code"] == "table_styling_missing" for e in result.errors)


def test_final_qa_table_styling_required_no_table_is_pass(tmp_path: Path) -> None:
    """POST-07: style_tables=True 但文档无表格 -> PASS。"""
    from docx import Document

    path = tmp_path / "no_table.docx"
    doc = Document()
    doc.add_paragraph("No tables here.")
    doc.save(str(path))

    qa = FinalArtifactQA()
    result = qa.run(
        path,
        config={"enable_cover": False, "toc": False, "style_tables": True},
    )
    assert result.status == "PASS"


def test_final_qa_table_grid_style_only_borders_pass(tmp_path: Path) -> None:
    """P0-TBL-03: Table Grid style-only borders（无 direct w:tblBorders）-> PASS。

    Word COM 保存时可能规范化移除冗余 direct w:tblBorders，边框仍由
    Table Grid style 继承；FinalArtifactQA 必须按 effective borders 判定。
    """
    from docx import Document

    path = tmp_path / "style_only_borders.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.save(str(path))

    qa = FinalArtifactQA()
    result = qa.run(
        path,
        config={"enable_cover": False, "toc": False, "style_tables": True},
    )
    assert result.status == "PASS"


def test_final_qa_truly_borderless_table_fails(tmp_path: Path) -> None:
    """P0-TBL-03 负向: 既无 Table Grid 样式也无 direct borders -> FAIL。"""
    from docx import Document

    path = tmp_path / "borderless.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.save(str(path))

    qa = FinalArtifactQA()
    result = qa.run(
        path,
        config={"enable_cover": False, "toc": False, "style_tables": True},
    )
    assert result.status == "FAIL"
    assert any("table_borders_missing" in e["message"] for e in result.errors)


def test_artifact_contract_from_config() -> None:
    """ArtifactContract 只消费 resolved configuration（F4）。"""
    contract = ArtifactContract.from_config(
        {"enable_cover": True, "toc": False, "style_tables": True}
    )
    assert contract.cover_required is True
    assert contract.toc_required is False
    assert contract.table_styling_required is True


def test_artifact_contract_requires_resolved_config() -> None:
    """缺失键（非 resolved configuration）-> ValueError，不允许第二套默认值。"""
    with pytest.raises(ValueError, match="missing keys"):
        ArtifactContract.from_config({"enable_cover": True, "toc": False})
    with pytest.raises(ValueError, match="missing keys"):
        ArtifactContract.from_config(None)
