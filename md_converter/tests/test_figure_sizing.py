"""
图形尺寸策略与渲染后几何验证（P12-CAND-002 / S/N 118 实现包 002）。

覆盖:
    - 主题 figure 策略解析（长度单位、content_width / available_page_height 语义）
    - 尺寸适配纯函数（保持宽高比、永不放大、不超内容区、最小宽度下限）
    - RenderedQA figure_overflow / figure_below_min_width 真实测量
    - 端到端渲染：超高图形被收缩到有效内容区内；最小宽度下限产生 RENDER005
"""

import base64
from pathlib import Path
from typing import Any, List, Tuple

import pytest
from docx import Document as DocxDocument
from docx.shared import Cm

from md_converter.compiler import CompilerContext
from md_converter.renderer.layout.figure_sizing import (
    FigureMeasurementError,
    figure_policy_from_data,
    figure_policy_from_theme,
    fit_figure_size,
    parse_length_cm,
)
from md_converter.renderer.layout.rendered_qa import RenderedQA
from md_converter.renderer.themes.v15_theme import V15Theme

#: 30 x 70 px PNG（宽高比 0.4286：收缩后宽度仍在 8cm 下限之上）
TALL_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAB4AAABGCAIAAABUj7NvAAAANklEQVR42u3MQQEAQAQAMC6J"
    "ULLJeym8bAGW1RM7XqxRq9VqtVqtVqvVarVarVar1Wq1Wn27/sDYAZZqC8CYAAAAAElFTkSuQmCC"
)

#: 10 x 100 px PNG（宽高比 0.1：收缩后宽度低于 8cm 下限）
NARROW_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAoAAABkCAIAAAAwrjbAAAAAKklEQVR42u3JQREAIAgAMCQJ"
    "ochmXhJgAG/77lTf2GU8aa211lprrX/pAVqZAdKDzZvyAAAAAElFTkSuQmCC"
)


def _write_png(tmp_path: Path, b64: str, name: str) -> Path:
    """将 base64 PNG 写入临时文件，返回路径。"""
    path = tmp_path / name
    path.write_bytes(base64.b64decode(b64))
    return path


def _data_uri(b64: str) -> str:
    """构造 PNG data URI。"""
    return f"data:image/png;base64,{b64}"


def _content_box_cm(doc: Any) -> Tuple[float, float]:
    """从文档 section 读取有效内容区（厘米）。"""
    section = doc.sections[-1]
    width = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
    height = section.page_height.cm - section.top_margin.cm - section.bottom_margin.cm
    return width, height


def _compile(body: str, tmp_path: Path) -> Tuple[Any, Path]:
    """编译 Markdown 并返回 (ctx, docx 路径)。"""
    output_path = tmp_path / "figure.docx"
    ctx = CompilerContext.create({"word_com": False, "verbose": False})
    ctx.compile(body, {}, output_path)
    return ctx, output_path


def _codes(ctx: Any) -> List[str]:
    """返回全部诊断 code。"""
    return [d.code for d in ctx.diag.diagnostics]


def _a4_document_with_margins() -> Any:
    """构造 A4 + 1in 页边距的空文档。"""
    doc = DocxDocument()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    return doc


def test_parse_length_cm_units() -> None:
    """主题长度单位解析。"""
    assert parse_length_cm("8cm") == pytest.approx(8.0)
    assert parse_length_cm("1in") == pytest.approx(2.54)
    assert parse_length_cm("10mm") == pytest.approx(1.0)
    assert parse_length_cm(12.7) == pytest.approx(12.7)
    assert parse_length_cm(None, 3.0) == pytest.approx(3.0)
    assert parse_length_cm("not-a-length", 5.0) == pytest.approx(5.0)


def test_figure_policy_reads_frozen_theme() -> None:
    """冻结主题的 figure 块解析为内容区边界 + 8cm 下限。"""
    policy = figure_policy_from_theme(V15Theme.load_default())

    assert policy.max_width_cm is None
    assert policy.max_height_cm is None
    assert policy.min_width_cm == pytest.approx(8.0)
    assert policy.preserve_aspect_ratio is True
    assert policy.alignment == "center"
    assert policy.keep_together is True


def test_figure_policy_defaults_without_declaration() -> None:
    """未声明 figure 块时使用内容区边界且不设下限。"""
    policy = figure_policy_from_data(None)

    assert policy.max_width_cm is None
    assert policy.max_height_cm is None
    assert policy.min_width_cm is None


def test_fit_figure_size_normal_fit() -> None:
    """常规图形保持目标宽度。"""
    fit = fit_figure_size(80, 40, 12.7, 15.92, 24.62, 8.0)

    assert fit.width_cm == pytest.approx(12.7)
    assert fit.height_cm == pytest.approx(6.35, abs=0.01)
    assert fit.scaled is False
    assert fit.below_min_width is False


def test_fit_figure_size_scales_tall_figure() -> None:
    """超高图形按内容区高度收缩并保持宽高比。"""
    fit = fit_figure_size(30, 70, 12.7, 15.92, 24.62, 8.0)

    assert fit.scaled is True
    assert fit.height_cm == pytest.approx(24.62, abs=0.01)
    assert fit.width_cm == pytest.approx(24.62 * 30 / 70, abs=0.01)
    assert fit.below_min_width is False
    assert fit.width_cm <= 12.7
    assert fit.height_cm <= 24.62


def test_fit_figure_size_marks_min_width_floor() -> None:
    """收缩后低于主题最小宽度下限时标记（仍按适配尺寸交付）。"""
    fit = fit_figure_size(10, 100, 12.7, 15.92, 24.62, 8.0)

    assert fit.scaled is True
    assert fit.below_min_width is True
    assert fit.width_cm == pytest.approx(2.462, abs=0.01)
    assert fit.height_cm == pytest.approx(24.62, abs=0.01)


def test_fit_figure_size_never_exceeds_bounds() -> None:
    """目标宽度超过内容区宽度时收敛到内容区宽度。"""
    fit = fit_figure_size(100, 100, 30.0, 15.92, 24.62, 8.0)

    assert fit.width_cm == pytest.approx(15.92)
    assert fit.height_cm == pytest.approx(15.92)
    assert fit.width_cm <= 15.92
    assert fit.height_cm <= 24.62


def test_fit_figure_size_rejects_invalid_measurements() -> None:
    """非法固有尺寸/边界必须显式失败。"""
    with pytest.raises(FigureMeasurementError):
        fit_figure_size(0, 10, 12.7, 15.92, 24.62, 8.0)
    with pytest.raises(FigureMeasurementError):
        fit_figure_size(10, 10, 12.7, 0, 24.62, 8.0)


def test_rendered_qa_detects_figure_overflow(tmp_path: Path) -> None:
    """交付图形超出内容区时 RenderedQA 报 error 并 FAIL。"""
    png = _write_png(tmp_path, TALL_PNG_B64, "tall.png")

    doc = _a4_document_with_margins()
    doc.add_paragraph().add_run().add_picture(str(png), width=Cm(12.7), height=Cm(30.0))

    result = RenderedQA(V15Theme.load_default()).run(doc)

    assert result.metrics["figure_overflow"] == 1
    assert result.status == "FAIL"
    assert any(issue["code"] == "figure_overflow" for issue in result.errors)


def test_rendered_qa_accepts_fitted_figure(tmp_path: Path) -> None:
    """符合内容区的图形不产生 overflow，也不误报最小宽度。"""
    png = _write_png(tmp_path, TALL_PNG_B64, "fitted.png")

    doc = _a4_document_with_margins()
    doc.add_paragraph().add_run().add_picture(str(png), width=Cm(10.55), height=Cm(24.62))

    result = RenderedQA(V15Theme.load_default()).run(doc)

    assert result.metrics["figure_overflow"] == 0
    assert result.metrics["figure_below_min_width"] == 0
    assert result.status == "PASS"


def test_rendered_qa_reports_min_width_floor(tmp_path: Path) -> None:
    """低于主题最小宽度下限时计入 warning 而不失败。"""
    png = _write_png(tmp_path, NARROW_PNG_B64, "narrow.png")

    doc = _a4_document_with_margins()
    doc.add_paragraph().add_run().add_picture(str(png), width=Cm(2.46), height=Cm(24.62))

    result = RenderedQA(V15Theme.load_default()).run(doc)

    assert result.metrics["figure_overflow"] == 0
    assert result.metrics["figure_below_min_width"] == 1
    assert result.status == "PASS_WITH_WARN"
    assert any(issue["code"] == "figure_below_min_width" for issue in result.warnings)


def test_compile_scales_tall_figure_into_content_box(tmp_path: Path) -> None:
    """端到端：超高图形被收缩到有效内容区内且保持宽高比。"""
    body = f"# Figure\n\n![tall]({_data_uri(TALL_PNG_B64)})\n"
    ctx, output_path = _compile(body, tmp_path)

    doc = DocxDocument(str(output_path))
    content_width, content_height = _content_box_cm(doc)
    shapes = list(doc.inline_shapes)

    assert shapes, "图形成端到端渲染中丢失"
    for shape in shapes:
        assert shape.width.cm <= content_width + 0.05
        assert shape.height.cm <= content_height + 0.05
    assert shapes[-1].height.cm == pytest.approx(content_height, abs=0.05)
    assert shapes[-1].height.cm / shapes[-1].width.cm == pytest.approx(70 / 30, rel=0.01)

    assert ctx.rendered_qa_result is not None
    assert ctx.rendered_qa_result.metrics["figure_overflow"] == 0
    assert "RENDER005" not in _codes(ctx)


def test_compile_warns_below_min_width(tmp_path: Path) -> None:
    """端到端：收缩后低于 8cm 下限时产生 RENDER005 且仍满足内容区约束。"""
    body = f"# Figure\n\n![narrow]({_data_uri(NARROW_PNG_B64)})\n"
    ctx, output_path = _compile(body, tmp_path)

    doc = DocxDocument(str(output_path))
    content_width, content_height = _content_box_cm(doc)
    shape = list(doc.inline_shapes)[-1]

    assert shape.width.cm <= content_width + 0.05
    assert shape.height.cm <= content_height + 0.05
    assert shape.width.cm < 8.0
    assert "RENDER005" in _codes(ctx)
    assert ctx.rendered_qa_result is not None
    assert ctx.rendered_qa_result.metrics["figure_below_min_width"] >= 1
    assert ctx.rendered_qa_result.status != "FAIL"
