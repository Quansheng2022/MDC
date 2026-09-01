"""
Pytest Configuration - 测试配置

提供 pytest 的 fixtures 和测试辅助功能。
包括 Golden Test 支持、临时目录、诊断收集器等。
"""

import shutil
import tempfile
from pathlib import Path
from typing import Any, Dict

import pytest

from md_converter.compiler import CompilerContext
from md_converter.diagnostics.collector import DiagnosticCollector

# ============================================================
# Fixtures
# ============================================================


@pytest.fixture
def temp_dir() -> Path:
    """
    创建临时目录。

    返回:
        Path: 临时目录路径
    """
    path = Path(tempfile.mkdtemp(prefix="md_converter_test_"))
    yield path
    # 清理
    shutil.rmtree(path, ignore_errors=True)


@pytest.fixture
def temp_file(temp_dir: Path) -> Path:
    """
    创建临时文件路径。

    参数:
        temp_dir: 临时目录

    返回:
        Path: 临时文件路径
    """
    return temp_dir / "test.md"


@pytest.fixture
def diagnostic_collector() -> DiagnosticCollector:
    """
    创建诊断收集器。

    返回:
        DiagnosticCollector: 诊断收集器实例
    """
    return DiagnosticCollector()


@pytest.fixture
def compiler_context(diagnostic_collector) -> CompilerContext:
    """
    创建编译器上下文。

    参数:
        diagnostic_collector: 诊断收集器

    返回:
        CompilerContext: 编译器上下文实例
    """
    config = {
        "verbose": True,
        "normalize": True,
        "diagram": True,
    }
    return CompilerContext.create(config)


@pytest.fixture
def sample_markdown() -> str:
    """
    示例 Markdown 内容。

    返回:
        str: Markdown 文本
    """
    return """---
title: Test Document
author: Test Author
date: 2024-01-01
tags: [test, markdown]
---

# Heading 1

This is a **bold** and *italic* text with `inline code`.

## Heading 2

- List item 1
- List item 2
  - Nested item 2.1
  - Nested item 2.2

### Heading 3

1. Ordered item 1
2. Ordered item 2

| Header 1 | Header 2 | Header 3 |
|----------|----------|----------|
| Cell 1   | Cell 2   | Cell 3   |
| Cell 4   | Cell 5   | Cell 6   |

```python
def hello():
    print("Hello, World!")
```
"""


# ============================================================
# DOCX 结构提取（Golden Test 支持）
# ============================================================


def extract_docx_structure(path) -> Dict[str, Any]:
    """
    提取 DOCX 文档的文本结构。

    参数:
        path: DOCX 文件路径

    返回:
        Dict[str, Any]: 文档结构（段落、表格、图片数量）
    """
    from docx import Document

    doc = Document(path)
    return {
        "paragraphs": [p.text for p in doc.paragraphs],
        "tables": [
            [[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables
        ],
        "inline_shapes": len(doc.inline_shapes),
    }
