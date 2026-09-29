"""Focused verification for TOC heading localization (THL).

Covers the mandated behavior matrix, the pure localization authority, and the
real DOCX artifacts (TOC title, field, entries, page break) for English,
Chinese, mixed and unknown/minimal documents - including a mixed-language batch
through the application service.
"""

from __future__ import annotations

import ast
from pathlib import Path
from typing import List, Set
from zipfile import ZipFile

import pytest
from docx import Document

from md_converter.application.conversion_request import ConversionRequest
from md_converter.application.conversion_service import ConversionService
from md_converter.compiler import compile_markdown
from md_converter.renderer.layout import language_detection
from md_converter.renderer.layout.toc_localization import (
    CHINESE_TOC_HEADING,
    DEFAULT_TOC_HEADING,
    toc_heading_for_document_text,
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
LAYOUT_DIR = PROJECT_ROOT / "md_converter" / "renderer" / "layout"
TOC_LOCALIZATION = LAYOUT_DIR / "toc_localization.py"
LANGUAGE_DETECTION = LAYOUT_DIR / "language_detection.py"

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

#: Representative documents for the frozen behavior matrix.
ENGLISH_DOC = "# Guide\n\nThis document explains how the converter works.\n"
CHINESE_DOC = "# 使用手册\n\n本文档说明转换器的工作方式。\n"
MIXED_DOC = "# Guide 使用手册\n\nThe 转换器 works with mixed content.\n"
UNKNOWN_DOC = "# 2026\n\n1.5 2.5 3.5 — ., 。、；\n"
ENGLISH_HEADING_CHINESE_BODY = "# Installation\n\n请按照以下步骤操作。\n"
CHINESE_HEADING_ENGLISH_BODY = "# 安装\n\nFollow the steps below.\n"


def _write_markdown(tmp_path: Path, name: str, body: str) -> Path:
    """Write ``body`` as UTF-8 Markdown and return the path."""
    path = tmp_path / name
    path.write_text(body, encoding="utf-8")
    return path


def _compile(tmp_path: Path, name: str, body: str, config: dict | None = None) -> Path:
    """Compile ``body`` to DOCX and return the artifact path."""
    target = tmp_path / f"{name}.docx"
    options = {"word_com": False, "verbose": False}
    options.update(config or {})
    compile_markdown(body, config=options, output_path=target)
    return target


def _toc_heading_paragraph(path: Path):
    """Return the TOC heading paragraph of a compiled artifact, or ``None``."""
    doc = Document(str(path))
    for paragraph in doc.paragraphs:
        style = paragraph.style
        if style is not None and style.name == "TOC Heading" and paragraph.text.strip():
            return paragraph
    return None


def _toc_entries(path: Path) -> List[str]:
    """Return the TOC entry texts (hyperlinked entry paragraphs) of an artifact."""
    doc = Document(str(path))
    return [
        paragraph.text.strip()
        for paragraph in doc.paragraphs
        if paragraph._element.find(f"{_W}hyperlink") is not None
    ]


def _zip_text(path: Path, member: str) -> str:
    """Read one member of the DOCX archive as UTF-8 text."""
    with ZipFile(path) as archive:
        return archive.read(member).decode("utf-8", "replace")


def _imported_modules(path: Path) -> Set[str]:
    """Return the import targets of ``path`` (relative imports included)."""
    tree = ast.parse(path.read_text(encoding="utf-8"))
    names: Set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            names.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom):
            module = node.module or ""
            if node.level:
                module = f"{'.' * node.level}{module}"
            if module:
                names.add(module)
    return names


def _code_string_constants(path: Path) -> List[str]:
    """Return string literals used in code (docstrings excluded)."""
    tree = ast.parse(path.read_text(encoding="utf-8"))
    docstrings = {
        ast.get_docstring(node, clean=False)
        for node in ast.walk(tree)
        if isinstance(node, (ast.Module, ast.ClassDef, ast.FunctionDef))
    }
    docstrings.discard(None)
    return [
        node.value
        for node in ast.walk(tree)
        if isinstance(node, ast.Constant)
        and isinstance(node.value, str)
        and node.value not in docstrings
    ]


# ============================================================
# Frozen behavior matrix (localization authority)
# ============================================================


def test_english_only_document_uses_english_heading() -> None:
    """English prose only -> "Table of Contents"."""
    assert toc_heading_for_document_text(ENGLISH_DOC) == DEFAULT_TOC_HEADING
    assert DEFAULT_TOC_HEADING == "Table of Contents"


def test_chinese_document_uses_chinese_heading() -> None:
    """Chinese prose -> "目录"."""
    assert toc_heading_for_document_text(CHINESE_DOC) == CHINESE_TOC_HEADING
    assert CHINESE_TOC_HEADING == "目录"


def test_mixed_document_uses_chinese_heading() -> None:
    """Mixed English + Chinese -> "目录"."""
    assert toc_heading_for_document_text(MIXED_DOC) == CHINESE_TOC_HEADING


@pytest.mark.parametrize("text", [None, "", "   ", "\n\n", "\t"])
def test_empty_or_language_unknown_uses_english_heading(text: object) -> None:
    """Empty / unknown language -> "Table of Contents"."""
    assert toc_heading_for_document_text(text) == DEFAULT_TOC_HEADING


def test_numbers_and_punctuation_use_english_heading() -> None:
    """Numbers, symbols and (full-width) punctuation alone are not Chinese."""
    assert toc_heading_for_document_text("1.5 2.5 3.5 — ., 。、；：！？") == DEFAULT_TOC_HEADING


def test_english_heading_with_chinese_body_uses_chinese_heading() -> None:
    """Detection considers the whole visible document, not just headings."""
    assert toc_heading_for_document_text(ENGLISH_HEADING_CHINESE_BODY) == CHINESE_TOC_HEADING


def test_chinese_heading_with_english_body_uses_chinese_heading() -> None:
    """A Chinese heading alone is enough for the Chinese heading."""
    assert toc_heading_for_document_text(CHINESE_HEADING_ENGLISH_BODY) == CHINESE_TOC_HEADING


def test_localization_is_deterministic() -> None:
    """The same input always yields the same heading (SPEC-INV-003)."""
    first = toc_heading_for_document_text(MIXED_DOC)

    assert first == toc_heading_for_document_text(MIXED_DOC)
    assert first == CHINESE_TOC_HEADING


# ============================================================
# Script predicate
# ============================================================


def test_han_predicate_detects_only_han_ideographs() -> None:
    """Kana, Hangul and CJK punctuation are not "meaningful Chinese"."""
    assert language_detection.contains_han_ideograph("中文") is True
    assert language_detection.contains_han_ideograph("Hello 世") is True
    assert language_detection.contains_han_ideograph("こんにちは") is False  # kana only
    assert language_detection.contains_han_ideograph("한국어") is False  # hangul only
    assert language_detection.contains_han_ideograph("。、；：！？") is False
    assert language_detection.contains_han_ideograph("Table of Contents") is False
    assert language_detection.contains_han_ideograph("2026 1.5 — .,") is False
    assert language_detection.contains_han_ideograph("") is False


# ============================================================
# Real DOCX evidence
# ============================================================


@pytest.mark.parametrize(
    ("name", "body", "expected"),
    [
        ("english", ENGLISH_DOC, DEFAULT_TOC_HEADING),
        ("chinese", CHINESE_DOC, CHINESE_TOC_HEADING),
        ("mixed", MIXED_DOC, CHINESE_TOC_HEADING),
        ("unknown", UNKNOWN_DOC, DEFAULT_TOC_HEADING),
        ("english_heading", ENGLISH_HEADING_CHINESE_BODY, CHINESE_TOC_HEADING),
        ("chinese_heading", CHINESE_HEADING_ENGLISH_BODY, CHINESE_TOC_HEADING),
    ],
)
def test_artifact_toc_heading_is_localized(
    tmp_path: Path, name: str, body: str, expected: str
) -> None:
    """The delivered DOCX carries the localized TOC heading."""
    artifact = _compile(tmp_path, name, body)
    heading = _toc_heading_paragraph(artifact)

    assert heading is not None
    assert heading.text.strip() == expected
    assert heading.style.name == "TOC Heading"
    assert heading.alignment is not None and int(heading.alignment) == 1  # CENTER
    assert heading.runs and heading.runs[0].font.bold is True
    assert heading.runs and heading.runs[0].font.size.pt == pytest.approx(18)


def test_toc_field_and_entries_are_unchanged(tmp_path: Path) -> None:
    """Localization changes the title text only: field, entries and break stay."""
    artifact = _compile(tmp_path, "entries", ENGLISH_DOC)
    document_xml = _zip_text(artifact, "word/document.xml")
    settings_xml = _zip_text(artifact, "word/settings.xml")

    # TOC field semantics unchanged
    assert 'TOC \\o "1-3" \\h \\z \\u' in document_xml
    assert 'w:fldCharType="begin"' in document_xml
    assert "w:dirty" not in document_xml
    assert "updateFields" not in settings_xml

    # Real cached entries, hyperlinked to heading bookmarks
    assert "w:bookmarkStart" in document_xml
    assert "_Toc_" in document_xml
    assert "w:hyperlink" in document_xml
    assert "w:anchor=" in document_xml

    entries = _toc_entries(artifact)
    assert entries == ["Guide"]
    assert Document(str(artifact)).paragraphs  # artifact opens successfully


def test_chinese_document_toc_entries_and_title(tmp_path: Path) -> None:
    """A Chinese document keeps its entries and gains the Chinese title."""
    artifact = _compile(tmp_path, "chinese_entries", CHINESE_DOC)

    assert _toc_heading_paragraph(artifact).text.strip() == CHINESE_TOC_HEADING
    assert _toc_entries(artifact) == ["使用手册"]


def test_page_break_after_toc_is_preserved(tmp_path: Path) -> None:
    """The page break after the last TOC entry is unaffected."""
    artifact = _compile(tmp_path, "page_break", ENGLISH_DOC)
    doc = Document(str(artifact))
    entries = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph._element.find(f"{_W}hyperlink") is not None
    ]

    assert entries
    following = entries[-1]._element.getnext()
    break_types = [br.get(f"{_W}type") for br in following.iter(f"{_W}br")]
    assert "page" in break_types


def test_toc_disabled_still_has_no_toc_heading(tmp_path: Path) -> None:
    """toc=False behavior is unchanged for both languages."""
    english = _compile(tmp_path, "no_toc_en", ENGLISH_DOC, {"toc": False})
    chinese = _compile(tmp_path, "no_toc_zh", CHINESE_DOC, {"toc": False})

    for artifact in (english, chinese):
        assert _toc_heading_paragraph(artifact) is None
        assert "instrText" not in _zip_text(artifact, "word/document.xml")


# ============================================================
# Orthogonality: profiles and batch
# ============================================================


@pytest.mark.parametrize("profile_id", ["professional_report", "business_report", "academic"])
def test_output_profiles_do_not_affect_localization(tmp_path: Path, profile_id: str) -> None:
    """Presentation profiles and TOC localization are independent."""
    english = _compile(tmp_path, f"en_{profile_id}", ENGLISH_DOC, {"output_profile": profile_id})
    chinese = _compile(tmp_path, f"zh_{profile_id}", CHINESE_DOC, {"output_profile": profile_id})

    assert _toc_heading_paragraph(english).text.strip() == DEFAULT_TOC_HEADING
    assert _toc_heading_paragraph(chinese).text.strip() == CHINESE_TOC_HEADING


def test_mixed_language_batch_localizes_each_document(tmp_path: Path) -> None:
    """One batch, several languages: each artifact gets its own correct title.

    Proves the decision is derived per document (no cross-document state), which
    is what keeps Serial Batch behavior unaffected.
    """
    service = ConversionService({"word_com": False})
    sources = [
        _write_markdown(tmp_path, "en.md", ENGLISH_DOC),
        _write_markdown(tmp_path, "zh.md", CHINESE_DOC),
        _write_markdown(tmp_path, "mixed.md", MIXED_DOC),
        _write_markdown(tmp_path, "unknown.md", UNKNOWN_DOC),
    ]

    titles = []
    for source in sources:
        result = service.convert(
            ConversionRequest(source_path=source, output_path=source.with_suffix(".docx"))
        )
        assert result.status.value in ("SUCCESS", "SUCCESS_WITH_WARNING")
        assert result.output_path is not None
        titles.append(_toc_heading_paragraph(result.output_path).text.strip())

    assert titles == [
        DEFAULT_TOC_HEADING,
        CHINESE_TOC_HEADING,
        CHINESE_TOC_HEADING,
        DEFAULT_TOC_HEADING,
    ]


# ============================================================
# Single authority / dependency boundary
# ============================================================


def test_localization_module_is_a_dependency_free_single_authority() -> None:
    """The authority has no IO, no Qt, no third-party and no duplicate rule."""
    imported = _imported_modules(TOC_LOCALIZATION)

    assert imported <= {"__future__", "typing", ".language_detection"}, imported
    assert not [name for name in imported if name.startswith(("PySide6", "os", "socket"))]


def test_toc_heading_texts_have_exactly_one_authority() -> None:
    """No other production module may hardcode a TOC heading text in code."""
    offenders: List[str] = []

    for module in (PROJECT_ROOT / "md_converter").rglob("*.py"):
        if module == TOC_LOCALIZATION or "tests" in module.parts:
            continue
        for literal in _code_string_constants(module):
            if literal in (CHINESE_TOC_HEADING, DEFAULT_TOC_HEADING):
                offenders.append(f"{module.relative_to(PROJECT_ROOT)}: {literal!r}")

    assert offenders == []
