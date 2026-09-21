"""
空标题行为（P12-CAND-003 / S/N 118 实现包 003）。

覆盖:
    - 空标题不渲染：无段落、无占位文本、无标题计数
    - AST 保留 source-truth Heading 节点
    - StaticQA 仍报告 semantic_empty_heading 警告
    - TOC 不出现合成条目；正常标题不受影响
"""

from pathlib import Path
from typing import Any, List

from docx import Document as DocxDocument

from md_converter.ast.nodes import Heading
from md_converter.compiler import CompilerContext
from md_converter.diagnostics.collector import DiagnosticCollector
from md_converter.parser.markdown_parser import parse_markdown
from md_converter.renderer.layout.static_qa import StaticQA


def _compile(body: str, tmp_path: Path, name: str = "heading.docx") -> Any:
    """编译 Markdown，返回 (ctx, docx 路径)。"""
    output_path = tmp_path / name
    ctx = CompilerContext.create({"word_com": False, "verbose": False})
    ctx.compile(body, {}, output_path)
    return ctx, output_path


def _paragraph_texts(doc: Any) -> List[str]:
    """全部段落文本。"""
    return [p.text for p in doc.paragraphs]


def _rendered_headings(doc: Any) -> List[str]:
    """渲染后的标题文本（Word Heading N 样式）。"""
    return [
        p.text for p in doc.paragraphs if p.style is not None and p.style.name.startswith("Heading")
    ]


def _toc_entries(doc: Any) -> List[str]:
    """目录条目文本（TOC 1..3 样式）。"""
    return [
        p.text
        for p in doc.paragraphs
        if p.style is not None and p.style.name in {"TOC 1", "TOC 2", "TOC 3"}
    ]


def test_ast_keeps_empty_heading_as_source_truth() -> None:
    """AST 保留空标题节点（source truth），便于诊断与追溯。"""
    doc = parse_markdown("## \n\nBody text.\n", DiagnosticCollector())

    headings = [node for node in doc.children if isinstance(node, Heading)]
    assert len(headings) == 1
    assert headings[0].to_plain_text().strip() == ""


def test_static_qa_still_reports_empty_heading() -> None:
    """既有 StaticQA 警告保持（WARN + DROP 的 WARN 部分）。"""
    diag = DiagnosticCollector()
    doc = parse_markdown("## \n\nBody text.\n", diag)

    result = StaticQA(diag=diag).run(doc)

    assert result.status == "PASS_WITH_WARN"
    assert any(issue["code"] == "semantic_empty_heading" for issue in result.warnings)


def test_renderer_drops_empty_heading(tmp_path: Path) -> None:
    """空标题不产生段落，也不注入占位文本 "Heading"。"""
    ctx, output_path = _compile("# Real Title\n\n## \n\nBody text.\n", tmp_path)
    doc = DocxDocument(str(output_path))
    texts = _paragraph_texts(doc)

    assert "Heading" not in texts
    assert not any("Heading" == text.strip() for text in texts)
    assert _rendered_headings(doc) == ["Real Title"]
    assert ctx.static_qa_result is not None
    assert any(issue["code"] == "semantic_empty_heading" for issue in ctx.static_qa_result.warnings)


def test_toc_has_no_synthetic_entry(tmp_path: Path) -> None:
    """目录条目数量等于真实标题数量，不出现合成条目。"""
    body = "# Real Title\n\n## \n\n### Section A\n\nBody text.\n"
    _ctx, output_path = _compile(body, tmp_path, "toc.docx")
    doc = DocxDocument(str(output_path))

    rendered = _rendered_headings(doc)
    entries = _toc_entries(doc)

    assert rendered == ["Real Title", "Section A"]
    assert len(entries) == len(rendered)
    assert not any(entry.strip() == "Heading" for entry in entries)
    assert any("Section A" in entry for entry in entries)


def test_whitespace_only_heading_variants_are_dropped(tmp_path: Path) -> None:
    """空白变体（空格/制表符）与空标题同样处理。"""
    body = "# Title\n\n##\t\n\n###   \n\nBody.\n"
    _ctx, output_path = _compile(body, tmp_path, "variants.docx")
    doc = DocxDocument(str(output_path))

    assert _rendered_headings(doc) == ["Title"]
    assert not any("Heading" in text for text in _paragraph_texts(doc))


def test_adjacent_headings_keep_levels(tmp_path: Path) -> None:
    """空标题与相邻标题：只渲染真实标题并保持层级。"""
    body = "## \n\n### Kept Heading\n\nBody.\n"
    _ctx, output_path = _compile(body, tmp_path, "adjacent.docx")
    doc = DocxDocument(str(output_path))

    headings = [
        (p.text, p.style.name)
        for p in doc.paragraphs
        if p.style is not None and p.style.name.startswith("Heading")
    ]
    assert headings == [("Kept Heading", "Heading 3")]
    assert len(_toc_entries(doc)) == 1


def test_normal_headings_unaffected(tmp_path: Path) -> None:
    """正常标题保持文本、样式与目录条目。"""
    body = "# One\n\n## Two\n\n### Three\n\nBody.\n"
    _ctx, output_path = _compile(body, tmp_path, "normal.docx")
    doc = DocxDocument(str(output_path))

    assert _rendered_headings(doc) == ["One", "Two", "Three"]
    assert len(_toc_entries(doc)) == 3
    assert "Body." in _paragraph_texts(doc)
