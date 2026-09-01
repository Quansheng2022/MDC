"""Tests for DOCX post-processor TOC insertion (native field, no win32com)."""

from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from md_converter.compiler import compile_markdown
from md_converter.renderer.post_processor import DocxPostProcessor

MARKDOWN = """# 标题

正文内容。
"""

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _read_zip_text(path: Path, member: str) -> str:
    """Read a member of a docx archive as UTF-8 text."""
    with ZipFile(path) as zf:
        return zf.read(member).decode("utf-8", "replace")


def test_toc_field_inserted_natively(temp_dir: Path) -> None:
    """TOC title, field, entry hyperlinks and page break are inserted without dirty flags."""
    out = temp_dir / "toc.docx"
    compile_markdown(
        MARKDOWN,
        config={"toc": True, "word_com": False},
        output_path=out,
    )

    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert "目录" in texts

    title_para = next(p for p in doc.paragraphs if p.text.strip() == "目录")
    assert title_para.style.name == "TOC Heading"
    assert title_para.alignment is not None and title_para.alignment == 1  # CENTER

    docxml = _read_zip_text(out, "word/document.xml")
    assert 'TOC \\o "1-3" \\h \\z \\u' in docxml

    # No global updateFields: it triggers Word's "update fields?" dialog on open.
    settings = _read_zip_text(out, "word/settings.xml")
    assert "updateFields" not in settings

    # No field-level w:dirty either: it also triggers Word's "update fields?" dialog.
    assert 'w:fldCharType="begin"' in docxml
    assert "w:dirty" not in docxml

    # 缓存结果包含真实条目：标题书签 + 超链接
    assert "w:bookmarkStart" in docxml
    assert "_Toc_" in docxml
    assert "w:hyperlink" in docxml
    assert "w:anchor=" in docxml
    assert "标题" in docxml  # 目录条目中的标题文本

    # A page break paragraph follows the last TOC entry
    entries = [p for p in doc.paragraphs if p._element.find(f"{_W}hyperlink") is not None]
    assert entries, "expected TOC entry hyperlink paragraphs"
    assert entries[0].style.name == "TOC 1"
    next_para = entries[-1]._element.getnext()
    br_types = [br.get(f"{_W}type") for br in next_para.iter(f"{_W}br")]
    assert "page" in br_types


def test_toc_skipped_when_disabled(temp_dir: Path) -> None:
    """toc=False leaves the document without a TOC field."""
    out = temp_dir / "no_toc.docx"
    compile_markdown(
        MARKDOWN,
        config={"toc": False, "word_com": False},
        output_path=out,
    )

    doc = Document(str(out))
    assert all(p.text.strip() != "目录" for p in doc.paragraphs)

    docxml = _read_zip_text(out, "word/document.xml")
    assert "instrText" not in docxml

    settings = _read_zip_text(out, "word/settings.xml")
    assert "updateFields" not in settings


def test_add_table_borders_propagates_failure() -> None:
    """P0-TBL-02: _add_table_borders 不再吞异常（SPEC-INV-006）。"""

    class _BrokenTable:
        """表对象：_tbl 不可访问，触发 AttributeError。"""

        def __init__(self) -> None:
            self._tbl = None

    with pytest.raises(AttributeError):
        DocxPostProcessor._add_table_borders(_BrokenTable())


def test_set_table_header_style_propagates_failure() -> None:
    """P0-TBL-02: _set_table_header_style 不再吞异常（SPEC-INV-006）。"""

    class _BrokenTable:
        """表对象：rows[0] 无 cells，触发 AttributeError。"""

        def __init__(self) -> None:
            self.rows = [None]

    with pytest.raises(AttributeError):
        DocxPostProcessor._set_table_header_style(_BrokenTable())
