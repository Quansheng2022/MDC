"""P11-MNT-008 / ISSUE-007: frozen A4 page size must be applied to portrait output.

The frozen theme (`QS-Word-Default-V1.5`) and default config both declare
``page.size = A4``; these tests pin the resulting OOXML geometry and the
existing landscape behaviour (wide tables) to the same authority.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from md_converter.compiler import compile_markdown
from md_converter.utils.helpers import parse_frontmatter

ACCEPTANCE_DIR = Path(__file__).parent / "acceptance"

# A4 (21.0 cm × 29.7 cm) in twips, as written by word_writer.set_section_orientation
A4_PORTRAIT = (11906, 16838)
A4_LANDSCAPE = (16838, 11906)


def _wide_table_columns(columns: int) -> str:
    """15 列时估算宽度 18.0 cm > portrait 正文宽度 + buffer → 触发 landscape。"""
    header = "|" + "|".join(f" H{i} " for i in range(columns)) + "|"
    separator = "|" + "|".join(" --- " for _ in range(columns)) + "|"
    row = "|" + "|".join(" " + "a" * 8 + " " for _ in range(columns)) + "|"
    return "\n".join(["# Wide", "", header, separator, row, ""])


WIDE_TABLE = _wide_table_columns(15)


def _section_geometry(doc: Document) -> list[tuple[int, int, bool]]:
    """Return (width_twips, height_twips, is_landscape) for every section."""
    out: list[tuple[int, int, bool]] = []
    for section in doc.sections:
        landscape = section.page_width > section.page_height
        out.append((section.page_width.twips, section.page_height.twips, landscape))
    return out


def test_ut_page_001_default_config_produces_a4_portrait(temp_dir: Path) -> None:
    """UT-PAGE-001: frozen default config → portrait w:pgSz == A4."""
    out = temp_dir / "a4_portrait.docx"
    compile_markdown(
        "# Heading\n\nBody text.\n",
        config={"word_com": False, "enable_cover": False, "toc": False},
        output_path=out,
    )

    doc = Document(str(out))
    assert _section_geometry(doc) == [(*A4_PORTRAIT, False)]

    # 页边距保持不变（1 in = 1440 twips，四种边距一致）
    section = doc.sections[0]
    assert section.left_margin.twips == 1440
    assert section.right_margin.twips == 1440
    assert section.top_margin.twips == 1440
    assert section.bottom_margin.twips == 1440


def test_representative_document_geometry_is_a4(temp_dir: Path) -> None:
    """Representative DOCX geometry check: tracked acceptance document is A4."""
    source = ACCEPTANCE_DIR / "AC011_long_document.md"
    body, meta = parse_frontmatter(source.read_text(encoding="utf-8"))
    out = temp_dir / "ac011_a4.docx"
    compile_markdown(body, config={"word_com": False}, metadata=meta, output_path=out)

    geometry = _section_geometry(Document(str(out)))
    assert geometry, "expected at least one section"
    assert geometry[0] == (*A4_PORTRAIT, False)
    assert all(not landscape for _, _, landscape in geometry)


def test_landscape_wide_table_section_uses_a4_landscape(temp_dir: Path) -> None:
    """Landscape regression check: wide tables still switch to A4 landscape."""
    out = temp_dir / "a4_landscape.docx"
    compile_markdown(
        WIDE_TABLE,
        config={"word_com": False, "enable_cover": False, "toc": False},
        output_path=out,
    )

    geometry = _section_geometry(Document(str(out)))
    landscapes = [g for g in geometry if g[2]]
    assert landscapes, "layout engine must select landscape for a 15-column table"
    assert landscapes[0] == (*A4_LANDSCAPE, True)
    for width, height, is_landscape in geometry:
        if not is_landscape:
            assert (width, height) == A4_PORTRAIT
