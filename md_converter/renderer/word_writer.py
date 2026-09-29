"""
Word Writer - Word 文档原子操作封装
"""

import base64
import os
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Optional, Sequence, Union

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.image.image import Image as DocxImage
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import Section
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from .layout.figure_sizing import FigureFit, fit_figure_size

#: 1 厘米 = 1/2.54 英寸 = 1440/2.54 twips（OOXML ``w:tblW`` / ``w:tcW`` 单位）
TWIPS_PER_CM = 1440.0 / 2.54


@dataclass(frozen=True)
class FigureBounds:
    """
    图形适配边界（有效内容区，厘米）。

    属性:
        target_width_cm: 目标宽度（不超过有效内容区宽度）。
        max_width_cm: 有效内容区宽度。
        max_height_cm: 有效内容区高度。
        min_width_cm: 主题声明的最小宽度下限（可选）。
    """

    target_width_cm: float
    max_width_cm: float
    max_height_cm: float
    min_width_cm: Optional[float] = None


@dataclass(frozen=True)
class ImagePlacement:
    """
    图形插入结果（P12-CAND-002）。

    属性:
        width_cm: 实际插入宽度（厘米）。
        height_cm: 实际插入高度（厘米）；未测量时为 ``None``。
        scaled: 是否发生了缩小。
        below_min_width: 是否低于主题声明的最小宽度下限。
        measured: 固有尺寸是否测量成功。
    """

    width_cm: float
    height_cm: Optional[float]
    scaled: bool
    below_min_width: bool
    measured: bool


def set_style_font(
    style,
    font_name: Optional[str] = None,
    *,
    ascii: Optional[str] = None,
    hAnsi: Optional[str] = None,
    eastAsia: Optional[str] = None,
    cs: Optional[str] = None,
    mapping: Optional[Dict[str, str]] = None,
):
    """
    设置样式字体，同时覆盖 ascii, hAnsi, eastAsia, cs。

    兼容旧调用：仅传 font_name 时四个槽位使用同一字体；
    传入 mapping 时按 V1.5 字体映射规则分别设置（中英混排）。
    """
    if not style:
        return

    if mapping:
        ascii = mapping.get("ascii", ascii)
        hAnsi = mapping.get("hAnsi", hAnsi)
        eastAsia = mapping.get("eastAsia", eastAsia)
        cs = mapping.get("cs", cs)

    if not any([font_name, ascii, hAnsi, eastAsia, cs]):
        return

    resolved_ascii = ascii or font_name
    resolved_h_ansi = hAnsi or font_name
    resolved_east_asia = eastAsia or font_name
    resolved_cs = cs or font_name

    # 设置字体名称
    if resolved_ascii:
        style.font.name = resolved_ascii

    # 确保 rPr 存在
    rPr = style.element.rPr
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        style.element.append(rPr)

    # 创建或获取 rFonts 元素
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)

    # 设置所有字符集的字体
    if resolved_ascii:
        rFonts.set(qn("w:ascii"), resolved_ascii)  # 英文/ASCII
    if resolved_h_ansi:
        rFonts.set(qn("w:hAnsi"), resolved_h_ansi)  # 高 ANSI
    if resolved_east_asia:
        rFonts.set(qn("w:eastAsia"), resolved_east_asia)  # 东亚字符
    if resolved_cs:
        rFonts.set(qn("w:cs"), resolved_cs)  # 复杂脚本


def set_run_font(
    run,
    font_name: Optional[str] = None,
    *,
    east_asia: Optional[str] = None,
    mapping: Optional[Dict[str, str]] = None,
):
    """
    设置 Run 字体，支持 ascii / hAnsi / eastAsia / cs 四槽位。

    参数:
        run: python-docx Run
        font_name: 西文字体（ascii/hAnsi/cs）
        east_asia: 东亚字体（eastAsia）
        mapping: 完整字体映射
    """
    if run is None:
        return
    if mapping:
        font_name = mapping.get("ascii", font_name)
        east_asia = mapping.get("eastAsia", east_asia)
    if font_name:
        run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    if font_name:
        rFonts.set(qn("w:ascii"), font_name)
        rFonts.set(qn("w:hAnsi"), font_name)
    if east_asia:
        rFonts.set(qn("w:eastAsia"), east_asia)
    if font_name and not rFonts.get(qn("w:cs")):
        rFonts.set(qn("w:cs"), font_name)


class WordWriter:
    """
    Word 文档原子操作封装。

    职责:
        1. 封装 python-docx 的所有底层操作
        2. 提供原子化的文档操作方法
        3. 不包含任何样式决策逻辑
        4. 所有样式参数由调用者传入
    """

    def __init__(self, document: Optional[Document] = None):
        self.doc = document or Document()
        self.current_paragraph: Optional[Paragraph] = None
        self.current_table: Optional[Table] = None
        self.current_row = None
        self.current_cell = None
        self.current_section: Optional[Section] = None

        # 表格行/列索引（用于访问预创建的行和单元格）
        self._row_idx: int = 0
        self._col_idx: int = 0

    @classmethod
    def create(cls, template_path: Optional[Union[str, Path]] = None) -> "WordWriter":
        if template_path and Path(template_path).exists():
            doc = Document(str(template_path))
        else:
            doc = Document()
        return cls(doc)

    def start_document(self) -> None:
        self.doc = self.doc or Document()
        self.current_paragraph = None
        self.current_table = None
        self.current_row = None
        self.current_cell = None
        self._row_idx = 0
        self._col_idx = 0

    def finish(self) -> Document:
        return self.doc

    def save(self, path: Union[str, Path]) -> None:
        self.doc.save(str(path))

    # ============================================================
    # 段落操作
    # ============================================================

    def add_paragraph(self, text: str = "", style: str = None) -> Paragraph:
        if style:
            self.current_paragraph = self.doc.add_paragraph(text, style=style)
        else:
            self.current_paragraph = self.doc.add_paragraph(text)
        return self.current_paragraph

    def wrap_runs_as_hyperlink(self, paragraph: Paragraph, start_index: int, href: str) -> bool:
        """把段落中新追加的 run 包裹为真实 Word 超链接（P11-MNT-007）。

        ``#anchor`` 形式使用文档内锚点（``w:anchor``，不创建外部关系）；
        其余目标创建 ``TargetMode="External"`` 关系，http/https/mailto 与
        相对路径均按 Markdown 源原样保留。既有 run 样式（颜色/下划线）不变。

        参数:
            paragraph: 目标段落
            start_index: 链接内容渲染前的直接 ``w:r`` 数量（包裹边界）
            href: Markdown 链接目标

        返回:
            bool: 是否创建了超链接（目标为空或无新增 run 时返回 False）
        """
        target = (href or "").strip()
        if not target:
            return False

        p_el = paragraph._p
        runs = list(p_el.findall(qn("w:r")))
        new_runs = runs[start_index:]
        if not new_runs:
            return False

        hyperlink = OxmlElement("w:hyperlink")
        if target.startswith("#"):
            anchor = target[1:].strip()
            if not anchor:
                return False
            hyperlink.set(qn("w:anchor"), anchor)
        else:
            hyperlink.set(
                qn("r:id"),
                paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True),
            )
        hyperlink.set(qn("w:history"), "1")

        new_runs[0].addprevious(hyperlink)
        for run in new_runs:
            p_el.remove(run)
            hyperlink.append(run)
        return True

    def add_run(
        self,
        text: str,
        bold: bool = False,
        italic: bool = False,
        underline: bool = False,
        font_name: Optional[str] = None,
        east_asia_font: Optional[str] = None,
        font_mapping: Optional[Dict[str, str]] = None,
        font_size: Optional[int] = None,
        color: Optional[RGBColor] = None,
        highlight: Optional[RGBColor] = None,
    ) -> None:
        """
        添加 Run，支持中英混排字体映射。

        参数:
            text: 文本内容
            bold: 加粗
            italic: 斜体
            underline: 下划线
            font_name: 西文字体
            east_asia_font: 东亚字体
            font_mapping: 完整字体映射（四槽位）
            font_size: 字号（磅）
            color: 颜色
            highlight: 高亮
        """
        if self.current_paragraph is None:
            self.add_paragraph()
        run = self.current_paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        run.underline = underline
        if font_name:
            run.font.name = font_name
        if font_mapping:
            set_run_font(run, mapping=font_mapping)
        elif east_asia_font:
            set_run_font(run, font_name=font_name, east_asia=east_asia_font)
        if font_size:
            run.font.size = Pt(font_size)
        if color:
            run.font.color.rgb = color
        if highlight:
            run.font.highlight_color = highlight

    def set_keep_with_next(self, paragraph: Optional[Paragraph] = None) -> None:
        """设置段落后与下一段同页（Keep with next）。"""
        target = paragraph or self.current_paragraph
        if target is not None:
            target.paragraph_format.keep_with_next = True

    def set_keep_lines(self, paragraph: Optional[Paragraph] = None) -> None:
        """设置段落行保持同页（Keep together）。"""
        target = paragraph or self.current_paragraph
        if target is not None:
            target.paragraph_format.keep_together = True

    def set_widow_control(self, paragraph: Optional[Paragraph] = None) -> None:
        """启用 Widow/Orphan 控制。"""
        target = paragraph or self.current_paragraph
        if target is not None:
            target.paragraph_format.widow_control = True

    def set_repeat_table_header(self, row) -> None:
        """设置表格行作为跨页重复表头。"""
        if row is None:
            return
        trPr = row._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        trPr.append(tbl_header)

    def set_row_cant_split(self, row) -> None:
        """禁止行内拆分（保持行完整性）。"""
        if row is None:
            return
        trPr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        trPr.append(cant_split)

    def set_cell_vertical_alignment(self, cell, value: str = "top") -> None:
        """设置单元格垂直对齐（top/center/bottom）。"""
        if cell is None:
            return
        tcPr = cell._tc.get_or_add_tcPr()
        v_align = tcPr.find(qn("w:vAlign"))
        if v_align is None:
            v_align = OxmlElement("w:vAlign")
            tcPr.append(v_align)
        v_align.set(qn("w:val"), value)

    def set_section_orientation(self, orientation: str = "portrait") -> None:
        """
        设置当前 Section 方向（A4 时交换宽高）。

        参数:
            orientation: 'portrait' 或 'landscape'
        """
        section = self.current_section or self.doc.sections[0]
        portrait = (21.0, 29.7)
        landscape = (29.7, 21.0)
        width_cm, height_cm = landscape if orientation == "landscape" else portrait
        section.page_width = Cm(width_cm)
        section.page_height = Cm(height_cm)
        section.orientation = (
            WD_ORIENT.LANDSCAPE if orientation == "landscape" else WD_ORIENT.PORTRAIT
        )

    def add_soft_break(self) -> None:
        if self.current_paragraph:
            self.current_paragraph.add_run().add_break()

    def add_hard_break(self) -> None:
        """Markdown 硬换行 → Word 行内换行（不是分页符）。"""
        if self.current_paragraph:
            self.current_paragraph.add_run().add_break(WD_BREAK.LINE)

    def add_horizontal_rule(self) -> None:
        p = self.add_paragraph()
        pPr = p._element.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "auto")
        pBdr.append(bottom)
        pPr.append(pBdr)

    def set_paragraph_alignment(self, alignment: int) -> None:
        if self.current_paragraph:
            self.current_paragraph.alignment = alignment

    def set_paragraph_spacing(
        self,
        before: Optional[Pt] = None,
        after: Optional[Pt] = None,
        line_spacing: Optional[float] = None,
    ) -> None:
        if self.current_paragraph:
            if before is not None:
                self.current_paragraph.paragraph_format.space_before = before
            if after is not None:
                self.current_paragraph.paragraph_format.space_after = after
            if line_spacing is not None:
                self.current_paragraph.paragraph_format.line_spacing = line_spacing

    def set_paragraph_indent(
        self, left: Optional[Cm] = None, right: Optional[Cm] = None, first_line: Optional[Cm] = None
    ) -> None:
        if self.current_paragraph:
            if left is not None:
                self.current_paragraph.paragraph_format.left_indent = left
            if right is not None:
                self.current_paragraph.paragraph_format.right_indent = right
            if first_line is not None:
                self.current_paragraph.paragraph_format.first_line_indent = first_line

    def set_paragraph_background(self, paragraph: Paragraph, color: RGBColor) -> None:
        pPr = paragraph._element.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), f"{color.rgb:06X}" if hasattr(color, "rgb") else str(color))
        pPr.append(shd)

    # ============================================================
    # 表格操作
    # ============================================================

    def start_table(self, rows: int = 0, cols: int = 0) -> Table:
        self.ensure_table_separator()
        if rows > 0 and cols > 0:
            self.current_table = self.doc.add_table(rows=rows, cols=cols)
        else:
            self.current_table = self.doc.add_table(rows=1, cols=1)
            if len(self.current_table.rows) > 0:
                self.current_table.rows[0]._element.getparent().remove(
                    self.current_table.rows[0]._element
                )
        self.current_table.style = "Table Grid"
        self._row_idx = 0
        self._col_idx = 0
        return self.current_table

    def ensure_table_separator(self) -> bool:
        """在紧跟表格之后插入 Word 稳定分隔段落（P11-MNT-009）。

        Word 在打开/保存文档时会合并**直接相邻**的 ``w:tbl``，从而丢失 Markdown
        源中由空行分隔的两个独立表格（表格数量、表头重复、gridSpan 结构被改写）。
        这里在新建表格前检查 body 末尾内容：若最后一个内容元素是 ``w:tbl``，
        则补一个空段落作为稳定分隔（``w:sectPr`` 不计入内容元素）。

        返回:
            bool: 是否插入了分隔段落
        """
        body = self.doc.element.body
        content = [child for child in body if child.tag != qn("w:sectPr")]
        if content and content[-1].tag == qn("w:tbl"):
            self.add_paragraph()
            return True
        return False

    def end_table(self) -> None:
        self.current_table = None
        self.current_row = None
        self.current_cell = None
        self._row_idx = 0
        self._col_idx = 0

    def start_row(self):
        if self.current_table is None:
            raise RuntimeError("No active table")
        if self._row_idx >= len(self.current_table.rows):
            raise IndexError(
                f"Row index {self._row_idx} out of range "
                f"(table has {len(self.current_table.rows)} rows)"
            )
        self.current_row = self.current_table.rows[self._row_idx]
        self._col_idx = 0
        return self.current_row

    def end_row(self) -> None:
        self._row_idx += 1
        self.current_row = None

    def start_cell(self):
        if self.current_row is None:
            raise RuntimeError("No active row")
        if self._col_idx >= len(self.current_row.cells):
            raise IndexError(
                f"Cell index {self._col_idx} out of range "
                f"(row has {len(self.current_row.cells)} cells)"
            )
        self.current_cell = self.current_row.cells[self._col_idx]
        p = self.current_cell.paragraphs[0]
        # 清除默认 Run
        for run in list(p.runs):
            p._element.remove(run._element)
        self.current_paragraph = p
        return self.current_cell

    def end_cell(self) -> None:
        self._col_idx += 1
        self.current_cell = None

    def set_cell_background(self, cell, color: RGBColor) -> None:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), f"{color.rgb:06X}" if hasattr(color, "rgb") else str(color))
        tcPr.append(shd)

    def merge_cells(self, start_row: int, start_col: int, end_row: int, end_col: int):
        if self.current_table is None:
            raise RuntimeError("No active table")
        return self.current_table.cell(start_row, start_col).merge(
            self.current_table.cell(end_row, end_col)
        )

    def set_table_alignment(self, alignment: int) -> None:
        if self.current_table:
            self.current_table.alignment = alignment

    def apply_table_column_widths(self, column_widths_cm: Sequence[float]) -> None:
        """
        应用确定的表格列宽（Program D / WP-D03）。

        由 :func:`md_converter.renderer.layout.table_fitting.plan_table_fit` 的决策驱动：
        本方法只做原子写入，不做任何适配决策（``SPEC-ARCH-009``：Renderer/Writer 只执行）。

        写入内容:
            - ``w:tblLayout type="fixed"``（固定列宽，Word 不再重新分配）；
            - ``w:tblW`` = 决策总宽（dxa），使表格宽度确定而不依赖 Word autofit；
            - ``w:tblGrid/w:gridCol`` 每列宽度；
            - 每个 ``w:tc`` 的 ``w:tcW``（与 gridCol 一致，保证 Word 一致解释）。

        文本、行/列结构、合并、样式与顺序均不被修改。

        参数:
            column_widths_cm: 每列宽度（厘米，按可见列顺序）；空序列时不做任何写入。
        """
        table = self.current_table
        if table is None or not column_widths_cm:
            return

        widths = [float(width) for width in column_widths_cm]
        table.autofit = False

        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            layout = tbl_pr.find(qn("w:tblLayout"))
            if layout is not None:
                layout.addprevious(tbl_w)
            else:
                tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(int(round(sum(widths) * TWIPS_PER_CM))))

        for index, width_cm in enumerate(widths):
            if index < len(table.columns):
                table.columns[index].width = Cm(width_cm)

        for row in table.rows:
            for index, cell in enumerate(row.cells):
                if index < len(widths):
                    cell.width = Cm(widths[index])

    def apply_table_fit(self, plan: Any) -> None:
        """
        应用表格适配决策（Program D / WP-D03）。

        参数:
            plan: ``table_fitting.TableFitPlan``（仅使用其列宽序列）。
        """
        self.apply_table_column_widths(getattr(plan, "column_widths_cm", ()))

    def set_table_style(self, style_name: str) -> None:
        if self.current_table:
            self.current_table.style = style_name

    # ============================================================
    # 图片操作
    # ============================================================

    def add_image(
        self,
        src: Union[str, Path],
        alt: str = "",
        width: Optional[Inches] = None,
        height: Optional[Inches] = None,
        bounds: Optional[FigureBounds] = None,
    ) -> Optional[ImagePlacement]:
        """
        插入图片。

        参数:
            src: 图片路径或 ``data:`` URI。
            alt: 替代文本（不可用时作为占位文本）。
            width / height: 传统用法下的显式尺寸（``bounds`` 为 ``None`` 时生效）。
            bounds: 有效内容区适配边界（P12-CAND-002）；给出时按内容区适配插入，
                并返回实际尺寸信息。

        返回:
            Optional[ImagePlacement]: 适配插入结果；传统用法或无法插入时返回 ``None``。
        """
        p = self.add_paragraph()
        try:
            if src.startswith("data:"):
                return self._add_image_from_data_uri(p, src, width, height, bounds)
            elif Path(src).exists():
                return self._add_image_from_file(p, src, width, height, bounds)
            else:
                p.add_run(f"[Image: {alt or src}]")
                return None
        except Exception as e:
            p.add_run(f"[Image error: {e}]")
            return None

    def _add_image_from_file(
        self,
        paragraph: Paragraph,
        src: str,
        width: Optional[Inches] = None,
        height: Optional[Inches] = None,
        bounds: Optional[FigureBounds] = None,
    ) -> Optional[ImagePlacement]:
        run = paragraph.add_run()
        if bounds is not None:
            return self._add_fitted_picture(run, src, None, bounds)
        if width:
            run.add_picture(src, width=width)
        elif height:
            run.add_picture(src, height=height)
        else:
            run.add_picture(src)
        return None

    def _add_image_from_data_uri(
        self,
        paragraph: Paragraph,
        data_uri: str,
        width: Optional[Inches] = None,
        height: Optional[Inches] = None,
        bounds: Optional[FigureBounds] = None,
    ) -> Optional[ImagePlacement]:
        header, data = data_uri.split(",", 1)

        if ";base64" in header:
            img_data = base64.b64decode(data)
        else:
            import urllib.parse

            img_data = urllib.parse.unquote_to_bytes(data)

        is_svg = "svg" in header

        if is_svg:
            png_data = None

            try:
                import cairosvg

                png_data = cairosvg.svg2png(bytestring=img_data)
                print("[WordWriter] SVG->PNG via cairosvg succeeded")
            except ImportError:
                print("[WordWriter] cairosvg not installed")
            except Exception as e:
                print(f"[WordWriter] cairosvg conversion failed: {e}")

            if png_data is None:
                try:
                    from wand.color import Color
                    from wand.image import Image

                    with Image(blob=img_data, format="svg") as img:
                        with Image(
                            width=img.width, height=img.height, background=Color("white")
                        ) as bg:
                            bg.format = "png"
                            bg.composite(img, 0, 0)
                            png_data = bg.make_blob("png")
                    print("[WordWriter] SVG->PNG via wand succeeded")
                except ImportError:
                    print("[WordWriter] wand not installed")
                except Exception as e:
                    print(f"[WordWriter] wand conversion failed: {e}")

            if png_data is None:
                paragraph.add_run(
                    "[SVG image cannot be rendered "
                    "(install cairosvg with Cairo or wand with ImageMagick)]"
                )
                return

            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
                f.write(png_data)
                temp_path = f.name
            raster_bytes = png_data
        else:
            ext = ".png"
            if "jpeg" in header or "jpg" in header:
                ext = ".jpg"
            elif "gif" in header:
                ext = ".gif"
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as f:
                f.write(img_data)
                temp_path = f.name
            raster_bytes = img_data

        try:
            run = paragraph.add_run()
            if bounds is not None:
                return self._add_fitted_picture(run, temp_path, raster_bytes, bounds)
            if width:
                run.add_picture(temp_path, width=width)
            elif height:
                run.add_picture(temp_path, height=height)
            else:
                run.add_picture(temp_path)
            return None
        finally:
            os.unlink(temp_path)

    def _add_fitted_picture(
        self,
        run: Any,
        picture_source: str,
        image_bytes: Optional[bytes],
        bounds: FigureBounds,
    ) -> ImagePlacement:
        """
        按有效内容区适配插入图片（P12-CAND-002）。

        保持宽高比、永不放大、永不超出内容区宽高；固有尺寸不可测量时，
        退化为按目标宽度插入（仍不超出内容区宽度）。

        参数:
            run: 目标 run。
            picture_source: 供 ``add_picture`` 使用的路径（文件路径或临时文件）。
            image_bytes: 已解码的位图字节（``data:`` URI 场景）；文件场景为 ``None``。
            bounds: 有效内容区适配边界。

        返回:
            ImagePlacement: 实际插入尺寸与标记。
        """
        fit: Optional[FigureFit] = None
        try:
            if image_bytes is not None:
                measured = DocxImage.from_blob(image_bytes)
            else:
                measured = DocxImage.from_file(picture_source)
            fit = fit_figure_size(
                px_width=int(measured.px_width),
                px_height=int(measured.px_height),
                target_width_cm=bounds.target_width_cm,
                max_width_cm=bounds.max_width_cm,
                max_height_cm=bounds.max_height_cm,
                min_width_cm=bounds.min_width_cm,
            )
        except Exception:
            fit = None

        if fit is None:
            run.add_picture(picture_source, width=Cm(bounds.target_width_cm))
            return ImagePlacement(
                width_cm=bounds.target_width_cm,
                height_cm=None,
                scaled=False,
                below_min_width=False,
                measured=False,
            )

        run.add_picture(
            picture_source,
            width=Cm(fit.width_cm),
            height=Cm(fit.height_cm),
        )
        return ImagePlacement(
            width_cm=fit.width_cm,
            height_cm=fit.height_cm,
            scaled=fit.scaled,
            below_min_width=fit.below_min_width,
            measured=True,
        )

    # ============================================================
    # 节（Section）操作
    # ============================================================

    def add_section(self) -> Section:
        self.current_section = self.doc.add_section()
        return self.current_section

    def set_page_margins(
        self, top: float = 2.5, bottom: float = 2.5, left: float = 2.5, right: float = 2.5
    ) -> None:
        section = self.current_section or self.doc.sections[0]
        section.top_margin = Cm(top)
        section.bottom_margin = Cm(bottom)
        section.left_margin = Cm(left)
        section.right_margin = Cm(right)

    def set_page_size(self, width: float, height: float) -> None:
        section = self.current_section or self.doc.sections[0]
        section.page_width = Cm(width)
        section.page_height = Cm(height)

    # ============================================================
    # 辅助方法
    # ============================================================

    def get_current_paragraph(self) -> Optional[Paragraph]:
        return self.current_paragraph

    def get_current_table(self) -> Optional[Table]:
        return self.current_table

    def get_current_row(self):
        return self.current_row

    def get_current_cell(self):
        return self.current_cell

    def clear(self) -> None:
        for p in self.doc.paragraphs:
            p._element.getparent().remove(p._element)
        for table in self.doc.tables:
            table._element.getparent().remove(table._element)
        self.current_paragraph = None
        self.current_table = None
        self.current_row = None
        self.current_cell = None
        self._row_idx = 0
        self._col_idx = 0
