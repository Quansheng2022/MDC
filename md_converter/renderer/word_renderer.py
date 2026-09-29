"""
Word Renderer - Word 文档渲染器

使用 Visitor 模式遍历 AST，通过 WordWriter 生成 Word 文档。
样式决策由 StyleResolver 负责，Writer 只负责原子操作。

V1.5 增强（QS-Word-Default-V1.5）:
    - 执行 LayoutPlan（决策与渲染分离）
    - 中英混排 Run 分割（CJK / Latin 分字体）
    - 语言自适应段落对齐（中文两端对齐、西文左对齐）
    - 标题 Keep with next / Keep together
    - 表格重复表头 + 按数据类型对齐
    - 代码 / ASCII 图专用样式
"""

from typing import Dict, List, Optional

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from ..ast.node_visitor import NodeVisitor
from ..ast.nodes import (
    BlockQuote,
    CodeBlock,
    Diagram,
    Document,
    Emphasis,
    HardBreak,
    Heading,
    HorizontalRule,
    Image,
    InlineCode,
    Link,
    ListBlock,
    ListItem,
    Node,
    Paragraph,
    SoftBreak,
    Strong,
    Table,
    TableCell,
    TableRow,
    Text,
)
from .inline_state import InlineState
from .layout.content_analyzer import infer_table_column_types
from .layout.decision_engine import DecisionEngine
from .layout.figure_sizing import CM_PER_INCH, parse_length_cm
from .layout.language_detection import segment_text
from .layout.layout_plan import LayoutPlan
from .layout.section_manager import PageGeometry
from .layout.table_fitting import (
    DEFAULT_MIN_COLUMN_WIDTH_CM,
    DEFAULT_READABILITY_FLOOR_PT,
    TableFitPlan,
    plan_table_fit,
)
from .render_context import RenderContext
from .style_resolver import StyleResolver
from .word_writer import FigureBounds, WordWriter, set_run_font, set_style_font

_ASCII_BOX_CHARS = set("┌┐└┘├┤┬┴┼─│┏┓┗┛┣┫╋╔╗╚╝║═◄►▼▲")


class WordRenderer(NodeVisitor):
    """
    Word 文档渲染器。

    使用 Visitor 模式遍历 AST，将每个节点转换为 Word 文档元素。
    样式决策通过 StyleResolver 完成，WordWriter 只负责底层操作。

    属性:
        ctx: 渲染上下文
        writer: Word 文档写入器
        style_resolver: 样式解析器
        inline_state: 行内样式状态栈
        layout_plan: 布局计划（V1.5）
    """

    def __init__(
        self,
        ctx: RenderContext,
        writer: WordWriter,
        style_resolver: Optional[StyleResolver] = None,
    ):
        """
        初始化渲染器。

        参数:
            ctx: 渲染上下文
            writer: Word 文档写入器
            style_resolver: 样式解析器（可选）
        """
        super().__init__()
        self.ctx = ctx
        self.writer = writer
        self.style_resolver = style_resolver or StyleResolver(ctx.theme)
        self.inline_state = InlineState()
        self._list_counters: Dict[int, int] = {}  # depth -> counter
        self._list_ordered: Dict[int, bool] = {}  # depth -> ordered
        self._current_list_depth = 0
        self._table_cell_align: Optional[str] = None
        self._table_column_types: List[str] = []
        self._table_font_size_override: Optional[float] = None
        self._block_font_kind = "body"
        self.layout_plan: Optional[LayoutPlan] = None
        self._plan_by_node: Dict[int, object] = {}

    # ============================================================
    # 主题辅助
    # ============================================================

    def _theme_value(self, name: str, default: object) -> object:
        """从主题提取属性，不存在时使用默认值。"""
        return getattr(self.ctx.theme, name, default)

    def _body_font_mapping(self) -> Dict[str, str]:
        if hasattr(self.ctx.theme, "font_mapping_for"):
            return self.ctx.theme.font_mapping_for("body")
        return {
            "ascii": self.ctx.theme.body_font,
            "hAnsi": self.ctx.theme.body_font,
            "eastAsia": self.ctx.theme.body_font,
            "cs": self.ctx.theme.body_font,
        }

    def _heading_font_mapping(self) -> Dict[str, str]:
        if hasattr(self.ctx.theme, "font_mapping_for"):
            return self.ctx.theme.font_mapping_for("heading")
        return {
            "ascii": self.ctx.theme.heading_font,
            "hAnsi": self.ctx.theme.heading_font,
            "eastAsia": self.ctx.theme.heading_font,
            "cs": self.ctx.theme.heading_font,
        }

    def _code_font_mapping(self) -> Dict[str, str]:
        font = self._theme_value("code_font", "Consolas")
        return {"ascii": font, "hAnsi": font, "eastAsia": font, "cs": font}

    # ============================================================
    # 主渲染入口
    # ============================================================

    def render(
        self,
        document: Document,
        layout_plan: Optional[LayoutPlan] = None,
    ):
        """
        渲染文档。

        参数:
            document: AST 文档节点
            layout_plan: 布局计划（V1.5）；未提供时自动构建

        返回:
            Document: python-docx 文档对象
        """
        if layout_plan is None:
            engine = DecisionEngine(theme=self.ctx.theme)
            layout_plan = engine.build_plan(document)
        self.layout_plan = layout_plan
        # 按文档顺序将块计划绑定到 AST 节点（HorizontalRule 不参与布局决策）
        plan_iter = iter(layout_plan.blocks)
        self._plan_by_node = {}
        for child in document.children:
            if isinstance(child, HorizontalRule):
                continue
            try:
                self._plan_by_node[id(child)] = next(plan_iter)
            except StopIteration:
                break

        self._configure_document_styles()
        self._configure_heading_styles()
        self.visit(document)

        return self.writer.doc

    def _configure_document_styles(self) -> None:
        """配置 Normal 与目录/标题等样式字体（V1.5 字体映射）。"""
        try:
            normal_style = self.writer.doc.styles["Normal"]
            set_style_font(normal_style, mapping=self._body_font_mapping())
            normal_style.font.size = Pt(self._theme_value("body_size", 10.5))
            normal_style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            normal_style.paragraph_format.line_spacing = self._theme_value("line_spacing", 1.15)
            normal_style.paragraph_format.space_after = Pt(
                self._theme_value("paragraph_space_after", 6)
            )
        except Exception:
            pass

        style_names = [
            "Title",
            "Subtitle",
            "TOC 1",
            "TOC 2",
            "TOC 3",
            "TOC 4",
            "TOC 5",
            "TOC 6",
            "TOC 7",
            "TOC 8",
            "TOC 9",
            "TOC Heading",
            "Caption",
            "List Paragraph",
            "Table Grid",
        ]
        for style_name in style_names:
            try:
                style = self.writer.doc.styles[style_name]
                set_style_font(style, mapping=self._heading_font_mapping())
            except Exception:
                continue

    def _configure_heading_styles(self) -> None:
        """将解析后的主题样式应用到 Word 内置标题样式。"""
        for level in range(1, 7):
            style = self.writer.doc.styles[f"Heading {level}"]
            heading_style = self.style_resolver.heading_style(level)

            set_style_font(style, mapping=self._heading_font_mapping())
            style.font.size = Pt(heading_style["font_size"])
            style.font.bold = heading_style.get("bold", False)
            style.font.italic = heading_style.get("italic", False)
            style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            style.paragraph_format.space_before = heading_style.get("space_before")
            style.paragraph_format.space_after = heading_style.get("space_after")
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.keep_together = True

    # ============================================================
    # 块节点访问方法
    # ============================================================

    def visit_Document(self, node: Document) -> None:
        """渲染文档根节点。"""
        self.writer.start_document()
        self._apply_page_size()
        self._apply_page_margins()

        for child in node.children:
            self.visit(child)

    def _apply_page_size(self) -> None:
        """应用冻结的页面尺寸（主题 ``page.size`` / 配置 ``page_width``；P11-MNT-008）。

        A4 几何由 ``word_writer.set_section_orientation`` 唯一定义：本方法只在
        portrait 文档初始化时应用它，landscape 表格仍通过 ``visit_Table`` 切换到
        同一个 A4 权威。修复前首节保留 python-docx 模板的 US Letter 默认尺寸。
        """
        try:
            size = self._theme_value("page_size", None) or self.ctx.config.get("page_width", "A4")
            orientation = "landscape" if "landscape" in str(size).lower() else "portrait"
            self.writer.set_section_orientation(orientation)
        except Exception as e:
            self.ctx.diag.warning(f"Failed to set page size: {e}", code="RENDER003")

    def visit_Heading(self, node: Heading) -> None:
        """
        渲染标题节点 - 使用 Word 内置标题样式。

        P12-CAND-003（WARN + DROP）：空标题不渲染（不产生段落、不产生占位文本、不递增
        标题计数，因而 TOC 与编号也不会出现合成条目）。AST 仍保留该节点作为源码事实，
        既有 StaticQA ``semantic_empty_heading`` 警告继续作为该状态的诊断。
        """
        text = node.to_plain_text().strip()
        if not text:
            return

        style_name = f"Heading {node.level}"
        self.writer.add_paragraph(text, style=style_name)
        p = self.writer.current_paragraph

        for run in p.runs:
            run.bold = True
            set_run_font(run, mapping=self._heading_font_mapping())

        # Keep with next / Keep together（第十二章）
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True

        # 记录标题编号（用于 TOC）
        self.ctx.heading_counts[node.level] = self.ctx.heading_counts.get(node.level, 0) + 1

    def visit_Paragraph(self, node: Paragraph) -> None:
        """渲染段落节点 - 语言自适应对齐。"""
        text = node.to_plain_text()
        has_line_breaks = any(isinstance(n, (SoftBreak, HardBreak)) for n in node.content)
        style = self.style_resolver.adaptive_paragraph_style(
            text,
            has_line_breaks=has_line_breaks,
        )
        self.writer.add_paragraph()
        p = self.writer.current_paragraph
        self._apply_paragraph_style(p, style)
        self._block_font_kind = "body"
        self._render_inline(node.content, font_kind="body")

        # 行距与段后间距（第十一章）
        p.paragraph_format.line_spacing = self._theme_value("line_spacing", 1.15)
        p.paragraph_format.space_after = Pt(self._theme_value("paragraph_space_after", 6))

        # 布局计划指令
        plan = self._plan_by_node.get(id(node))
        if plan is not None:
            if plan.keep_with_next:
                p.paragraph_format.keep_with_next = True
            if plan.widow_orphan_control:
                p.paragraph_format.widow_control = True

    def visit_BlockQuote(self, node: BlockQuote) -> None:
        """渲染引用块节点 - 增加左边距和斜体。"""
        for child in node.children:
            self.visit(child)
            if self.writer.current_paragraph:
                p = self.writer.current_paragraph
                p.paragraph_format.left_indent = Cm(1.0)
                for run in p.runs:
                    run.italic = True

    def visit_ListBlock(self, node: ListBlock) -> None:
        """渲染列表块节点 - 正确处理列表项。"""
        self._current_list_depth += 1
        depth = self._current_list_depth
        self._list_ordered[depth] = node.ordered
        # 每个列表块独立编号：进入时重置计数器，保证从 1 开始
        self._list_counters[depth] = 0

        for item in node.items:
            self.visit(item)

        self._current_list_depth -= 1

    def visit_ListItem(self, node: ListItem) -> None:
        """渲染列表项节点。"""
        depth = self._current_list_depth
        ordered = self._list_ordered.get(depth, False)

        if ordered:
            self._list_counters[depth] = self._list_counters.get(depth, 0) + 1
            prefix = f"{self._list_counters[depth]}."
        else:
            symbols = ["•", "◦", "▪", "▫"]
            idx = min(depth - 1, len(symbols) - 1)
            prefix = symbols[idx]

        self.writer.add_paragraph()
        p = self.writer.current_paragraph
        p.paragraph_format.left_indent = Cm(0.5 * depth)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.space_after = Pt(3)

        self.writer.add_run(
            prefix + " ",
            font_mapping=self._body_font_mapping(),
            font_size=self._theme_value("body_size", 10.5),
        )

        self._block_font_kind = "body"
        for child in node.children:
            if isinstance(child, Paragraph):
                self._render_inline(child.content, font_kind="body")
            elif isinstance(child, ListBlock):
                self.visit(child)
            else:
                self.visit(child)

    def visit_Table(self, node: Table) -> None:
        """渲染表格节点 - 列宽适配 + 重复表头 + 按数据类型对齐。

        Program D（WP-D03）：在表格创建后、写入单元格前应用
        :func:`plan_table_fit` 的确定性列宽决策（有效内容宽度来自实际 section 几何）。
        单元格文本、行/列结构、样式、边框与顺序均不被修改。
        """
        if not node.rows:
            return

        max_cols = max(len(row.cells) for row in node.rows)
        if max_cols == 0:
            return

        plan = self._find_plan_for(node)
        landscape = bool(plan is not None and plan.landscape)
        if landscape:
            self.writer.add_section()
            self.writer.set_section_orientation("landscape")

        self.writer.start_table(rows=len(node.rows), cols=max_cols)
        self.writer.current_table.style = self.ctx.theme.table_style or "Table Grid"

        self._apply_table_fit(node)

        # 按列推断数据类型（第九章 9.2）
        self._table_column_types = infer_table_column_types(node)

        try:
            for row in node.rows:
                self.visit(row)
        finally:
            self.writer.end_table()
            self._table_font_size_override = None
            if landscape:
                self.writer.add_section()
                self.writer.set_section_orientation("portrait")

    def _apply_table_fit(self, node: Table) -> Optional[TableFitPlan]:
        """计算并应用表格适配决策（Program D / WP-D03）。

        有效内容宽度取自当前 section 几何（``_effective_content_size_cm``），最小列宽与
        字号下限取自解析后的主题值。任何无法计算的情况都降级为「保持原有表格」并产生
        结构化 WARNING，绝不静默修改文档内容。

        参数:
            node: 表格 AST 节点（只读取其单元格文本作为宽度信号）。

        返回:
            Optional[TableFitPlan]: 已应用的决策；未应用时为 ``None``。
        """
        try:
            content_width_cm, _ = self._effective_content_size_cm()
            min_column_width_cm, readability_floor_pt = self._table_fitting_limits()
            rows = [[cell.to_plain_text() for cell in row.cells] for row in node.rows]
            fit = plan_table_fit(
                rows,
                content_width_cm=content_width_cm,
                base_font_size_pt=float(self._theme_value("table_font_size", 9.5)),
                min_column_width_cm=min_column_width_cm,
                readability_floor_pt=readability_floor_pt,
            )
        except Exception as e:  # noqa: BLE001 - 降级必须可见且不致命（SPEC-INV-006）
            self.ctx.diag.warning(
                f"Table fitting skipped, keeping the original table: {e}",
                code="RENDER007",
                location=node.span,
                suggestion=(
                    "Check the effective page geometry and the theme table constraints "
                    "if this table should be fitted"
                ),
            )
            return None

        self.writer.apply_table_fit(fit)
        self._table_font_size_override = fit.font_size_pt

        if fit.squeezed:
            self.ctx.diag.warning(
                "Table cannot fit the effective content width at the minimum column width; "
                "content is preserved and text wraps",
                code="RENDER006",
                location=node.span,
                suggestion=(
                    "Reduce the number of columns, shorten the table, or use a wider output "
                    "profile if the table is unreadable"
                ),
                data=fit.to_dict(),
            )
        return fit

    def _table_fitting_limits(self) -> tuple[float, float]:
        """读取解析后的表格适配下限（最小列宽 / 字号可读性下限）。

        返回:
            tuple[float, float]: ``(min_column_width_cm, readability_floor_pt)``；主题未声明
            时退回冻结主题的默认值。
        """
        min_column_width_cm = DEFAULT_MIN_COLUMN_WIDTH_CM
        data = getattr(self.ctx.theme, "data", None)
        if isinstance(data, dict):
            table_block = data.get("table", {})
            constraints = (
                table_block.get("constraints", {}) if isinstance(table_block, dict) else {}
            )
            parsed = parse_length_cm(constraints.get("min_width"), None)
            if parsed:
                min_column_width_cm = float(parsed)

        readability_floor_pt = DEFAULT_READABILITY_FLOOR_PT
        minimums = getattr(self.ctx.theme, "readability_minimums", None) or {}
        if minimums.get("table_font"):
            readability_floor_pt = float(minimums["table_font"])
        return min_column_width_cm, readability_floor_pt

    def visit_TableRow(self, node: TableRow) -> None:
        """渲染表格行节点。"""
        self.writer.start_row()
        if self.writer._row_idx == 0:
            # 跨页重复表头（第九章 9.4）
            self.writer.set_repeat_table_header(self.writer.current_row)
        for cell in node.cells:
            self.visit(cell)
        self.writer.end_row()

    def visit_TableCell(self, node: TableCell) -> None:
        """渲染表格单元格节点。"""
        self.writer.start_cell()
        col_idx = self.writer._col_idx
        data_type = (
            self._table_column_types[col_idx] if col_idx < len(self._table_column_types) else "text"
        )
        cell_style = self.style_resolver.table_cell_style_by_type(data_type)

        self.writer.set_cell_vertical_alignment(
            self.writer.current_cell, self._theme_value("table_vertical_alignment", "top")
        )
        if self.writer.current_paragraph is not None:
            self.writer.current_paragraph.alignment = cell_style["alignment"]

        self._block_font_kind = "table"
        self._render_inline(node.content, font_kind="table")
        if self.writer._row_idx == 0:
            self._apply_header_cell_style()
        self.writer.end_cell()

    def visit_CodeBlock(self, node: CodeBlock) -> None:
        """渲染代码块节点 - 支持 ASCII 图样式。"""
        if self._looks_like_ascii(node):
            self._render_ascii_block(node)
            return

        style = self.style_resolver.code_style()
        self.writer.add_paragraph()
        p = self.writer.current_paragraph
        p.paragraph_format.left_indent = style.get("indent", Cm(0.5))
        p.paragraph_format.space_before = style.get("space_before", Pt(6))
        p.paragraph_format.space_after = style.get("space_after", Pt(6))
        p.paragraph_format.line_spacing = self._theme_value("code_line_spacing", 1.0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        self.writer.add_run(
            node.text,
            font_mapping=self._code_font_mapping(),
            font_size=style.get("font_size", 10),
        )

        if style.get("background"):
            self.writer.set_paragraph_background(p, style["background"])

        # 短代码块整块同页（第八章）
        if len(node.text.splitlines()) <= self._theme_value("code_keep_together_if_lines", 25):
            self.writer.set_keep_lines(p)

    def _render_ascii_block(self, node: CodeBlock) -> None:
        """渲染 ASCII 图：Consolas 8.5pt、居中。

        注意：不要设置禁止换行（w:noWrap / w:wordWrap=0），
        该属性会导致 Word 在长文档分页时卡死。
        """
        style = self.style_resolver.ascii_style()
        self.writer.add_paragraph()
        p = self.writer.current_paragraph
        p.paragraph_format.space_before = style.get("space_before", Pt(6))
        p.paragraph_format.space_after = style.get("space_after", Pt(6))
        p.paragraph_format.line_spacing = 1.0
        p.alignment = style.get("alignment", WD_ALIGN_PARAGRAPH.LEFT)
        self.writer.set_keep_lines(p)

        for line in node.text.splitlines():
            self.writer.add_run(
                line,
                font_mapping=self._code_font_mapping(),
                font_size=style.get("font_size", 8.5),
            )
            self.writer.add_soft_break()

    def _looks_like_ascii(self, node: CodeBlock) -> bool:
        """判断代码块是否为 ASCII 图（含方框字符）。"""
        if node.language and node.language.lower() not in ("", "text", "plain", "ascii"):
            return False
        return any(ch in _ASCII_BOX_CHARS for ch in node.text)

    def visit_Diagram(self, node: Diagram) -> None:
        """渲染图表节点（通常由 DiagramPass 处理）。"""
        if node.diagram_type == "ascii":
            self.ctx.diag.warning(
                "Unconverted ASCII diagram rendered as monospace text",
                code="RENDER004",
                location=node.span,
                suggestion="Enable AsciiToMermaidPass for automatic conversion",
            )
            self._render_ascii_block(CodeBlock(language="ascii", text=node.content, span=node.span))
            return
        self.ctx.diag.warning(
            "Unprocessed diagram node found in AST",
            code="RENDER001",
            location=node.span,
            suggestion="DiagramPass should have converted this to Image",
        )
        self.writer.add_paragraph()
        lines = node.content.splitlines() if node.content else []
        self.writer.add_run(f"[Diagram: {len(lines)} lines]")

    def visit_Image(self, node: Image) -> None:
        """渲染图片节点 - 按有效内容区适配（P12-CAND-002；Program D / WP-D05）。

        适配决策由 ``figure_sizing.plan_figure_fit`` 唯一给出；本方法与
        ``WordWriter._add_fitted_picture`` 只做集成（无重复适配算术）。
        """
        try:
            bounds = self._figure_bounds_cm()
            self.writer.add_paragraph()
            placement = self.writer.add_image(node.src, node.alt, bounds=bounds)

            if placement is not None and placement.below_min_width:
                self.ctx.diag.warning(
                    "Figure scaled below the theme minimum width",
                    code="RENDER005",
                    location=node.span,
                    suggestion=(
                        "Reduce the figure aspect ratio, or split the source figure, "
                        "if the scaled size is not readable"
                    ),
                    data={
                        "width_cm": round(placement.width_cm, 2),
                        "height_cm": (
                            round(placement.height_cm, 2)
                            if placement.height_cm is not None
                            else None
                        ),
                        "aspect_ratio": (
                            round(placement.aspect_ratio, 6)
                            if placement.aspect_ratio is not None
                            else None
                        ),
                        "height_limited": placement.height_limited,
                        "width_limited": placement.width_limited,
                        "min_width_cm": bounds.min_width_cm,
                    },
                )

            if self.writer.current_paragraph:
                figure_style = self.style_resolver.figure_style()
                if figure_style.get("alignment") is not None:
                    self.writer.current_paragraph.alignment = figure_style["alignment"]
                if figure_style.get("keep_together", True):
                    self.writer.set_keep_lines(self.writer.current_paragraph)

        except Exception as e:
            self.ctx.diag.warning(
                f"Failed to render image: {e}", code="RENDER002", location=node.span
            )
            self.writer.add_paragraph()
            self.writer.add_run(f"[Image: {node.alt}]")

    def _figure_bounds_cm(self) -> FigureBounds:
        """
        计算图形适配边界（P12-CAND-002 / CLAR-02）。

        有效内容区优先取自当前 section 的实际几何（页宽/页高减去页边距）；
        仅在 section 几何不可用时，退回主题/配置的参考 A4 几何（``PageGeometry``）。

        返回:
            FigureBounds: 目标宽度与有效内容区宽高（厘米）。
        """
        policy = self.style_resolver.figure_size_policy()
        available_width, available_height = self._effective_content_size_cm()

        max_width = policy.max_width_cm or available_width
        max_height = policy.max_height_cm or available_height
        configured = self.ctx.config.get("image_width", 5)
        target = min(float(configured) * CM_PER_INCH, float(max_width))

        return FigureBounds(
            target_width_cm=target,
            max_width_cm=float(max_width),
            max_height_cm=float(max_height),
            min_width_cm=policy.min_width_cm,
        )

    def _effective_content_size_cm(self) -> tuple[float, float]:
        """
        返回有效内容区尺寸（厘米）——Program D 的**唯一**页面几何取值点。

        优先使用当前（最后）section 的实际几何：``页宽 − 左右页边距`` 与
        ``页高 − 上下页边距``；section 几何不可用或非正时退回冻结主题的参考 A4 几何
        （``PageGeometry`` + 解析后的主题页边距）。表格适配与图形适配共用本方法，
        因此两者看到的有效宽度始终一致（``SPEC-FUNC-023`` / ``Program D`` §9）。

        返回:
            tuple[float, float]: ``(content_width_cm, content_height_cm)``。
        """
        try:
            section = self.writer.doc.sections[-1]
            available_width = (
                section.page_width.cm - section.left_margin.cm - section.right_margin.cm
            )
            available_height = (
                section.page_height.cm - section.top_margin.cm - section.bottom_margin.cm
            )
            if available_width > 0 and available_height > 0:
                return float(available_width), float(available_height)
        except Exception:
            pass

        geometry = PageGeometry()
        margins = self._theme_value("page_margins_cm", None) or {}
        return (
            geometry.portrait_width_cm
            - float(margins.get("left", 2.54))
            - float(margins.get("right", 2.54)),
            geometry.portrait_height_cm
            - float(margins.get("top", 2.54))
            - float(margins.get("bottom", 2.54)),
        )

    def visit_HorizontalRule(self, node: HorizontalRule) -> None:
        """渲染水平分割线。"""
        self.writer.add_horizontal_rule()

    # ============================================================
    # 行内渲染
    # ============================================================

    def _render_inline(self, nodes: List[Node], font_kind: str = "body") -> None:
        """渲染行内节点列表。"""
        for node in nodes:
            if isinstance(node, Text):
                self._render_text(node, font_kind)
            elif isinstance(node, Strong):
                self.inline_state.push(bold=True)
                self._render_inline(node.content, font_kind)
                self.inline_state.pop()
            elif isinstance(node, Emphasis):
                self.inline_state.push(italic=True)
                self._render_inline(node.content, font_kind)
                self.inline_state.pop()
            elif isinstance(node, InlineCode):
                self._render_inline_code(node)
            elif isinstance(node, Link):
                self._render_link(node, font_kind)
            elif isinstance(node, Image):
                self.visit_Image(node)
            elif isinstance(node, SoftBreak):
                self.writer.add_soft_break()
            elif isinstance(node, HardBreak):
                self.writer.add_hard_break()
            else:
                self._render_inline(list(node.iter_children()), font_kind)

    def _render_text(self, node: Text, font_kind: str = "body") -> None:
        """渲染文本节点 - 中英混排 Run 分割。"""
        style = self.inline_state.current_style()
        mapping = self._mapping_for(font_kind)
        font_size = self._size_for(font_kind)

        for seg in segment_text(node.content):
            self.writer.add_run(
                seg.text,
                bold=style.get("bold", False),
                italic=style.get("italic", False),
                underline=style.get("underline", False),
                font_mapping=mapping,
                font_size=font_size,
                color=style.get("color"),
            )

    def _mapping_for(self, font_kind: str) -> Dict[str, str]:
        if font_kind == "heading":
            return self._heading_font_mapping()
        if font_kind == "code":
            return self._code_font_mapping()
        return self._body_font_mapping()

    def _size_for(self, font_kind: str) -> float:
        if font_kind == "table":
            if self._table_font_size_override is not None:
                return self._table_font_size_override
            return self._theme_value("table_font_size", 9.5)
        if font_kind == "code":
            return self._theme_value("code_size", 10)
        return self._theme_value("body_size", 10.5)

    def _render_inline_code(self, node: InlineCode) -> None:
        style = self.inline_state.current_style()
        self.writer.add_run(
            node.text,
            bold=style.get("bold", False),
            italic=style.get("italic", False),
            font_mapping=self._code_font_mapping(),
            font_size=self._theme_value("code_size", 10),
            color=RGBColor(0x80, 0x00, 0x00),
        )

    def _render_link(self, node: Link, font_kind: str = "body") -> None:
        """渲染链接：保留可见文本/样式，并写入真实超链接目标（P11-MNT-007）。

        文本与样式沿用既有行为；渲染完成后把新增 run 包裹进 ``w:hyperlink``，
        使 ``http/https``、``mailto`` 等目标在最终 DOCX 中可恢复。
        """
        paragraph = self.writer.current_paragraph
        start_index = len(paragraph.runs) if paragraph is not None else 0
        self.inline_state.push(color=RGBColor(0x00, 0x00, 0xFF), underline=True)
        self._render_inline(node.content, font_kind)
        self.inline_state.pop()
        if paragraph is not None:
            self.writer.wrap_runs_as_hyperlink(paragraph, start_index, node.href)

    # ============================================================
    # 样式辅助方法
    # ============================================================

    def _apply_page_margins(self) -> None:
        """应用页面边距（V1.5：1in）。"""
        try:
            margins = self._theme_value("page_margins_cm", None)
            if margins is None:
                margins = self.ctx.config.get("page_margins", {})
                margins = {
                    "top": margins.get("top", 2.54),
                    "bottom": margins.get("bottom", 2.54),
                    "left": margins.get("left", 2.54),
                    "right": margins.get("right", 2.54),
                }
            section = self.writer.doc.sections[0]
            section.top_margin = Cm(margins["top"])
            section.bottom_margin = Cm(margins["bottom"])
            section.left_margin = Cm(margins["left"])
            section.right_margin = Cm(margins["right"])
        except Exception as e:
            self.ctx.diag.warning(f"Failed to set page margins: {e}", code="RENDER003")

    def _apply_paragraph_style(self, paragraph, style: Dict) -> None:
        """应用段落样式字典。"""
        if style.get("alignment") is not None:
            paragraph.alignment = style["alignment"]
        if style.get("space_before") is not None:
            paragraph.paragraph_format.space_before = style["space_before"]
        if style.get("space_after") is not None:
            paragraph.paragraph_format.space_after = style["space_after"]
        if style.get("left_indent") is not None:
            paragraph.paragraph_format.left_indent = style["left_indent"]
        if style.get("line_spacing") is not None:
            paragraph.paragraph_format.line_spacing = style["line_spacing"]

    def _apply_header_cell_style(self) -> None:
        if self.writer.current_cell:
            cell = self.writer.current_cell
            self.writer.set_cell_background(cell, self.ctx.theme.table_header_bg)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = self.ctx.theme.table_header_fg
                    run.bold = True

    def _find_plan_for(self, node: Node):
        """按 AST 节点对象身份查找布局计划。"""
        return self._plan_by_node.get(id(node))

    def generic_visit(self, node: Node) -> None:
        for child in node.iter_children():
            self.visit(child)
