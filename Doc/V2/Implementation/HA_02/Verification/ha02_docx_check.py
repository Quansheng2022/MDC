"""HA-02 DOCX Mermaid verification (verification-only).

Inspects a converted DOCX and reports whether its Mermaid figures are real
renders or the product's raw-source text fallback.

The fallback image produced by ``DiagramPass._create_text_fallback_image`` is an
SVG with a ``#f8f9fa`` background and a "rendering failed, showing raw text"
warning line, rasterised to PNG before embedding. A real Mermaid render has a
white background, far more distinct colours, and no raw source text.

Usage:
    python ha02_docx_check.py <output.docx> [--media-dir <dir>] [--json]
"""

from __future__ import annotations

import argparse
import collections
import json
import struct
import sys
import zipfile
import zlib
from pathlib import Path

FALLBACK_BACKGROUND = (0xF8, 0xF9, 0xFA)
#: Unambiguous Mermaid source syntax.  Bare words such as "subgraph" also
#: occur in ordinary prose, so they are deliberately not used as markers.
RAW_SOURCE_MARKERS = ("flowchart TB", "flowchart TD", "flowchart LR", "-->", "~~~")
FALLBACK_TEXT_MARKERS = ("rendering failed", "showing raw text")


def load_png(path: Path) -> tuple[int, int, int, bytes]:
    """Decode a non-interlaced PNG into (width, height, channels, pixels)."""
    data = path.read_bytes()
    pos, idat = 8, b""
    width = height = channels = 0
    while pos < len(data):
        length = struct.unpack(">I", data[pos : pos + 4])[0]
        kind = data[pos + 4 : pos + 8]
        chunk = data[pos + 8 : pos + 8 + length]
        if kind == b"IHDR":
            width, height, _depth, colour, _comp, _filt, _inter = struct.unpack(">IIBBBBB", chunk)
            channels = {0: 1, 2: 3, 3: 1, 4: 2, 6: 4}[colour]
        elif kind == b"IDAT":
            idat += chunk
        elif kind == b"IEND":
            break
        pos += 12 + length
    raw = zlib.decompress(idat)
    stride = width * channels
    out = bytearray()
    prev = bytearray(stride)
    index = 0
    for _ in range(height):
        filt = raw[index]
        index += 1
        line = bytearray(raw[index : index + stride])
        index += stride
        for x in range(stride):
            left = line[x - channels] if x >= channels else 0
            up = prev[x]
            upleft = prev[x - channels] if x >= channels else 0
            if filt == 1:
                line[x] = (line[x] + left) & 255
            elif filt == 2:
                line[x] = (line[x] + up) & 255
            elif filt == 3:
                line[x] = (line[x] + (left + up) // 2) & 255
            elif filt == 4:
                est = left + up - upleft
                pa, pb, pc = abs(est - left), abs(est - up), abs(est - upleft)
                pred = left if (pa <= pb and pa <= pc) else (up if pb <= pc else upleft)
                line[x] = (line[x] + pred) & 255
        out += line
        prev = line
    return width, height, channels, bytes(out)


def image_stats(path: Path) -> dict:
    """Summarise one embedded image."""
    width, height, channels, pixels = load_png(path)
    colours: collections.Counter = collections.Counter()
    for y in range(0, height, 2):
        row = y * width * channels
        for x in range(0, width, 2):
            offset = row + x * channels
            colours[tuple(pixels[offset : offset + 3])] += 1
    sampled = sum(colours.values())
    top = colours.most_common(3)
    non_white = sum(v for k, v in colours.items() if not all(c > 240 for c in k))
    background = top[0][0] if top else (0, 0, 0)
    return {
        "file": path.name,
        "pixels": f"{width}x{height}",
        "distinct_colours": len(colours),
        "non_white_ratio": round(non_white / sampled, 4) if sampled else 0.0,
        "background": "#%02x%02x%02x" % background,
        "fallback_background": background == FALLBACK_BACKGROUND,
    }


def check(docx_path: Path, media_dir: Path | None) -> dict:
    """Report the Mermaid rendering facts for one DOCX."""
    import docx  # imported here so the module stays importable without the dep

    document = docx.Document(str(docx_path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    raw_markers = {marker: text.count(marker) for marker in RAW_SOURCE_MARKERS}
    fallback_markers = {marker: text.count(marker) for marker in FALLBACK_TEXT_MARKERS}

    archive = zipfile.ZipFile(docx_path)
    media = sorted(n for n in archive.namelist() if n.startswith("word/media/"))
    stats: list[dict] = []
    if media_dir is not None:
        media_dir.mkdir(parents=True, exist_ok=True)
        for name in media:
            target = media_dir / Path(name).name
            target.write_bytes(archive.read(name))
            stats.append(image_stats(target))

    sizes = [
        {
            "width_cm": round(shape.width.cm, 3),
            "height_cm": round(shape.height.cm, 3),
        }
        for shape in document.inline_shapes
    ]
    fallback_images = [s["file"] for s in stats if s["fallback_background"]]
    return {
        "docx": str(docx_path),
        "paragraphs": len(document.paragraphs),
        "inline_images": len(document.inline_shapes),
        "image_sizes_cm": sizes,
        "media_files": media,
        "media_stats": stats,
        "raw_source_marker_counts": raw_markers,
        "fallback_text_marker_counts": fallback_markers,
        "fallback_images": fallback_images,
        "rendered_mermaid": bool(
            document.inline_shapes
            and media
            and not fallback_images
            and not any(raw_markers.values())
            and not any(fallback_markers.values())
        ),
    }


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx")
    parser.add_argument("--media-dir", default=None)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)
    report = check(Path(args.docx), Path(args.media_dir) if args.media_dir else None)
    if args.json:
        print(json.dumps(report, indent=2))
    else:
        images = report["inline_images"]
        rendered = report["rendered_mermaid"]
        print(f"{report['docx']}: images={images} rendered={rendered}")
        for stat in report["media_stats"]:
            print(
                f"  {stat['file']}: {stat['pixels']} bg={stat['background']} "
                f"colours={stat['distinct_colours']} fallback={stat['fallback_background']}"
            )
    return 0 if report["rendered_mermaid"] else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
