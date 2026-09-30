"""R2-V03 native-Word acceptance-set assembly and artifact index.

Verification-only helper for the R2-V03 gate. It copies the accepted R2-V02
packaged outputs (plus one already-accepted Program D table fixture) into a
single neutral acceptance-set directory and records a compact structural index
that the native Word inspection refers to.

The script never touches the repository product source, the installed product,
or any product setting. It reads artifacts, copies them, and writes one JSON
index.

Usage::

    python r2v03_acceptance_set.py --dest <dir> --index <json> --item role=<path>
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Tuple

from docx import Document
from lxml import etree

__all__ = ["main"]


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W_P = f"{{{W}}}p"
W_PPR = f"{{{W}}}pPr"
W_PSTYLE = f"{{{W}}}pStyle"
W_VAL = f"{{{W}}}val"
W_T = f"{{{W}}}t"
W_TAB = f"{{{W}}}tab"
W_INSTR = f"{{{W}}}instrText"
W_TBLGRID = f"{{{W}}}tblGrid"
W_GRIDCOL = f"{{{W}}}gridCol"
W_W = f"{{{W}}}w"

A_DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"

TWIPS_PER_CM = 1440.0 / 2.54
EMU_PER_CM = 360000.0


def sha256_of(path: Path) -> str:
    """Return the uppercase SHA-256 hex digest of ``path``."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def document_xml(path: Path) -> str:
    """Return ``word/document.xml`` decoded as text (``""`` when absent)."""
    with zipfile.ZipFile(path) as archive:
        if "word/document.xml" not in archive.namelist():
            return ""
        return archive.read("word/document.xml").decode("utf-8", "replace")


def valid_docx(path: Path) -> bool:
    """Return True when ``path`` is a readable OOXML package with a body."""
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names:
                return False
            root = etree.fromstring(archive.read("word/document.xml"))
        return root.find(f"{{{W}}}body") is not None
    except Exception:  # noqa: BLE001 - integrity probe must never raise
        return False


def toc_field_instructions(path: Path) -> List[str]:
    """Return every TOC field instruction string found in the document body."""
    root = etree.fromstring(document_xml(path).encode("utf-8"))
    instructions: List[str] = []
    for node in root.iter(W_INSTR):
        text = (node.text or "").strip()
        if text.upper().startswith("TOC"):
            instructions.append(text)
    return instructions


def _paragraph_style_id(paragraph: etree._Element) -> str:
    properties = paragraph.find(W_PPR)
    if properties is None:
        return ""
    style = properties.find(W_PSTYLE)
    if style is None:
        return ""
    return (style.get(W_VAL) or "").strip().lower()


def _paragraph_text_to_tab(paragraph: etree._Element) -> str:
    """Return the paragraph text up to the tab that precedes its page number.

    When Word is installed, the packaged product refreshes the inserted TOC
    through Word, and Word rewrites every cached entry as a hyperlink.  The
    entry runs therefore sit inside ``w:hyperlink`` rather than directly under
    the paragraph, so the whole paragraph subtree is walked in document order
    and the first ``w:tab`` (the page-number tab) terminates the entry text.
    """
    parts: List[str] = []
    for node in paragraph.iter():
        if node.tag == W_PPR or W_PPR in [ancestor.tag for ancestor in node.iterancestors()]:
            continue
        if node.tag == W_TAB:
            break
        if node.tag == W_T:
            parts.append(node.text or "")
    return "".join(parts)


def toc_layout(path: Path) -> Dict[str, Any]:
    """Return the TOC heading, entry count and first entries from the XML."""
    root = etree.fromstring(document_xml(path).encode("utf-8"))
    heading = ""
    entries: List[str] = []
    for paragraph in root.iter(W_P):
        style_id = _paragraph_style_id(paragraph)
        if style_id == "tocheading":
            heading = "".join(node.text or "" for node in paragraph.iter(W_T))
            continue
        if style_id.startswith("toc"):
            entries.append(_paragraph_text_to_tab(paragraph))
    return {
        "heading": heading,
        "entry_count": len(entries),
        "entries": [entry for entry in entries if entry][:12],
        "field_instructions": toc_field_instructions(path),
    }


def table_summary(document: Any) -> List[Dict[str, Any]]:
    """Return column count and grid widths (cm) for every table."""
    summaries: List[Dict[str, Any]] = []
    for table in document.tables:
        widths: List[float] = []
        grid = table._tbl.find(W_TBLGRID)
        if grid is not None:
            for column in grid.iter(W_GRIDCOL):
                raw = column.get(W_W)
                if raw and raw.isdigit():
                    widths.append(round(int(raw) / TWIPS_PER_CM, 4))
        if not widths:
            widths = [round(col.width.cm, 4) for col in table.columns if col.width]
        summaries.append(
            {
                "rows": len(table.rows),
                "columns": len(table.columns),
                "grid_widths_cm": widths,
                "total_width_cm": round(sum(widths), 4) if widths else None,
                "style": table.style.name if table.style is not None else None,
            }
        )
    return summaries


def figure_summary(path: Path, document: Any) -> Dict[str, Any]:
    """Return inline figure sizes in cm plus the raw drawing count."""
    shapes: List[Dict[str, Any]] = []
    for shape in document.inline_shapes:
        width_cm = round(shape.width / EMU_PER_CM, 4)
        height_cm = round(shape.height / EMU_PER_CM, 4)
        shapes.append(
            {
                "width_cm": width_cm,
                "height_cm": height_cm,
                "ratio": round(width_cm / height_cm, 4) if height_cm else None,
            }
        )
    root = etree.fromstring(document_xml(path).encode("utf-8"))
    drawings = len(list(root.iter(f"{{{A_DRAWING}}}inline"))) + len(
        list(root.iter(f"{{{A_DRAWING}}}anchor"))
    )
    return {"inline_shapes": shapes, "drawing_elements": drawings}


def measure(path: Path) -> Dict[str, Any]:
    """Return the compact structural measurement of one DOCX artifact."""
    document = Document(str(path))
    sections = [
        {
            "left_cm": round(section.left_margin.cm, 3),
            "right_cm": round(section.right_margin.cm, 3),
            "top_cm": round(section.top_margin.cm, 3),
            "bottom_cm": round(section.bottom_margin.cm, 3),
            "page_width_cm": round(section.page_width.cm, 3),
            "page_height_cm": round(section.page_height.cm, 3),
            "orientation": str(section.orientation),
        }
        for section in document.sections
    ]
    content_width = None
    if sections:
        primary = sections[-1] if len(sections) > 1 else sections[0]
        content_width = round(
            primary["page_width_cm"] - primary["left_cm"] - primary["right_cm"], 3
        )
    return {
        "paragraph_count": len(document.paragraphs),
        "sections": sections,
        "primary_content_width_cm": content_width,
        "toc": toc_layout(path),
        "tables": table_summary(document),
        "figures": figure_summary(path, document),
    }


def parse_item(raw: str) -> Tuple[str, Path]:
    """Parse one ``role=path`` acceptance-set item."""
    if "=" not in raw:
        raise argparse.ArgumentTypeError(f"expected role=path, got {raw!r}")
    role, _, value = raw.partition("=")
    return role.strip(), Path(value.strip())


def main() -> int:
    """Assemble the acceptance set and write its index JSON."""
    parser = argparse.ArgumentParser(description="R2-V03 acceptance-set assembly")
    parser.add_argument("--dest", required=True, type=Path)
    parser.add_argument("--index", required=True, type=Path)
    parser.add_argument("--item", required=True, action="append", type=parse_item)
    args = parser.parse_args()

    dest: Path = args.dest
    dest.mkdir(parents=True, exist_ok=True)

    items: List[Dict[str, Any]] = []
    failures: List[str] = []
    for role, source in args.item:
        if not source.is_file():
            failures.append(f"missing source for {role}: {source}")
            continue
        target = dest / source.name
        shutil.copy2(source, target)
        ok = valid_docx(target)
        entry: Dict[str, Any] = {
            "role": role,
            "name": source.name,
            "source_path": str(source),
            "acceptance_path": str(target),
            "bytes": target.stat().st_size,
            "sha256": sha256_of(target),
            "source_sha256": sha256_of(source),
            "valid_docx": ok,
        }
        if ok:
            entry["structure"] = measure(target)
        items.append(entry)

    index = {
        "program": "R2-V03",
        "work_package": "R2V03-01",
        "purpose": "native Word openability / visual acceptance set",
        "acceptance_directory": str(dest),
        "item_count": len(items),
        "invalid_count": sum(1 for item in items if not item["valid_docx"]),
        "failures": failures,
        "items": items,
    }
    args.index.parent.mkdir(parents=True, exist_ok=True)
    args.index.write_text(
        json.dumps(index, indent=4, ensure_ascii=False), encoding="utf-8"
    )
    print(f"items={len(items)} invalid={index['invalid_count']} failures={len(failures)}")
    print(f"index={args.index}")
    return 0 if not failures and index["invalid_count"] == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
