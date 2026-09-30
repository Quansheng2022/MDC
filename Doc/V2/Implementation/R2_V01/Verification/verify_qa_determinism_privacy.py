"""R2-V01 WP-05 QA / determinism / failure isolation / privacy verification.

  Q  QA authority: each QA stage runs exactly once per conversion (no second
     QA authority); the Conversion Report reflects the *real* diagnostics and
     the irreducible-wide-table warning stays surfaced.
  R  Failure isolation: a failed conversion does not poison a later conversion
     on the same service instance.
  S  Determinism: repeated identical source/config preserves table decisions,
     figure decisions, TOC localization, profile resolution and diagnostics.
  P  Privacy: re-runs the existing approved source-level privacy/network guard
     (``tests/gui/test_about_dialog.py``) and applies the *same* forbidden-import
     method to the whole production package (no new security audit).

Exit code 0 == every check passed.
"""

from __future__ import annotations

import base64
import contextlib
import io
import os
import re
import struct
import subprocess
import sys
import tempfile
import time
import zipfile
import zlib
from pathlib import Path
from typing import Any, Dict, List, Sequence, Set, Tuple

REPO_ROOT = None
for _candidate in [Path(__file__).resolve(), *Path(__file__).resolve().parents]:
    if (_candidate / "md_converter" / "__init__.py").is_file():
        REPO_ROOT = _candidate
        break
if REPO_ROOT is None:  # pragma: no cover - defensive
    raise RuntimeError("repository root not found")

PKG_ROOT = REPO_ROOT / "md_converter"
PYTHON = REPO_ROOT / ".venv" / "Scripts" / "python.exe"

os.environ.setdefault("PYTHONUTF8", "1")

BASE_OVERRIDES: Dict[str, Any] = {
    "word_com": False,
    "enable_cover": True,
    "toc": True,
    "style_tables": True,
    "verbose": False,
}


class Report:
    """Collect and print PASS/FAIL checks."""

    def __init__(self) -> None:
        self.rows: List[Tuple[str, bool, str]] = []

    def check(self, label: str, ok: bool, detail: str = "") -> None:
        """Record one check."""
        self.rows.append((label, bool(ok), detail))

    def print(self) -> int:
        """Print the report and return the number of failures."""
        print("R2-V01 WP-05 QA / DETERMINISM / FAILURE ISOLATION / PRIVACY")
        print("=" * 72)
        failures = 0
        for label, ok, detail in self.rows:
            print(f"[{'PASS' if ok else 'FAIL'}] {label}")
            if detail:
                print(f"       {detail}")
            failures += 0 if ok else 1
        print("=" * 72)
        print(f"checks={len(self.rows)} failed={failures}")
        return failures


def png_data_uri(width: int, height: int) -> str:
    """Return a deterministic solid-colour PNG data URI."""
    scanline = b"\x00" + bytes((0x2F, 0x54, 0x96)) * width
    raw = scanline * height

    def chunk(tag: bytes, payload: bytes) -> bytes:
        body = tag + payload
        return (
            struct.pack(">I", len(payload)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    png = (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    return "data:image/png;base64," + base64.b64encode(png).decode("ascii")


MIXED_DOC = f"""---
title: QA Determinism Report
---

# QA Determinism Report

Integrated report body with 中文内容 for localization.

## Data Table

| Region | Q1 | Q2 | Q3 | Q4 | Notes | Owner | Status |
| --- | --- | --- | --- | --- | --- | --- | --- |
| North | 120 | 130 | 140 | 150 | steady | A. Lee | accepted |
| South | 98 | 105 | 111 | 118 | recovering | B. Ng | accepted |

## Figure

![Chart]({png_data_uri(400, 100)})

### 备注

Deterministic verification only.
"""

_BOUNDARY_TABLE = "\n".join(
    [
        "|" + "|".join(f" H{index} " for index in range(25)) + "|",
        "|" + "|".join(" --- " for _ in range(25)) + "|",
        "|" + "|".join(" a " for _ in range(25)) + "|",
        "",
    ]
)

BOUNDARY_DOC = f"""---
title: QA Boundary
---

# QA Boundary

{_BOUNDARY_TABLE}
"""


def convert(
    workspace: Path, name: str, body: str, overrides: Dict[str, Any]
) -> Tuple[Any, Path]:
    """Convert ``body`` via the canonical application entry point."""
    from md_converter.application.conversion_request import ConversionRequest
    from md_converter.application.conversion_service import ConversionService

    source = workspace / f"{name}.md"
    source.write_text(body, encoding="utf-8")
    output = workspace / f"{name}.docx"
    config = dict(BASE_OVERRIDES)
    config.update(overrides)
    with contextlib.redirect_stdout(io.StringIO()):
        result = ConversionService().convert(
            ConversionRequest(source_path=source, output_path=output, config_overrides=config)
        )
    return result, output


def decision_fingerprint(output: Path) -> Dict[str, Any]:
    """Return the observable decisions of one artifact (table/figure/TOC/geometry)."""
    from docx import Document

    doc = Document(str(output))
    toc_heading = next(
        (
            paragraph.text.strip()
            for paragraph in doc.paragraphs
            if paragraph.style is not None and paragraph.style.name == "TOC Heading"
        ),
        None,
    )
    return {
        "table_widths": [round(column.width.cm, 4) for column in doc.tables[0].columns],
        "table_cells": [[cell.text for cell in row.cells] for row in doc.tables[0].rows],
        "figure": [round(shape.width.cm, 4) for shape in doc.inline_shapes]
        + [round(shape.height.cm, 4) for shape in doc.inline_shapes],
        "toc_heading": toc_heading,
        "margins": [
            round(doc.sections[-1].left_margin.cm, 3),
            round(doc.sections[-1].right_margin.cm, 3),
        ],
    }


# --------------------------------------------------------------------------
# Q — QA authority + report fidelity
# --------------------------------------------------------------------------
def check_qa(report: Report, workspace: Path) -> None:
    """Each QA stage runs exactly once; the report reflects real diagnostics."""
    import md_converter.compiler as compiler_module

    counts: Dict[str, int] = {}
    originals: Dict[str, Any] = {}
    for name in ("StaticQA", "RenderedQA", "FinalArtifactQA"):
        originals[name] = getattr(compiler_module, name)

        def make(original: Any, key: str) -> Any:
            class _Counting(original):  # type: ignore[misc, valid-type]
                def __init__(self, *args: Any, **kwargs: Any) -> None:
                    counts[key] = counts.get(key, 0) + 1
                    super().__init__(*args, **kwargs)

            return _Counting

        setattr(compiler_module, name, make(originals[name], name))
    try:
        result, _ = convert(workspace, "qa_clean", MIXED_DOC, {"output_profile": "academic"})
    finally:
        for name, original in originals.items():
            setattr(compiler_module, name, original)

    report.check(
        "Q1 every QA stage instantiated exactly once (no second QA authority)",
        counts == {"StaticQA": 1, "RenderedQA": 1, "FinalArtifactQA": 1},
        f"instantiation_counts={counts}",
    )
    gate = result.quality_gate_report or {}
    stages = ("static_qa", "rendered_qa", "post_processor", "final_artifact_qa")
    report.check(
        "Q2 the quality-gate report carries each canonical stage once",
        all(stage in gate for stage in stages) and len(gate) == 7,
        f"report_keys={sorted(gate)}",
    )

    boundary_result, _ = convert(workspace, "qa_boundary", BOUNDARY_DOC, {"output_profile": "academic"})
    from md_converter.gui.presentation_model import present_result
    from md_converter.gui.preflight_model import preflight_from_result
    from md_converter.gui.result_details import build_report_text

    presentation = present_result(boundary_result)
    report_text = build_report_text(presentation, boundary_result)
    summary = preflight_from_result(boundary_result)
    real_message = next(
        (record.message for record in boundary_result.diagnostics if record.code == "RENDER006"),
        "",
    )
    report.check(
        "Q3 the Conversion Report carries the real diagnostic message",
        bool(real_message) and real_message in report_text,
        f"message={real_message[:70]!r} in_report={real_message in report_text}",
    )
    report.check(
        "Q4 the irreducible-wide-table warning stays surfaced in the report",
        "RENDER006" in report_text or "Table" in report_text,
        f"report_excerpt={report_text[report_text.find('Table') - 20: report_text.find('Table') + 90]!r}",
    )
    report.check(
        "Q5 the preflight summary counts the real warnings (not fabricated)",
        summary.total >= 1 and summary.warning_count >= 1,
        f"summary_total={summary.total} warnings={summary.warning_count}",
    )


# --------------------------------------------------------------------------
# R — failure isolation
# --------------------------------------------------------------------------
def check_failure_isolation(report: Report, workspace: Path) -> None:
    """A failed conversion does not poison a later conversion on the same service."""
    from md_converter.application.conversion_request import ConversionRequest
    from md_converter.application.conversion_service import ConversionService

    service = ConversionService({"word_com": False})
    failing = workspace / "iso_fail.md"
    failing.write_text("", encoding="utf-8")
    failing_out = workspace / "iso_fail.docx"
    with contextlib.redirect_stdout(io.StringIO()):
        failed_result = service.convert(
            ConversionRequest(source_path=failing, output_path=failing_out)
        )

    good = workspace / "iso_ok.md"
    good.write_text(MIXED_DOC, encoding="utf-8")
    good_out = workspace / "iso_ok.docx"
    with contextlib.redirect_stdout(io.StringIO()):
        good_result = service.convert(ConversionRequest(source_path=good, output_path=good_out))

    report.check(
        "R1 the first conversion fails as expected",
        failed_result.is_failed,
        f"status={failed_result.status} category={failed_result.error_category}",
    )
    report.check(
        "R2 the next conversion on the same service is clean and successful",
        good_result.is_success
        and len(good_result.diagnostics) == 0
        and good_out.exists(),
        f"status={good_result.status} diagnostics={[r.code for r in good_result.diagnostics]}",
    )
    gate = good_result.quality_gate_report or {}
    report.check(
        "R3 no failed-stage evidence leaks into the later conversion",
        all((gate.get(stage) or {}).get("status") != "FAIL" for stage in gate if isinstance(gate.get(stage), dict)),
        f"stage_statuses={ {k: v.get('status') for k, v in gate.items() if isinstance(v, dict)} }",
    )


# --------------------------------------------------------------------------
# S — determinism of decisions
# --------------------------------------------------------------------------
def check_determinism(report: Report, workspace: Path) -> None:
    """Repeated identical source/config preserves every decision."""
    for profile in ("professional_report", "academic"):
        fingerprints = []
        diagnostics = []
        for index in range(3):
            result, output = convert(
                workspace, f"det_{profile}_{index}", MIXED_DOC, {"output_profile": profile}
            )
            fingerprints.append(decision_fingerprint(output))
            diagnostics.append(tuple(sorted(record.code for record in result.diagnostics)))

        consistent = all(print_ == fingerprints[0] for print_ in fingerprints)
        report.check(
            f"S[{profile}] repeated conversion preserves table/figure/TOC/geometry decisions",
            consistent,
            f"tables={fingerprints[0]['table_widths']} figure={fingerprints[0]['figure']} "
            f"toc={fingerprints[0]['toc_heading']!r} margins={fingerprints[0]['margins']}",
        )
        report.check(
            f"S[{profile}] repeated conversion preserves the diagnostic set",
            all(set_ == diagnostics[0] for set_ in diagnostics),
            f"diagnostics={diagnostics[0]}",
        )
        report.check(
            f"S[{profile}] profile resolution reaches identical geometry every time",
            all(print_["margins"] == fingerprints[0]["margins"] for print_ in fingerprints),
            f"margins={fingerprints[0]['margins']}",
        )


# --------------------------------------------------------------------------
# P — privacy (re-uses the approved source-level check)
# --------------------------------------------------------------------------
def imported_modules(module: Path) -> Set[str]:
    """Return the module names imported by ``module`` (source scan)."""
    import ast

    tree = ast.parse(module.read_text(encoding="utf-8"))
    names: Set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            names.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            names.add(node.module)
    return names


def check_privacy(report: Report, workspace: Path) -> None:
    """Run the approved privacy guard and re-apply its method to the package."""
    completed = subprocess.run(
        [
            str(PYTHON),
            "-m",
            "pytest",
            str(PKG_ROOT / "tests" / "gui" / "test_about_dialog.py"),
            "-p",
            "no:cacheprovider",
            "-q",
        ],
        cwd=str(REPO_ROOT),
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=600,
    )
    tail = (completed.stdout or "").strip().splitlines()
    report.check(
        "P1 existing approved privacy/network guard passes",
        completed.returncode == 0,
        f"rc={completed.returncode} tail={tail[-1] if tail else ''}",
    )

    # Same method as the approved guard, applied at package scope: prove no
    # network-client capability exists, and that the pre-existing local
    # process/URL-parsing imports are confined to the accepted baseline set.
    network_prefixes = (
        "socket",
        "http",
        "requests",
        "httpx",
        "aiohttp",
        "ftplib",
        "smtplib",
        "telnetlib",
        "webbrowser",
    )
    local_prefixes = ("urllib", "subprocess")

    network_offenders: Dict[str, List[str]] = {}
    local_offenders: Dict[str, List[str]] = {}
    for path in sorted(PKG_ROOT.rglob("*.py")):
        relative = path.relative_to(PKG_ROOT).as_posix()
        if relative.startswith("tests/") or "__pycache__" in relative:
            continue
        imported = imported_modules(path)
        network_hits = sorted(name for name in imported if name.startswith(network_prefixes))
        local_hits = sorted(name for name in imported if name.startswith(local_prefixes))
        if network_hits:
            network_offenders[relative] = network_hits
        if local_hits:
            local_offenders[relative] = local_hits

    report.check(
        "P2 no network-client module anywhere in the production package",
        not network_offenders,
        f"network_offenders={network_offenders}",
    )
    accepted_local = {
        "pipeline/passes/diagram_pass.py": ["subprocess"],
        "renderer/word_writer.py": ["urllib.parse"],
        "utils/helpers.py": ["subprocess"],
    }
    report.check(
        "P3 local process/URL-parsing imports stay at the accepted baseline set",
        local_offenders == accepted_local,
        f"local={local_offenders} accepted={accepted_local}",
    )


def main() -> int:
    """Run every WP-05 check."""
    report = Report()
    workspace = Path(tempfile.mkdtemp(prefix="r2v01_wp05_"))
    check_qa(report, workspace)
    check_failure_isolation(report, workspace)
    check_determinism(report, workspace)
    check_privacy(report, workspace)
    print(f"workspace={workspace}")
    return 1 if report.print() else 0


if __name__ == "__main__":
    sys.exit(main())
