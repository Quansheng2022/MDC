"""R2-V01 WP-03 cross-feature document integration verification.

Converts representative multi-feature documents through the accepted
application entry point (``ConversionService``) and *measures* the produced
DOCX to prove the integrated capabilities still compose:

  A  English professional report: TOC + normal/wide tables + figure +
     explicit ``image_width`` + (no) diagnostics.
  B  Chinese / mixed report: localized TOC + table/figure fitting.
  C  The same representative document across all five output profiles.
  D  Boundary case: irreducible wide table (fit floor) + page-capped figure
     (height limit), with fit-reason evidence from the single authorities.

Asserted: no content/cell mutation, correct TOC heading + field, table floor
preserved, aspect ratio preserved, effective profile geometry respected,
``image_width`` respected/capped, diagnostics QA-derived, deterministic repeat,
and no profile-name branching.

Exit code 0 == every check passed.
"""

from __future__ import annotations

import base64
import contextlib
import hashlib
import io
import re
import struct
import sys
import tempfile
import zipfile
import zlib
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from docx import Document

from md_converter.application.conversion_request import ConversionRequest
from md_converter.application.conversion_service import ConversionService
from md_converter.profiles import all_profiles
from md_converter.renderer.layout.figure_sizing import plan_figure_fit
from md_converter.renderer.layout.table_fitting import plan_table_fit

A4_WIDTH_CM = 21.0
A4_HEIGHT_CM = 29.7
TABLE_MIN_COLUMN_CM = 1.2
TABLE_READABILITY_FLOOR_PT = 8.5

#: DOCX stores lengths in twips, so a column written as exactly 1.2 cm reads
#: back as ~1.1994 cm.  The repository's own guards use a 1e-3 cm tolerance.
MEASURE_TOL_CM = 1e-3


# --------------------------------------------------------------------------
# deterministic PNG fixture (no third-party dependency)
# --------------------------------------------------------------------------
def png_data_uri(width: int, height: int) -> str:
    """Return a deterministic solid-colour PNG data URI of ``width`` x ``height``."""
    scanline = b"\x00" + bytes((0x2F, 0x54, 0x96)) * width
    raw = scanline * height

    def chunk(tag: bytes, payload: bytes) -> bytes:
        body = tag + payload
        return (
            struct.pack(">I", len(payload)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    png = (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    return "data:image/png;base64," + base64.b64encode(png).decode("ascii")


def _wide_table(columns: int, cell_text: str) -> str:
    """Return a Markdown table with ``columns`` columns (irreducible when wide)."""
    header = "|" + "|".join(f" H{index} " for index in range(columns)) + "|"
    separator = "|" + "|".join(" --- " for _ in range(columns)) + "|"
    row = "|" + "|".join(f" {cell_text} " for _ in range(columns)) + "|"
    return "\n".join([header, separator, row, ""])


# --------------------------------------------------------------------------
# fixtures
# --------------------------------------------------------------------------
EN_REPORT = f"""---
title: Competitive Foundation Report
---

# Competitive Foundation Report

Integrated source-level verification of the competitive foundation.

## Architecture

| Component | Authority | Status |
| --- | --- | --- |
| Parser | markdown-it | frozen |
| Renderer | WordRenderer | frozen |
| QA | compiler | frozen |

## Financial Overview

| Region | Q1 | Q2 | Q3 | Q4 | Notes | Owner | Status |
| --- | --- | --- | --- | --- | --- | --- | --- |
| North | 120 | 130 | 140 | 150 | steady | A. Lee | accepted |
| South | 98 | 105 | 111 | 118 | recovering | B. Ng | accepted |
| East | 210 | 214 | 219 | 225 | expanding | C. Tan | accepted |
| West | 64 | 70 | 75 | 81 | new | D. Rao | accepted |

## Projected Growth

![Growth chart]({png_data_uri(400, 100)})

### Notes

Deterministic verification only.
"""

EN_TABLE_CELLS = [
    ["Component", "Authority", "Status"],
    ["Parser", "markdown-it", "frozen"],
    ["Renderer", "WordRenderer", "frozen"],
    ["QA", "compiler", "frozen"],
]

CN_REPORT = f"""---
title: 竞争基础集成验证
---

# 竞争基础集成验证

本报告用于验证竞争基础能力的集成一致性。

## 架构

| 组件 | 权威 | 状态 |
| --- | --- | --- |
| 解析器 | markdown-it | 冻结 |
| 渲染器 | WordRenderer | 冻结 |
| 质量门 | compiler | 冻结 |

## 图形

![增长图]({png_data_uri(300, 150)})

### 备注

仅用于确定性验证。
"""

BOUNDARY = f"""---
title: Boundary Cases
---

# Boundary Cases

## Irreducible Table

{_wide_table(25, "a")}

## Page Capped Figure

![Tall figure]({png_data_uri(20, 200)})
"""


# --------------------------------------------------------------------------
# conversion + measurement helpers
# --------------------------------------------------------------------------
BASE_OVERRIDES: Dict[str, Any] = {
    "word_com": False,  # sandbox has no Word; TOC field is still inserted
    "enable_cover": True,
    "toc": True,
    "style_tables": True,
    "verbose": False,
}


def convert(
    workspace: Path,
    name: str,
    body: str,
    overrides: Optional[Dict[str, Any]] = None,
) -> Tuple[Any, Path]:
    """Convert ``body`` via ``ConversionService`` and return ``(result, docx)``."""
    source = workspace / f"{name}.md"
    source.write_text(body, encoding="utf-8")
    output = workspace / f"{name}.docx"
    config = dict(BASE_OVERRIDES)
    config.update(overrides or {})
    service = ConversionService()
    with contextlib.redirect_stdout(io.StringIO()):
        result = service.convert(
            ConversionRequest(source_path=source, output_path=output, config_overrides=config)
        )
    return result, output


def doc_xml(path: Path) -> str:
    """Return ``word/document.xml`` text of a DOCX."""
    with zipfile.ZipFile(path) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def document_xml_sha256(path: Path) -> str:
    """Return the SHA-256 of ``word/document.xml`` (determinism probe)."""
    return hashlib.sha256(doc_xml(path).encode("utf-8")).hexdigest()


def normalize_document_xml(xml: str) -> str:
    """Mask the data-URI picture name (a per-run temp stem), nothing else."""
    return re.sub(r'(<pic:cNvPr id="\d+" name=")[^"]*(")', r"\1<image>\2", xml)


def convert_file_image(
    workspace: Path, name: str, body: str, overrides: Optional[Dict[str, Any]] = None
) -> Tuple[Any, Path]:
    """Convert ``body`` with its single data-URI image replaced by a file image."""
    uri_match = re.search(r"!\[[^\]]*\]\((data:image/png;base64,[^)]+)\)", body)
    if uri_match is None:
        raise ValueError("fixture has no data-URI image to externalise")
    png_bytes = base64.b64decode(uri_match.group(1).split(",", 1)[1])
    (workspace / "growth.png").write_bytes(png_bytes)
    return convert(workspace, name, body.replace(uri_match.group(1), "growth.png"), overrides)


def toc_heading(doc: Any) -> Optional[str]:
    """Return the TOC heading paragraph text, when present."""
    for paragraph in doc.paragraphs:
        if paragraph.style is not None and paragraph.style.name == "TOC Heading":
            return paragraph.text.strip()
    return None


def toc_entry_texts(doc: Any) -> List[str]:
    """Return the cached TOC entry texts (styles ``TOC 1``..``TOC 3``)."""
    entries: List[str] = []
    for paragraph in doc.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("TOC ") and style != "TOC Heading":
            entries.append(paragraph.text.strip())
    return entries


def toc_field_present(path: Path) -> bool:
    """Return whether the canonical TOC field instruction is present."""
    xml = doc_xml(path)
    return bool(re.search(r'TOC \\o "1-3" \\h \\z \\u', xml))


def table_matrix(table: Any) -> List[List[str]]:
    """Return the visible cell text matrix of a table."""
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def table_widths_cm(table: Any) -> List[float]:
    """Return the explicit column widths (cm) of a table."""
    return [round(column.width.cm, 4) for column in table.columns]


def table_layout_is_fixed(path: Path) -> bool:
    """Return whether the artifact requests fixed table layout."""
    return bool(re.search(r'w:tblLayout[^>]*w:type="fixed"', doc_xml(path)))


def section_margins_cm(doc: Any) -> Tuple[float, float, float, float]:
    """Return ``(left, right, top, bottom)`` margins in cm of the last section."""
    section = doc.sections[-1]
    return (
        round(section.left_margin.cm, 3),
        round(section.right_margin.cm, 3),
        round(section.top_margin.cm, 3),
        round(section.bottom_margin.cm, 3),
    )


def content_box_cm(doc: Any) -> Tuple[float, float]:
    """Return the effective content box ``(width, height)`` in cm."""
    section = doc.sections[-1]
    width = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
    height = section.page_height.cm - section.top_margin.cm - section.bottom_margin.cm
    return round(width, 3), round(height, 3)


def diagnostics_by_code(result: Any) -> Dict[str, List[Any]]:
    """Group application diagnostics by canonical code."""
    grouped: Dict[str, List[Any]] = {}
    for record in result.diagnostics:
        grouped.setdefault(record.code, []).append(record)
    return grouped


# --------------------------------------------------------------------------
# checks
# --------------------------------------------------------------------------
class Report:
    """Collect and print PASS/FAIL checks."""

    def __init__(self) -> None:
        self.rows: List[Tuple[str, bool, str]] = []

    def check(self, label: str, ok: bool, detail: str = "") -> None:
        """Record one check."""
        self.rows.append((label, bool(ok), detail))

    def print(self) -> int:
        """Print the report and return the number of failures."""
        print("R2-V01 WP-03 CROSS-FEATURE DOCUMENT INTEGRATION")
        print("=" * 72)
        failures = 0
        for label, ok, detail in self.rows:
            print(f"[{'PASS' if ok else 'FAIL'}] {label}")
            if detail:
                print(f"       {detail}")
            failures += 0 if ok else 1
        print("=" * 72)
        print(f"checks={len(self.rows)} failed={failures}")
        return failures


def check_scenario_a(workspace: Path, report: Report) -> None:
    """English professional report: TOC + normal/wide tables + figure + image_width."""
    result, output = convert(workspace, "en_report", EN_REPORT, {"image_width": 4})
    doc = Document(str(output))

    report.check("A1 status is SUCCESS (no warnings)", result.is_success, f"status={result.status}")
    report.check(
        "A2 TOC heading is the English title",
        toc_heading(doc) == "Table of Contents",
        f"heading={toc_heading(doc)!r}",
    )
    report.check("A3 canonical TOC field present", toc_field_present(output))
    entries = toc_entry_texts(doc)
    report.check(
        "A4 TOC caches the document headings",
        {"Competitive Foundation Report", "Architecture", "Financial Overview", "Projected Growth"}
        <= {entry.rstrip() for entry in entries},
        f"entries={entries}",
    )

    normal = table_matrix(doc.tables[0])
    report.check(
        "A5 normal table cells unchanged",
        normal == EN_TABLE_CELLS,
        f"cells={normal}",
    )

    content_w, content_h = content_box_cm(doc)
    wide = doc.tables[1]
    widths = table_widths_cm(wide)
    report.check(
        "A6 wide table has 8 explicit columns",
        len(widths) == 8,
        f"widths={widths}",
    )
    report.check(
        "A7 wide table total width <= effective content width",
        sum(widths) <= content_w + 0.01,
        f"total={round(sum(widths),4)} content={content_w}",
    )
    report.check(
        "A8 wide table respects the 1.2cm column floor",
        min(widths) >= TABLE_MIN_COLUMN_CM - MEASURE_TOL_CM,
        f"min_column={min(widths)}",
    )
    report.check("A9 fixed table layout requested", table_layout_is_fixed(output))
    report.check(
        "A10 wide table keeps every cell (no mutation)",
        table_matrix(wide)[1][0] == "North" and len(wide.rows) == 5 and len(wide.columns) == 8,
        f"rows={len(wide.rows)} cols={len(wide.columns)} first_cell={table_matrix(wide)[1][0]!r}",
    )

    shapes = list(doc.inline_shapes)
    report.check("A11 exactly one figure rendered", len(shapes) == 1, f"figures={len(shapes)}")
    shape = shapes[0]
    target = min(4 * 2.54, content_w)
    report.check(
        "A12 figure width == min(image_width, content width)",
        abs(shape.width.cm - target) < 0.05,
        f"width={round(shape.width.cm,3)} target={round(target,3)}",
    )
    report.check(
        "A13 figure aspect ratio preserved (4:1)",
        abs(shape.width.cm / shape.height.cm - 4.0) < 0.02,
        f"ratio={round(shape.width.cm / shape.height.cm, 4)}",
    )
    report.check(
        "A14 figure stays inside the content box",
        shape.width.cm <= content_w + 0.05 and shape.height.cm <= content_h + 0.05,
        f"figure={round(shape.width.cm,3)}x{round(shape.height.cm,3)} box={content_w}x{content_h}",
    )
    report.check(
        "A15 no diagnostics for a clean document",
        len(result.diagnostics) == 0,
        f"codes={[record.code for record in result.diagnostics]}",
    )


def check_scenario_b(workspace: Path, report: Report) -> None:
    """Chinese / mixed report: localized TOC + table/figure fitting."""
    result, output = convert(workspace, "cn_report", CN_REPORT, {"image_width": 4})
    doc = Document(str(output))

    report.check("B1 status is SUCCESS", result.is_success, f"status={result.status}")
    report.check(
        "B2 TOC heading localized to Chinese",
        toc_heading(doc) == "目录",
        f"heading={toc_heading(doc)!r}",
    )
    report.check("B3 canonical TOC field present", toc_field_present(output))
    entries = toc_entry_texts(doc)
    report.check(
        "B4 TOC caches the Chinese headings",
        any("竞争基础集成验证" in entry for entry in entries) and any("架构" in entry for entry in entries),
        f"entries={entries}",
    )

    content_w, _ = content_box_cm(doc)
    table = doc.tables[0]
    widths = table_widths_cm(table)
    report.check(
        "B5 Chinese table cells unchanged",
        table_matrix(table)
        == [["组件", "权威", "状态"], ["解析器", "markdown-it", "冻结"], ["渲染器", "WordRenderer", "冻结"], ["质量门", "compiler", "冻结"]],
        f"cells={table_matrix(table)}",
    )
    report.check(
        "B6 Chinese table fitted inside the content width",
        sum(widths) <= content_w + 0.01 and min(widths) >= TABLE_MIN_COLUMN_CM - MEASURE_TOL_CM,
        f"total={round(sum(widths),4)} min={min(widths)} content={content_w}",
    )
    shapes = list(doc.inline_shapes)
    report.check("B7 Chinese report figure rendered", len(shapes) == 1, f"figures={len(shapes)}")
    shape = shapes[0]
    report.check(
        "B8 Chinese report figure aspect ratio preserved (2:1)",
        abs(shape.width.cm / shape.height.cm - 2.0) < 0.02,
        f"ratio={round(shape.width.cm / shape.height.cm, 4)}",
    )


def check_scenario_c(workspace: Path, report: Report) -> None:
    """The same representative document across all five output profiles."""
    for profile in all_profiles():
        result, output = convert(
            workspace,
            f"profile_{profile.id}",
            EN_REPORT,
            {"image_width": 4, "output_profile": profile.id},
        )
        doc = Document(str(output))
        left, right, top, bottom = section_margins_cm(doc)
        expected = (
            profile.presentation.page_margins.left_cm,
            profile.presentation.page_margins.right_cm,
            profile.presentation.page_margins.top_cm,
            profile.presentation.page_margins.bottom_cm,
        )
        report.check(
            f"C[{profile.id}] effective page margins == declared profile geometry",
            all(abs(a - b) < 0.02 for a, b in zip((left, right, top, bottom), expected)),
            f"artifact={(left, right, top, bottom)} declared={expected}",
        )
        content_w, content_h = content_box_cm(doc)
        table_ok = all(
            sum(table_widths_cm(table)) <= content_w + 0.01
            for table in doc.tables
        )
        shape = list(doc.inline_shapes)[0]
        figure_ok = (
            abs(shape.width.cm - min(4 * 2.54, content_w)) < 0.05
            and shape.width.cm <= content_w + 0.05
            and shape.height.cm <= content_h + 0.05
        )
        report.check(
            f"C[{profile.id}] tables and figure respect the profile content box",
            table_ok and figure_ok,
            f"content_box={content_w}x{content_h} figure={round(shape.width.cm,3)}x{round(shape.height.cm,3)}",
        )
        report.check(
            f"C[{profile.id}] TOC heading unchanged by profile",
            toc_heading(doc) == "Table of Contents" and result.is_success,
            f"heading={toc_heading(doc)!r} status={result.status}",
        )


def check_scenario_d(workspace: Path, report: Report) -> None:
    """Boundary case: irreducible table floor + page-capped figure, with reasons."""
    result, output = convert(workspace, "boundary", BOUNDARY, {"image_width": 4})
    doc = Document(str(output))
    content_w, content_h = content_box_cm(doc)

    table = doc.tables[0]
    widths = table_widths_cm(table)
    matrix = table_matrix(table)
    report.check(
        "D1 irreducible table keeps all 25 columns and every cell",
        len(widths) == 25 and len(matrix[0]) == 25 and matrix[0][0] == "H0",
        f"columns={len(widths)} first_cell={matrix[0][0]!r}",
    )
    report.check(
        "D2 irreducible table columns sit exactly at the 1.2cm floor",
        all(abs(width - TABLE_MIN_COLUMN_CM) < MEASURE_TOL_CM for width in widths),
        f"min={min(widths)} max={max(widths)}",
    )
    plan = plan_table_fit(
        [[f"H{i}" for i in range(25)], ["a" for _ in range(25)]],
        content_width_cm=content_w,
        base_font_size_pt=9.5,
    )
    report.check(
        "D3 table fit-reason: squeezed=True (floor reached, content kept)",
        plan.squeezed and plan.total_width_cm >= content_w,
        f"plan={plan.to_dict()}",
    )
    report.check(
        "D4 bounded font step stays at/above the 8.5pt readability floor",
        plan.font_size_pt >= TABLE_READABILITY_FLOOR_PT,
        f"font_pt={plan.font_size_pt}",
    )

    shape = list(doc.inline_shapes)[0]
    intrinsic = (20, 200)
    fit = plan_figure_fit(
        intrinsic_width_px=intrinsic[0],
        intrinsic_height_px=intrinsic[1],
        target_width_cm=min(4 * 2.54, content_w),
        content_width_cm=content_w,
        content_height_cm=content_h,
    )
    report.check(
        "D5 figure fit-reason: height-limited by the effective content height",
        fit.height_limited and not fit.width_limited,
        f"plan={fit.to_dict()}",
    )
    report.check(
        "D6 page-capped figure height == content height, aspect preserved",
        abs(shape.height.cm - content_h) < 0.05
        and abs(shape.width.cm / shape.height.cm - intrinsic[0] / intrinsic[1]) < 1e-3,
        f"figure={round(shape.width.cm,3)}x{round(shape.height.cm,3)} box_h={content_h}",
    )
    report.check(
        "D7 delivered figure never exceeds the effective content box",
        shape.width.cm <= content_w + 0.05 and shape.height.cm <= content_h + 0.05,
        f"figure={round(shape.width.cm,3)}x{round(shape.height.cm,3)}",
    )

    codes = diagnostics_by_code(result)
    render006 = codes.get("RENDER006", [])
    fit_details = render006[0].details if render006 else None
    report.check(
        "D8 irreducible wide table surfaces the RENDER006 warning",
        bool(render006),
        f"codes={sorted(codes)}",
    )
    report.check(
        "D9 RENDER006 carries the single-authority fit evidence",
        bool(fit_details)
        and fit_details.get("squeezed") is True
        and abs(fit_details.get("min_column_width_cm", 0) - TABLE_MIN_COLUMN_CM) < 1e-6,
        f"details={fit_details}",
    )
    report.check(
        "D10 warning is QA-derived (rendered_qa record present)",
        any(record.stage == "rendered_qa" for record in result.diagnostics)
        and result.quality_gate_report is not None,
        f"stages={sorted({record.stage for record in result.diagnostics if record.stage})}",
    )
    gate = result.quality_gate_report or {}
    report.check(
        "D11 every quality-gate stage is non-FAIL",
        all(
            (gate.get(stage) or {}).get("status") not in (None, "FAIL")
            for stage in ("static_qa", "rendered_qa", "post_processor", "final_artifact_qa")
        ),
        f"stages={ {s: (gate.get(s) or {}).get('status') for s in ('static_qa','rendered_qa','post_processor','final_artifact_qa')} }",
    )


def check_determinism(workspace: Path, report: Report) -> None:
    """Repeated identical conversion preserves the produced document XML."""
    _, first = convert(workspace, "det_en_1", EN_REPORT, {"image_width": 4})
    _, second = convert(workspace, "det_en_2", EN_REPORT, {"image_width": 4})
    _, boundary_1 = convert(workspace, "det_bd_1", BOUNDARY, {"image_width": 4})
    _, boundary_2 = convert(workspace, "det_bd_2", BOUNDARY, {"image_width": 4})
    _, file_1 = convert_file_image(workspace, "det_file_1", EN_REPORT, {"image_width": 4})
    _, file_2 = convert_file_image(workspace, "det_file_2", EN_REPORT, {"image_width": 4})

    report.check(
        "E1 file-image report: repeated conversion is byte-identical (document.xml)",
        document_xml_sha256(file_1) == document_xml_sha256(file_2),
        f"sha={document_xml_sha256(file_1)[:16]}",
    )
    report.check(
        "E2 data-URI report: repeated conversion is identical apart from the picture name",
        normalize_document_xml(doc_xml(first)) == normalize_document_xml(doc_xml(second)),
        f"sha={document_xml_sha256(first)[:16]}",
    )
    report.check(
        "E3 boundary report: repeated conversion is identical apart from the picture name",
        normalize_document_xml(doc_xml(boundary_1))
        == normalize_document_xml(doc_xml(boundary_2)),
        f"sha={document_xml_sha256(boundary_1)[:16]}",
    )


def check_no_profile_branching(report: Report) -> None:
    """The renderer contains no output-profile identifier branch."""
    renderer_root = Path(__file__).resolve().parents[5] / "md_converter" / "renderer"
    offenders: List[str] = []
    for path in sorted(renderer_root.rglob("*.py")):
        text = path.read_text(encoding="utf-8")
        if re.search(r"from\s+\.+profiles\s+import|md_converter\.profiles", text):
            offenders.append(path.name)
    report.check(
        "F1 renderer imports no output-profile module (no profile-ID branch)",
        not offenders,
        f"offenders={offenders}",
    )


def main() -> int:
    """Run every WP-03 check."""
    report = Report()
    workspace = Path(tempfile.mkdtemp(prefix="r2v01_wp03_"))
    check_scenario_a(workspace, report)
    check_scenario_b(workspace, report)
    check_scenario_c(workspace, report)
    check_scenario_d(workspace, report)
    check_determinism(workspace, report)
    check_no_profile_branching(report)
    print(f"workspace={workspace}")
    return 1 if report.print() else 0


if __name__ == "__main__":
    sys.exit(main())
